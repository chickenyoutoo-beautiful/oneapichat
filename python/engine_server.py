#!/usr/bin/env python3
"""
OneAPIChat 后台引擎 - 心跳 / Cron / 子代理
"""
import asyncio
import json
import os
import sys
import time
import threading
import uuid
import subprocess
import requests
import re
from datetime import datetime, timedelta
from pathlib import Path
import sqlite3
import tempfile
import glob
from collections import deque

# ── 代理配置 ────────────────────────────────────────────
def _load_proxy_config():
    """从用户配置中加载可用代理，不修改进程级环境变量。"""
    try:
        # 读取 localStorage 持久化的配置
        import glob as _glob
        config_files = _glob.glob(os.path.join(PROJECT_ROOT, 'chat_data/config_user_*.json'))
        for cf in config_files:
            try:
                with open(cf, 'r') as f:
                    cfg = json.load(f)
                if cfg.get('proxyEnabled') == '1' and cfg.get('proxyUrl'):
                    proxy_url = cfg['proxyUrl']
                    print('[Engine] 出站代理已启用')
                    return proxy_url
            except Exception:
                pass
        # ★ 回退: proxyEnabled=1 但无 proxyUrl → 自动使用 ECS1→GCP 私网链路
        for cf in config_files:
            try:
                with open(cf, 'r') as f:
                    cfg = json.load(f)
                if cfg.get('proxyEnabled') == '1':
                    # ECS1:8890 只绑定 ZeroTier 地址，并透传到 WireGuard 内的 GCP HTTP 代理。
                    _gcp_proxy = 'http://192.168.195.226:8890'
                    print('[Engine] proxyEnabled=1, 自动路由已启用')
                    return _gcp_proxy
            except Exception:
                pass
    except Exception as e:
        print(f'[Engine] proxy configuration load failed: {type(e).__name__}')
    return None

# Cross-platform: fcntl is Unix-only
try:
    import fcntl
    HAS_FCNTL = True
except ImportError:
    HAS_FCNTL = False

# ── Project root detection ────────────────────────────
PROJECT_ROOT = str(Path(__file__).parent.parent.resolve())
_SYNC_TRACE_LOG = Path(PROJECT_ROOT) / 'logs' / 'multidevice-sync.jsonl'
_SYNC_TRACE_LOCK = threading.Lock()

def _sync_trace(stage: str, data: dict | None = None):
    data = dict(data or {})
    trace_id = str(data.get('trace_id') or '')[:96]
    if not trace_id:
        return
    entry = {
        'ts': int(time.time() * 1000),
        'component': 'engine_server.py',
        'stage': stage,
        'trace_id': re.sub(r'[^a-zA-Z0-9_.:-]', '', trace_id),
    }
    for key, value in data.items():
        if key in {'content', 'text', 'messages', 'files', 'body', 'token'}:
            continue
        if value is None or isinstance(value, (str, int, float, bool)):
            entry[key] = value
    try:
        _SYNC_TRACE_LOG.parent.mkdir(parents=True, exist_ok=True)
        line = json.dumps(entry, ensure_ascii=False, separators=(',', ':'))
        with _SYNC_TRACE_LOCK:
            with _SYNC_TRACE_LOG.open('a', encoding='utf-8') as fh:
                fh.write(line + '\n')
    except Exception:
        pass

# ★ _PROXY_URL 依赖 PROJECT_ROOT，必须在 PROJECT_ROOT 之后初始化
_PROXY_URL = _load_proxy_config()

# ── 全局 Session (带代理) ──────────────────────────────
# ★ 无代理 Session: 引擎访问自身(127.0.0.1:8766)子代理转发 server_* 工具时使用
#   避免出站代理导致 localhost 自调用 RemoteDisconnected
_http_session_no_proxy = requests.Session()
_http_session_no_proxy.trust_env = False

def _get_proxies():
    """获取请求代理字典"""
    if _PROXY_URL:
        return {'http': _PROXY_URL, 'https': _PROXY_URL}
    return None

def _bypass_auto_proxy_for_base(base_url: str) -> bool:
    """自有 API 入口一律直连，避免 primary→ECS1→GCP→ECS1 回环。"""
    try:
        from urllib.parse import urlparse
        host = (urlparse(base_url or '').hostname or '').lower()
        return host in {'gpt.naujtrats.xyz', 'aliyun.naujtrats.xyz'}
    except Exception:
        return False

def _requires_auto_proxy(base_url: str) -> bool:
    """仅对国内通常无法稳定直连的境外 API 启用 ECS1→GCP。"""
    if not _PROXY_URL or _bypass_auto_proxy_for_base(base_url):
        return False
    try:
        from urllib.parse import urlparse
        host = (urlparse(base_url or '').hostname or '').lower().rstrip('.')
    except Exception:
        return False
    restricted = (
        'openai.com', 'chatgpt.com', 'oaistatic.com', 'oaiusercontent.com',
        'anthropic.com', 'claude.ai', 'x.ai', 'openrouter.ai',
        'googleapis.com', 'google.com', 'googleusercontent.com',
        'nvidia.com',
        'api.search.brave.com', 'brave.com', 'tavily.com', 'api.tavily.com',
    )
    return any(host == domain or host.endswith('.' + domain) for domain in restricted)

class _SelectiveProxySession(requests.Session):
    """共享工具 Session 按目标域名分流，不设置进程级代理。"""
    def __init__(self):
        super().__init__()
        self.trust_env = False

    def request(self, method, url, **kwargs):
        if 'proxies' not in kwargs and _requires_auto_proxy(url):
            kwargs['proxies'] = {'http': _PROXY_URL, 'https': _PROXY_URL}
        return super().request(method, url, **kwargs)

# ── 全局 Session（域名分流）──
_http_session = _SelectiveProxySession()

def _repair_tool_json(raw):
    """容错修复工具参数JSON: 未转义引号 → \\\"、未转义换行 → \\\\n、截断 → 补齐引号/花括号"""
    if not isinstance(raw, str):
        return '{}'
    out = []
    in_str = False
    had_inner = False
    i = 0
    n = len(raw)
    while i < n:
        ch = raw[i]
        if not in_str:
            out.append(ch)
            if ch == '"':
                in_str = True
                had_inner = False
            i += 1
            continue
        if ch == '\\':
            out.append(ch)
            if i + 1 < n:
                out.append(raw[i + 1])
                i += 2
            else:
                i += 1
            continue
        if ch == '\n':
            out.append('\\n')
            i += 1
            continue
        if ch == '\r':
            out.append('\\r')
            i += 1
            continue
        if ch == '\t':
            out.append('\\t')
            i += 1
            continue
        if ord(ch) < 32:
            out.append(' ')
            i += 1
            continue
        if ch == '"':
            j = i + 1
            while j < n and raw[j] in ' \t':
                j += 1
            nxt = raw[j] if j < n else ''
            if nxt == ',':
                k = j + 1
                while k < n and raw[k] in ' \t':
                    k += 1
                after_comma = raw[k] if k < n else ''
                if after_comma in ('"', '}', ']', ''):
                    if had_inner:
                        # 值内引号对刚闭合, 逗号前补一个真正的JSON收尾引号
                        out.append('\\"')
                        out.append('"')
                        in_str = False
                    else:
                        out.append('"')
                        in_str = False
                else:
                    out.append('\\"')
                    had_inner = True
            elif nxt in ('}', ']', ':', ''):
                out.append('"')
                in_str = False
            else:
                out.append('\\"')
                had_inner = True
            i += 1
            continue
        out.append(ch)
        i += 1
    if in_str:
        out.append('"')
    fixed = ''.join(out)
    ob = fixed.count('{')
    cb = fixed.count('}')
    if cb < ob:
        fixed += '}' * (ob - cb)
    return fixed

def _tolerant_tool_args(tool_name, raw):
    """模型输出的工具参数JSON不合法时,尽量恢复 server_exec/server_python 的命令/脚本(兼容未转义引号)"""
    if not isinstance(raw, str):
        return None
    try:
        data = json.loads(_repair_tool_json(raw))
        if isinstance(data, dict):
            if tool_name == "server_exec":
                val = data.get("cmd") or data.get("command") or data.get("query")
            elif tool_name == "server_python":
                val = data.get("script") or data.get("code") or data.get("cmd")
            if val:
                return {("cmd" if tool_name == "server_exec" else "script"): str(val)}
            return data
    except Exception:
        pass
    if tool_name == "server_exec":
        keys = ["cmd", "command", "query"]
    elif tool_name == "server_python":
        keys = ["script", "code", "cmd"]
    else:
        return None
    target = "cmd" if tool_name == "server_exec" else "script"
    # 1) 标准 JSON 转义的值
    for key in keys:
        m = re.search(r'"' + key + r'"\s*:\s*"((?:[^"\\]|\\.)*)"', raw)
        if m:
            val = m.group(1)
            return {target: val.replace('\\n', '\n').replace('\\r', '\r').replace('\\"', '"').replace('\\\\', '\\')}
    # 2) 容错: 值内含未转义的引号 → 取 "key": " 之后全部内容, 从末尾找真正的收尾引号
    for key in keys:
        m2 = re.search(r'"' + key + r'"\s*:\s*"([\s\S]*)$', raw)
        if not m2:
            continue
        val = m2.group(1).rstrip()
        cut = -1
        for i in range(len(val) - 1, -1, -1):
            if val[i] == '"':
                rest = val[i+1:].lstrip()
                if rest == '' or rest.startswith('}') or rest.startswith(','):
                    cut = i
                    break
        if cut >= 0:
            val = val[:cut]
        val = val.strip()
        return {target: val.replace('\\n', '\n').replace('\\r', '\r').replace('\\"', '"').replace('\\\\', '\\')}
    return None
sys.path.insert(0, PROJECT_ROOT)
sys.path.insert(0, str(Path(__file__).parent.resolve()))
sys.path.insert(0, os.path.join(tempfile.gettempdir(), 'pylib'))

try:
    from fastapi import FastAPI, Query, HTTPException, Request, WebSocket, WebSocketDisconnect
    from fastapi.responses import StreamingResponse, JSONResponse
    from fastapi.middleware.cors import CORSMiddleware
    import uvicorn
except Exception:
    print("[引擎] 需要安装 fastapi/uvicorn: pip install fastapi uvicorn --break-system-packages")
    sys.exit(1)

# ── 引擎层模块 ────────────────────────────────────────────
from engine.exec_policy import ExecPolicy, ExecDecision, Priority
from engine.speculation import SpeculationEngine, SpeculationState
from engine.retry import RetryEngine, RetryStatus
from engine.tool_registry import ToolRegistry, ToolDef, Capability, ApprovalKind, get_global_registry
from engine.event_frame import EventFlowBuilder, EventType, EventLog
from engine.store import (EngineStore, ChatStore, get_ns as _store_get_ns,
                          get_chat_store as _store_get_chat_store, harden_all_chat_stores)
from engine.rag_engine import (rag_list_collections, rag_create_collection, rag_delete_collection,
                                rag_upload_document, rag_search, rag_list_documents, rag_delete_document,
                                _get_embedding, _load_json, _save_json, _docs_path, RAG_DIR)
from engine.video_edit import (SUBTITLE_FONTS, DEFAULT_FONT, generate_srt as _video_generate_srt,
    str_to_rgb, color_to_ass, ypos_to_alignment, hex_to_rgba, draw_rounded_rect, init_video_context,
    _apply_subtitle, _apply_filter, _apply_transition, _apply_tts,
    _apply_voice_to_video, _apply_crop, _apply_reverse, _apply_mute,
    _apply_bgm, _apply_enhance, _apply_gif, _apply_silent_cut,
    _apply_subtitle_style, _apply_ffmpeg_filter, _apply_ffmpeg_transition, _apply_compose,
    _apply_stt, _apply_stt_to_timeline)
from engine.cron import _run_cron_job, _start_cron_job as _cron_start, _stop_cron_job as _cron_stop
from engine.agent_roles import AGENT_ROLES, filter_tools_by_role as _filter_tools_by_role, cleanup_old_agents as _cleanup_old_agents
from engine.server_tools import register_server_tools
from engine.agent_endpoints import register_agent_endpoints
from engine.memory_endpoints import register_memory_endpoints
from engine.crypto import load_encryption_key, get_aes_key, decrypt_xor
from engine.agent_memory import read_memory_json, write_memory_json
from engine.workflow import create_workflow, run_workflow, list_workflows, status_workflow, delete_workflow, get_roles as _wf_get_roles
from engine.agent_errors import AgentRuntimeError, ErrorCode, classify_provider_error, redact_secrets
from engine.provider_runtime import (detect_provider, normalize_usage, prepare_openai_request,
                                     safe_request_snapshot, complete_chat)
from engine.runtime_auth import EngineAuthMiddleware, get_internal_bridge_secret, require_scope_owner
from engine.runtime_store import get_agent_runtime_store
from engine.runtime_api import register_runtime_endpoints
from engine.tool_pipeline import ToolExecutionContext, ToolExecutionPipeline
from engine.stream_retention import compact_stream_files
from engine.self_description import admit_self_context
from engine.resource_guard import ResourceOwnerGuard
from engine.goal_runner import DurableGoalRunner
from engine.browser import prune_browser_managers
from engine.upload_retention import cleanup_uploads
from engine.chat_projection import ChatProjectionStore

_INTERNAL_BRIDGE_SECRET = get_internal_bridge_secret(PROJECT_ROOT)
resource_owners = ResourceOwnerGuard(PROJECT_ROOT, Path(PROJECT_ROOT) / '.engine')
if _INTERNAL_BRIDGE_SECRET:
    _http_session_no_proxy.headers.update({'X-OneAPIChat-Internal': _INTERNAL_BRIDGE_SECRET})

app = FastAPI(title="OneAPIChat Engine")
app.add_middleware(CORSMiddleware, allow_origins=[
    "https://naujtrats.xyz",
    "https://www.naujtrats.xyz",
    "https://localmodels.naujtrats.xyz",
], allow_credentials=True, allow_methods=["*"], allow_headers=["*"])
# Public /engine routes must authenticate against the same session source as PHP.
# Trusted loopback calls without reverse-proxy headers remain available to engine_api.php.
app.add_middleware(
    EngineAuthMiddleware,
    project_root=PROJECT_ROOT,
    public_paths={
        "/engine/health",
        "/engine/stock_realtime",
        "/engine/stock_market_overview",
        "/engine/stock_kline",
        "/engine/stock_sector_flow",
        "/engine/stock_dragon_tiger",
        "/engine/stock_north_flow",
        "/engine/stock_diagnosis",
        "/engine/stock_indicators",
        "/engine/stock_chart",
    },
)

ENGINE_DIR = Path(PROJECT_ROOT) / ".engine"
ENGINE_DIR.mkdir(parents=True, exist_ok=True)
STREAM_DIR = ENGINE_DIR / "streams"
STREAM_DIR.mkdir(parents=True, exist_ok=True)
_stream_maintenance = compact_stream_files(STREAM_DIR, retention_days=30, max_chunks=2048)
if _stream_maintenance.get('deleted') or _stream_maintenance.get('compacted') or _stream_maintenance.get('errors'):
    print(f"[Engine] stream maintenance: {_stream_maintenance}")
_chat_store_hardening = harden_all_chat_stores(ENGINE_DIR, retention_days=30)
if _chat_store_hardening.get('scrubbed') or _chat_store_hardening.get('pruned') or _chat_store_hardening.get('errors'):
    print(f"[Engine] chat store migration: {_chat_store_hardening}")
TEMP_DIR = Path(tempfile.gettempdir())
init_video_context(PROJECT_ROOT, TEMP_DIR, _http_session)

# ── 引擎层全局实例 ──────────────────────────────────────────
_exec_policies: dict[str, ExecPolicy] = {}
_runtime_cache_access: dict[str, float] = {}
_exec_policies_lock = threading.RLock()

def _get_exec_policy(user_id: str) -> ExecPolicy:
    owner = str(user_id or '').strip()
    if not owner:
        raise AgentRuntimeError(ErrorCode.UNAUTHORIZED, 'authenticated user required', status=401)
    key = uuid.uuid5(uuid.NAMESPACE_URL, 'oneapichat-exec-policy:' + owner).hex
    with _exec_policies_lock:
        _runtime_cache_access[owner] = time.monotonic()
        policy = _exec_policies.get(owner)
        if policy is None:
            path = ENGINE_DIR / f'exec_policy_{key}.json'
            legacy = ENGINE_DIR / 'exec_policy.json'
            if not path.exists() and legacy.exists() and resource_owners.owner('global_runtime_admin') == owner:
                path.write_bytes(legacy.read_bytes())
                os.chmod(path, 0o600)
            policy = ExecPolicy(rules_file=str(path))
            if policy._rules_file and not policy._rules_file.exists():
                policy.save()
            try:
                os.chmod(path, 0o600)
            except OSError:
                pass
            _exec_policies[owner] = policy
        return policy

_speculation_engines: dict[str, SpeculationEngine] = {}
_retry_engines: dict[str, RetryEngine] = {}
_user_engine_lock = threading.RLock()

def _get_speculation_engine(user_id: str) -> SpeculationEngine:
    owner = str(user_id or '').strip()
    if not owner:
        raise AgentRuntimeError(ErrorCode.UNAUTHORIZED, 'authenticated user required', status=401)
    with _user_engine_lock:
        _runtime_cache_access[owner] = time.monotonic()
        return _speculation_engines.setdefault(owner, SpeculationEngine())

def _get_retry_engine(user_id: str) -> RetryEngine:
    owner = str(user_id or '').strip()
    if not owner:
        raise AgentRuntimeError(ErrorCode.UNAUTHORIZED, 'authenticated user required', status=401)
    with _user_engine_lock:
        _runtime_cache_access[owner] = time.monotonic()
        return _retry_engines.setdefault(owner, RetryEngine(max_attempts=5, backoff_base_ms=500))

def _prune_user_runtime_caches(max_idle_seconds: float = 3600.0) -> int:
    cutoff = float('inf') if float(max_idle_seconds) <= 0 else time.monotonic() - float(max_idle_seconds)
    removed = 0
    with _user_engine_lock, _exec_policies_lock:
        for owner, accessed in list(_runtime_cache_access.items()):
            if accessed >= cutoff:
                continue
            _runtime_cache_access.pop(owner, None)
            removed += int(_speculation_engines.pop(owner, None) is not None)
            removed += int(_retry_engines.pop(owner, None) is not None)
            removed += int(_exec_policies.pop(owner, None) is not None)
    return removed

tool_registry = get_global_registry()
event_log = EventLog()
agent_runtime = get_agent_runtime_store(ENGINE_DIR)
tool_pipeline = ToolExecutionPipeline(tool_registry, runtime_store=agent_runtime)
chat_projection_store = ChatProjectionStore(PROJECT_ROOT)
_goal_runner = None

def _dispatch_runtime_goal(goal):
    return _goal_runner.dispatch(goal) if _goal_runner is not None else False

register_runtime_endpoints(app, agent_runtime, tool_registry, goal_dispatcher=_dispatch_runtime_goal)

# ==================== 存储实例 (EngineStore 由 engine.store 导入) ====================
cron_store = EngineStore(ENGINE_DIR / "cron.json")
agent_store = EngineStore(ENGINE_DIR / "agents.json")
heartbeat_store = EngineStore(ENGINE_DIR / "heartbeat.json")

# SUBTITLE_FONTS / DEFAULT_FONT → engine.video_edit (imported above)

# 视频工具函数全部迁移到 engine.video_edit (通过 import 直接引用)









# _apply_* 视频函数全部迁移到 engine.video_edit (通过 import 直接引用)

# ═══════════════════════════════════════════════════════
# 新增视频处理功能 (2026-05-28)
# ═══════════════════════════════════════════════════════









# -- 存储工厂函数 (包装 engine.store, 自动注入 ENGINE_DIR) --
def get_ns(suffix: str, user_id: str = "") -> EngineStore:
    """获取用户隔离的 store 实例"""
    return _store_get_ns(ENGINE_DIR, suffix, user_id)

def get_chat_store(user_id: str = "") -> ChatStore:
    """获取用户隔离的 ChatStore 单例"""
    return _store_get_chat_store(ENGINE_DIR, user_id)

# ==================== 心跳 ====================
@app.get("/engine/health")
def engine_health():
    return {"status": "ok", "time": datetime.now().isoformat()}

@app.get("/engine/heartbeat")
def heartbeat(user_id: str = Query("")):
    """客户端心跳上报"""
    store = get_ns("heartbeat", user_id)
    client = "web"
    data = store.get()
    data[client] = {
        "last_seen": time.time(),
        "time": datetime.now().isoformat()
    }
    store.set(data)
    # 返回待处理的消息
    pending = data.get("pending_messages", [])
    result = {"ok": True, "pending": pending}
    if pending:
        data["pending_messages"] = []
        store.set(data)
    return result

@app.get("/engine/heartbeat/push")
def heartbeat_push(msg: str = Query(...), user_id: str = Query("")):
    """向客户端推送消息(通过心跳带回 + SSE实时推送)"""
    store = get_ns("heartbeat", user_id)
    data = store.get()
    pending = data.get("pending_messages", [])
    entry = {"msg": msg, "time": datetime.now().isoformat()}
    pending.append(entry)
    data["pending_messages"] = pending
    store.set(data)
    # Also push via SSE for instant delivery
    _broadcast_to_user(user_id, 'heartbeat:push', entry)
    return {"ok": True}

# ==================== 子代理并发锁 ====================
# per-user 的写锁,防止并行子代理写入冲突
_agent_store_locks: dict = {}
_agent_store_lock_lock = threading.Lock()
_agent_cancel_events: dict[tuple[str, str], threading.Event] = {}
_agent_cancel_events_lock = threading.Lock()

def _get_agent_cancel_event(user_id: str, name: str, *, reset: bool = False) -> threading.Event:
    key = (str(user_id), str(name))
    with _agent_cancel_events_lock:
        if reset or key not in _agent_cancel_events:
            _agent_cancel_events[key] = threading.Event()
        return _agent_cancel_events[key]

def _forget_agent_cancel_event(user_id: str, name: str) -> None:
    with _agent_cancel_events_lock:
        _agent_cancel_events.pop((str(user_id), str(name)), None)

def _get_agent_store_lock(user_id: str) -> threading.Lock:
    """获取用户级别的写锁(线程安全)"""
    with _agent_store_lock_lock:
        if user_id not in _agent_store_locks:
            _agent_store_locks[user_id] = threading.Lock()
        return _agent_store_locks[user_id]

# ==================== Cron 任务 (后台逻辑→engine.cron) ====================

@app.get("/engine/cron/list")
def cron_list(user_id: str = Query("")):
    store = get_ns("cron", user_id)
    return store.get()

@app.get("/engine/cron/create")
def cron_create(
    name: str = Query(...),
    interval: int = Query(...),
    action: str = Query(...),
    user_id: str = Query("")
):
    # ★ 防护: 拒绝空值或前端透传的 undefined 字符串
    if not action or not action.strip() or action.strip().lower() == 'undefined':
        return {"ok": False, "error": "action 参数无效(空值或 undefined)"}
    if interval < 10:
        return {"ok": False, "error": "interval 不能小于10秒"}
    store = get_ns("cron", user_id)
    jobs = store.get()
    jobs[name] = {
        "name": name, "interval": interval, "action": action.strip(),
        "enabled": True, "created": datetime.now().isoformat()
    }
    store.set(jobs)
    _cron_start(name, user_id, get_ns)
    return {"ok": True, "job": name}

@app.get("/engine/cron/delete")
def cron_delete(name: str = Query(...), user_id: str = Query("")):
    _cron_stop(name, user_id, get_ns)
    store = get_ns("cron", user_id)
    store.delete(name)
    return {"ok": True}



# ==================== 子代理 ====================


@app.get("/engine/agent/list")
def agent_list(user_id: str = Query("")):
    store = get_ns("agents", user_id)
    agents = store.get()
    cleaned = _cleanup_old_agents(agents)
    if cleaned:
        store.set(agents)
    return agents

@app.get("/engine/agent/create")
def agent_create(
    name: str = Query(...),
    prompt: str = Query(...),
    role: str = Query("general"),
    model: str = Query(""),
    base_url: str = Query(""),
    provider: str = Query(""),
    user_id: str = Query(""),
    proxy_url: str = Query(""),
    proxy_enabled: str = Query("")
):
    name = str(name or "").strip()
    prompt = str(prompt or "").strip()
    if not name or not prompt:
        raise HTTPException(400, "agent name and prompt are required")
    store = get_ns("agents", user_id)
    agents = store.get()
    # 自动清理过时子代理
    _cleanup_old_agents(agents)
    # 验证角色名
    if role not in AGENT_ROLES:
        role = "general"
    # ★ 注入当前日期到 prompt,让子代理知道真实时间,避免搜出过时信息
    now_cn = datetime.now().strftime("%Y年%m月%d日")
    tz_str = "Asia/Shanghai (UTC+8)"
    time_tag = f"\n\n[系统] 当前日期: {now_cn} (时区: {tz_str}), 所有搜索关键词应包含最新年份日期。"
    agent_data = {
        "name": name,
        "prompt": prompt + time_tag,
        "role": role,
        "status": "idle",
        "created": datetime.now().isoformat(),
        "model": model or "",
        "base_url": base_url or "",
        "provider": provider or "",
        "proxy_url": proxy_url or "",
        "proxy_enabled": proxy_enabled or ""
    }
    agents[name] = agent_data
    store.set(agents)
    return {"ok": True, "agent": name, "role": role}

@app.post("/engine/agent/create")
async def agent_create_post(request: Request, user_id: str = Query("")):
    body = await request.json()
    return agent_create(
        name=str(body.get("name") or ""), prompt=str(body.get("prompt") or ""),
        role=str(body.get("role") or "general"), model=str(body.get("model") or ""),
        base_url=str(body.get("base_url") or ""), provider=str(body.get("provider") or ""),
        user_id=user_id,
        proxy_url=str(body.get("proxy_url") or ""),
        proxy_enabled="1" if body.get("proxy_enabled") in (True, 1, "1") else "",
    )

@app.get("/engine/agent/run")
def agent_run(name: str = Query(...), user_id: str = Query(""), message: str = Query(""), from_ask: str = Query("")):
    """运行子代理(调用AI完成指定任务)
    message - 追加的消息内容(用于agent_ask)
    from_ask - 如果是agent_ask触发,消息追加到prompt后"""
    store = get_ns("agents", user_id)
    agents = store.get()
    agent = agents.get(name)
    if not agent:
        raise HTTPException(404, f"Agent {name} not found")
    # Running agents cannot be started twice; completed agents remain continuable.
    _current_status = agent.get("status", "")
    _is_followup = bool(from_ask and message)
    if _current_status == "running":
        return {"ok": False, "error": "Agent 已在运行，请勿重复启动"}
    if _current_status == "completed" and not _is_followup:
        return {"ok": False, "error": "Agent 已完成；可通过 agent_ask 继续同一会话"}
    _base_prompt = str(agent.get("prompt", ""))
    _previous_result = str(agent.get("result", ""))
    _run_prompt = str(message) if _is_followup else _base_prompt

    # ★ 所有 agent 统一从主聊天配置同步
    main_config = _get_main_chat_config(user_id)
    api_key = main_config.get("api_key", "") or os.getenv("OPENAI_API_KEY", "")

    # 子代理路由与模型对齐
    _agent_has_override = bool(agent.get("independent_route") or agent.get("use_custom_route"))
    agent_model = str(agent.get("model") or "").strip()
    agent_base = str(agent.get("base_url") or "").strip()

    if _agent_has_override or (agent_model and agent_base):
        base_url = agent_base or main_config.get("base_url", "")
        model = agent_model or main_config.get("model", "")
    else:
        base_url = main_config.get("base_url", "") or agent_base or os.getenv("OPENAI_BASE_URL", "") or "https://api.deepseek.com/v1"
        model = main_config.get("model", "") or agent_model or "deepseek-chat"

    # ★ 终极防串扰保护: 保证 base_url 与 model 严格兼容，杜绝历史跨提供商残留
    if "api.longcat.chat" in base_url and "longcat" not in model.lower():
        # LongCat endpoint 不接受异构模型，重置为 main_config 真实路由或 LongCat-2.0
        if main_config.get("base_url") and "api.longcat.chat" not in main_config.get("base_url"):
            base_url = main_config.get("base_url")
            model = main_config.get("model") or model
        else:
            model = "LongCat-2.0"
    if "api.minimaxi.com" in base_url and "minimax" not in model.lower():
        if main_config.get("base_url") and "api.minimaxi.com" not in main_config.get("base_url"):
            base_url = main_config.get("base_url")
            model = main_config.get("model") or model
        else:
            model = "MiniMax-M2.7"
    if "api.deepseek.com" in base_url and "deepseek" not in model.lower():
        if main_config.get("base_url") and "api.deepseek.com" not in main_config.get("base_url"):
            base_url = main_config.get("base_url")
            model = main_config.get("model") or model
        else:
            model = "deepseek-chat"

    # ★ 防御: model 为无效值(如"加载中...")时回退到默认
    if not model or model.startswith("加载中") or len(model) < 3:
        model = "deepseek-chat" if "deepseek" in base_url else "MiniMax-M2.7"

    if not api_key:
        return {"error": "未配置API Key,请在聊天设置中配置后重试"}

    _provider_info = detect_provider(base_url, model, anthropic_format=bool(main_config.get("anthropic_format")))
    _cancel_event = _get_agent_cancel_event(user_id, name, reset=True)
    _parent = agent_runtime.ensure_session(
        user_id, f"legacy-agent-parent:{name}", origin="legacy-subagent",
        provider=_provider_info.name, model=model,
    )
    _runtime_id = str(agent.get("runtime_agent_id") or "")
    _runtime_subagent = None
    if _runtime_id:
        try:
            agent_runtime.get_subagent(_runtime_id, user_id=user_id)
            _runtime_subagent = agent_runtime.update_subagent(
                _runtime_id, "queued", user_id=user_id, provider=_provider_info.name,
                model=model, prompt=_run_prompt,
            )
        except AgentRuntimeError:
            _runtime_subagent = None
    if _runtime_subagent is None:
        _runtime_subagent = agent_runtime.create_subagent(
            user_id, _parent["session_id"], name, _run_prompt, role=agent.get("role", "general"),
            provider=_provider_info.name, model=model,
        )
    agent["runtime_agent_id"] = _runtime_subagent["agent_id"]
    agent["runtime_session_id"] = _runtime_subagent["child_session_id"]
    agent["provider"] = _provider_info.name
    agent["model"] = model
    agents[name] = agent
    store.set(agents)

    # ★ 根据角色选择工具集(最小权限原则)
    agent_role = agent.get("role", "general")

    # ★ 应用子代理的代理配置（用 http_client 而非全局 env 避免并发竞争）
    _agent_proxy_url = agent.get("proxy_url", "")
    _agent_proxy_enabled = agent.get("proxy_enabled", "")
    _agent_http_client = None
    if _agent_proxy_enabled == '1' and _agent_proxy_url:
        import httpx as _httpx
        _agent_http_client = _httpx.Client(proxy=_agent_proxy_url, trust_env=False)
    elif _requires_auto_proxy(base_url):
        import httpx as _httpx
        _agent_http_client = _httpx.Client(proxy=_PROXY_URL, trust_env=False)
        print(f'[Agent {name}] 境外 API 定向代理已启用')

    role_config = AGENT_ROLES.get(agent_role, AGENT_ROLES["general"])
    TOOLS = _filter_tools_by_role(agent_role)

    # ★ 角色级别模型选择:cheap 角色用轻量模型节省开销
    if role_config["model_tier"] == "cheap":
        # 尝试用 deepseek-chat 或站内最便宜的模型
        cheap_model = main_config.get("cheap_model", "") or os.getenv("CHEAP_MODEL", "")
        if cheap_model:
            model = cheap_model
        elif "minimaxi" in model.lower():
            model = "MiniMax-M2.7"  # MiniMax 本身已经是便宜模型
        # 对于 explorer/planner 减少 max_tokens 节省token
    max_agent_rounds = role_config["max_rounds"]

    def _exec_generate_image_local(uid, tool_name, args):
        """生图工具本地多提供商分发 (读 DB 配置, 支持 MiniMax/OpenAI/xAI/Custom)"""
        import sqlite3, urllib.request, urllib.error, base64 as _b64

        # 从 DB 读取用户图像配置
        db_path = os.path.join(PROJECT_ROOT, "users", "oneapichat.db")
        cfg = {}
        try:
            conn = sqlite3.connect(db_path)
            cur = conn.cursor()
            cur.execute("SELECT config_json FROM user_config WHERE user_id = ?", (uid,))
            row = cur.fetchone()
            if row:
                cfg = json.loads(row[0]) or {}
            conn.close()
        except Exception as e:
            return json.dumps({"error": f"读取配置失败: {e}"})

        provider = (cfg.get("imageProvider") or "minimax").lower()
        prompt = args.get("prompt", "").strip()
        if not prompt:
            return json.dumps({"error": "prompt is required"})

        # ── MiniMax ──
        if provider == "minimax":
            mmx_cfg_path = os.path.join(PROJECT_ROOT, "config", ".mmx_config.json")
            mmx_key = ""
            try:
                with open(mmx_cfg_path) as f:
                    mmx_key = json.load(f).get("api_key", "")
            except Exception:
                pass
            if not mmx_key:
                return json.dumps({"error": "MiniMax API key not configured"})
            api_url = "https://api.minimaxi.com/v1/image_generation"
            body = json.dumps({"model": "image-01", "prompt": prompt, "n": 1, "response_format": "url"}).encode()
            req = urllib.request.Request(api_url, data=body, headers={"Content-Type": "application/json", "Authorization": f"Bearer {mmx_key}"})
            try:
                with urllib.request.urlopen(req, timeout=120) as resp:
                    data = json.loads(resp.read())
                imgs = []
                for d in (data.get("data") or []):
                    if isinstance(d, dict) and d.get("image_url"):
                        imgs.append(d["image_url"])
                if imgs:
                    return json.dumps({"images": imgs, "status": "ok", "provider": "minimax"})
                return json.dumps({"error": "MiniMax 未返回图片", "raw": str(data)[:300]})
            except Exception as e:
                return json.dumps({"error": f"MiniMax 请求失败: {e}"})

        # ── OpenAI 兼容 (openrouter/openai/custom) ──
        key_map = {
            "openrouter": {"key": "imageApiKeyOpenrouter", "url": "imageBaseUrlOpenrouter", "model": "imageModel_openrouter", "default_url": "https://openrouter.ai/api/v1", "default_model": "openai/gpt-5.4-image-2"},
            "openai":     {"key": "imageApiKeyOpenai",     "url": "imageBaseUrlOpenai",     "model": "imageModel_openai",     "default_url": "https://api.openai.com/v1",        "default_model": "gpt-image-1"},
            "custom":     {"key": "imageApiKeyCustom",     "url": "imageBaseUrlCustom",     "model": "imageModel_custom",     "default_url": "",                              "default_model": ""},
        }
        if provider not in key_map:
            return json.dumps({"error": f"Unknown provider: {provider}"})
        km = key_map[provider]

        # 解密 v2 密钥
        enc_key = cfg.get(km["key"]) or ""
        api_key = ""
        if enc_key.startswith("v2:"):
            try:
                api_key = _decrypt_xor(enc_key) or ""
            except Exception:
                api_key = enc_key
        else:
            api_key = enc_key

        base_url = (cfg.get(km["url"]) or km["default_url"]).rstrip("/")
        model = (cfg.get(km["model"]) or km["default_model"]).strip()

        if not api_key:
            return json.dumps({"error": f"未配置 {provider} 的 API Key"})
        if not base_url:
            return json.dumps({"error": f"未配置 {provider} 的 API Base URL"})
        if not base_url.endswith("/v1"):
            base_url += "/v1"

        # xAI (Grok) 不支持 size 参数
        is_xai = "x.ai" in base_url.lower() or "grok" in model.lower()
        n = min(int(args.get("n", 1)), 10)
        body = {"model": model, "prompt": prompt, "n": n, "response_format": "b64_json"}
        if not is_xai:
            size_map = {"1:1": "1024x1024", "16:9": "1792x1024", "9:16": "1024x1792"}
            body["size"] = size_map.get(args.get("aspect_ratio", "1:1"), "1024x1024")

        api_url = base_url + "/images/generations"
        req = urllib.request.Request(api_url, data=json.dumps(body).encode(), headers={"Content-Type": "application/json", "Authorization": f"Bearer {api_key}"})
        try:
            with urllib.request.urlopen(req, timeout=120) as resp:
                data = json.loads(resp.read())
            imgs = []
            for item in (data.get("data") or []):
                if isinstance(item, dict):
                    if item.get("url"):
                        imgs.append(item["url"])
                    elif item.get("b64_json"):
                        imgs.append("data:image/png;base64," + item["b64_json"])
            if imgs:
                return json.dumps({"images": imgs, "status": "ok", "provider": provider, "model": model})
            return json.dumps({"error": "未返回图片", "raw": str(data)[:300]})
        except urllib.error.HTTPError as e:
            err_body = e.read().decode() if e.fp else ""
            return json.dumps({"error": f"生图失败 ({e.code}): {err_body[:300]}", "provider": provider})
        except Exception as e:
            return json.dumps({"error": f"生图请求异常: {e}", "provider": provider})

    # Register missing role-filtered tools for schema validation and telemetry.
    for _legacy_spec in TOOLS:
        _legacy_fn = _legacy_spec.get("function", {})
        _legacy_name = str(_legacy_fn.get("name") or "")
        if _legacy_name and tool_registry.get(_legacy_name) is None:
            tool_registry.register_from_dict({
                "name": _legacy_name,
                "description": _legacy_fn.get("description", ""),
                "parameters": _legacy_fn.get("parameters", {"type": "object", "properties": {}}),
                "approval": "auto",
                "timeout_ms": 120000,
            })

    def _execute_tool(tool_name, args):
        """执行子代理工具调用。文件操作必须回流至 server_tools，禁止绕过统一权限/观察/验证策略。"""
        def _legacy_file_request(action, payload):
            try:
                from engine.runtime_auth import get_internal_bridge_secret
                headers = {"X-OneAPIChat-Internal": get_internal_bridge_secret(PROJECT_ROOT)}
                payload = payload if isinstance(payload, dict) else {}
                path = payload.get("path") or payload.get("file_path") or ""
                common = {"user_id": user_id, "path": path}
                if payload.get("cwd"): common["cwd"] = payload.get("cwd")
                if action == "read":
                    common["max_lines"] = int(payload.get("max_lines") or payload.get("limit") or 200)
                    response = _http_session_no_proxy.get("http://127.0.0.1:8766/engine/file/read", params=common, headers=headers, timeout=30)
                elif action in {"write", "append"}:
                    common["append"] = "true" if action == "append" else "false"
                    response = _http_session_no_proxy.post("http://127.0.0.1:8766/engine/file/write", params=common, data=str(payload.get("content", "")).encode("utf-8"), headers={**headers, "Content-Type": "text/plain"}, timeout=45)
                else:
                    return "错误:未知文件工具操作"
                parsed = response.json()
                return json.dumps(parsed, ensure_ascii=False)
            except Exception as exc:
                return f"文件工具桥接失败: {type(exc).__name__}"

        if tool_name == "web_search":
            query = args.get("query", "")
            if not query:
                return "错误:缺少 query 参数"
            try:
                # 从主聊配置读取搜索 Provider 和对应的 API Key
                search_provider = ""  # 从配置读取；无配置时使用 Tavily/MiniMax 回退链
                search_api_key = ""
                try:
                    config_path = os.path.join(PROJECT_ROOT, f"chat_data/config_user_{user_id}.json")
                    with open(config_path) as f:
                        raw_cfg = json.load(f)
                    # 读取搜索 Provider (用户可能在配置中选择 brave/tavily/deepseek/minimax)
                    raw_provider = raw_cfg.get("searchProvider", "") or ""
                    if raw_provider and raw_provider != "not-needed":
                        search_provider = raw_provider
                except Exception:
                    pass
                # ★ 未配置或读取失败时，优先使用 Tavily 配置
                if not search_provider:
                    search_provider = "tavily"
                # 读取对应 Provider 的 API Key(优先专用字段,回退到通用 searchApiKey)
                try:
                    provider_key_fields = {
                        "tavily": ["searchApiKeyTavily", "searchApiKey"],
                        "brave": ["searchApiKeyBrave", "searchApiKey"],
                        "google": ["searchApiKeyGoogle", "searchApiKey"],
                    }
                    key_fields = provider_key_fields.get(search_provider, ["searchApiKey"])
                    # ★ 依次尝试所有候选字段,第一个非空且解密成功的胜出
                    for key_field in key_fields:
                        stored = raw_cfg.get(key_field, "")
                        if not stored:
                            continue
                        decrypted = _decrypt_xor(stored)
                        if decrypted:
                            search_api_key = decrypted
                            break
                        elif not stored.startswith("v2:"):
                            search_api_key = stored  # 明文存储
                            break
                except Exception:
                    search_api_key = ""
                # ★ 无 Key 时改走已有的 MiniMax 搜索回退，不再调用不可达的 DuckDuckGo
                if not search_api_key and search_provider in {"tavily", "brave", "google"}:
                    print(f"[web_search] provider={search_provider} 无有效 API Key,回退 MiniMax", flush=True)
                    search_provider = "minimax"
                print(f"[web_search] provider={search_provider} credential_configured={bool(search_api_key)}", flush=True)

                # 兼容旧配置：已下线的搜索引擎转入现有搜索回退链。
                if search_provider in {"duckduckgo", "google"}:
                    search_provider = "minimax"
                # Tavily 搜索 (失败时自动回退到 PHP 代理 → MiniMax CLI)
                def _try_tavily(q):
                    import re as _re  # ★ 嵌套函数必须内部 import,否则闭包找不到模块级变量
                    # ★ 路径1: 直接调用 Tavily API
                    if search_api_key:
                        try:
                            _depth = "basic"
                            if search_api_key.startswith("tvly-") and not search_api_key.startswith("tvly-dev-") and not search_api_key.startswith("tvly-free-"):
                                _depth = "advanced"
                            r = _http_session.post(
                                "https://api.tavily.com/search",
                                json={"api_key": search_api_key, "query": q, "search_depth": _depth, "max_results": 10, "include_answer": True},
                                timeout=20
                            )
                            if r.status_code == 200:
                                data = r.json()
                                results = data.get("results", [])
                                answer = data.get("answer", "") or ""
                                if results or answer:
                                    if not results and answer:
                                        return f"搜索结果 (query: {q}):\n[摘要] {answer[:500]}"
                                    lines = []
                                    for res in results[:8]:
                                        title = res.get("title", "")
                                        url = res.get("url", "")
                                        content = res.get("content", "")[:200].replace("\n", " ")
                                        content = _re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', content)
                                        lines.append(f"- [{title}]({url})\n  {content}")
                                    return f"搜索结果 (provider: tavily, query: {q}):\n" + "\n\n".join(lines) + "\n\n注: 如需查看详情请使用 web_fetch 工具抓取网页内容。"
                        except Exception:
                            pass
                        except Exception as _e:
                            with open("/tmp/engine_tool_debug.log", "a") as _f:
                                _f.write(f"[TAVILY_XXX] exception: {_e}\n")
                    # ★ 路径2: PHP 代理回退(与父代理相同路径)
                    try:
                        _api_key_param = f"&api_key={search_api_key}" if search_api_key else ""
                        _proxy_url = f"https://127.0.0.1/oneapichat/api/engine_api.php?action=tavily_search&q={requests.utils.quote(q)}&limit=10{_api_key_param}"
                        _pr = _http_session.get(_proxy_url, headers={"Host": "naujtrats.xyz"}, timeout=20, verify=False)
                        if _pr.status_code == 200:
                            _pd = _pr.json()
                            if _pd.get("results") or _pd.get("answer"):
                                results = _pd.get("results", [])
                                answer = _pd.get("answer", "")
                                if answer and not results:
                                    return f"搜索结果 (query: {q}):\n[摘要] {answer[:500]}"
                                lines = []
                                for res in results[:8]:
                                    if isinstance(res, dict):
                                        title = res.get("title", "")
                                        url = res.get("url", "")
                                        content = res.get("content", "")[:200].replace("\n", " ")
                                        content = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', content)
                                        lines.append(f"- [{title}]({url})\n  {content}")
                                if lines:
                                    return f"搜索结果 (provider: tavily_proxy, query: {q}):\n" + "\n\n".join(lines) + "\n\n注: 如需查看详情请使用 web_fetch 工具抓取网页内容。"
                            if _pr.text and len(_pr.text) > 50 and '<!doctype' not in _pr.text.lower():
                                return f"搜索结果 (provider: proxy, query: {q}):\n{_pr.text[:2000]}"
                    except Exception:
                        pass
                    return None

                # MiniMax CLI 搜索回退
                def _try_minimax_search(q):
                    try:
                        import subprocess as _subprocess
                        # ★ 自动给搜索词追加当前日期,提升搜索相关性
                        today_str = datetime.now().strftime("%Y年%m月%d日")
                        if today_str[:4] not in q:
                            q = q + f" {today_str}"
                        r = _subprocess.run(
                            ["mmx", "search", "query", "--q", q, "--output", "json"],
                            capture_output=True, text=True, timeout=30
                        )
                        if r.returncode != 0:
                            return None
                        data = json.loads(r.stdout)
                        organic = data.get("organic", [])
                        if not organic:
                            return None
                        lines = []
                        for res in organic[:8]:
                            title = res.get("title", "")
                            url = res.get("link", "")
                            content = res.get("snippet", res.get("body", ""))[:200].replace("\n", " ")
                            content = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', content)
                            lines.append(f"- [{title}]({url})\n  {content}")
                        return f"搜索结果 (provider: minimax, query: {q}):\n" + "\n\n".join(lines) + "\n\n注: 如需查看详情请使用 web_fetch 工具抓取网页内容。"
                    except Exception as _mmx_e2:
                        return None

                if search_provider == "tavily":
                    result = _try_tavily(query)
                    if result:
                        return result
                    # Tavily 失败 → 回退 MiniMax CLI
                    mmx_result = _try_minimax_search(query)
                    if mmx_result:
                        return mmx_result
                    return f'搜索 "{query}" 无结果。搜索引擎不可用,请用 web_fetch 或换其他方式获取信息。'

                # Brave 搜索 (官方接口 + 代理)
                if search_provider == "brave":
                    if not search_api_key:
                        return f"搜索出错: 未找到 Brave API Key (请先在设置中配置搜索API Key)"
                    headers = {
                        "Accept": "application/json",
                        "Accept-Encoding": "gzip",
                        "X-Subscription-Token": search_api_key
                    }
                    try:
                        r = _http_session.get(
                            f"https://api.search.brave.com/res/v1/web/search?q={requests.utils.quote(query)}&count=8&safesearch=off&text_decorations=0",
                            headers=headers, timeout=15
                        )
                        if r.status_code == 200:
                            data = r.json()
                            results = data.get("web", {}).get("results", [])
                            if results:
                                lines = []
                                for res in results[:8]:
                                    title = res.get("title", "")
                                    url = res.get("url", "")
                                    desc = res.get("description", "")
                                    extra = res.get("extra_snippets", [])
                                    if extra and isinstance(extra, list):
                                        desc = desc + " " + " ".join(extra)
                                    content = desc[:300].replace("\n", " ")
                                    content = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', content)
                                    lines.append(f"- [{title}]({url})\n  {content}")
                                return f"搜索结果 (provider: {search_provider}, query: {query}):\n" + "\n\n".join(lines) + "\n\n注: 如需查看详情请使用 web_fetch 工具抓取网页内容。"
                        else:
                            print(f"[web_search] Brave API 返回状态码 {r.status_code}: {r.text[:200]}", flush=True)
                    except Exception as _brave_err:
                        print(f"[web_search] Brave 请求失败: {str(_brave_err)}", flush=True)

                # 兜底: 不支持的 provider (如 minimax) 统一先用 MiniMax CLI 搜索
                print(f"[web_search] 走兜底, search_provider={search_provider}, query_size={len(query)}", flush=True)
                try:
                    mmx_result = _try_minimax_search(query)
                    print(f"[web_search] MiniMax fallback returned={bool(mmx_result)} size={len(str(mmx_result or ''))}", flush=True)
                    if mmx_result:
                        return mmx_result
                except Exception as _mmx_e:
                    print(f"[web_search] MiniMax fallback error={type(_mmx_e).__name__}", flush=True)
                # MiniMax 也无结果时,再试试 Tavily
                print(f"[web_search] 尝试 Tavily 兜底", flush=True)
                tavily_result = _try_tavily(query)
                print(f"[web_search] Tavily fallback returned={bool(tavily_result)} size={len(str(tavily_result or ''))}", flush=True)
                if tavily_result:
                    return tavily_result
                print(f"[web_search] 全部搜索失败,返回无结果", flush=True)
                return f'搜索 "{query}" 无结果。搜索引擎不可用,请用 web_fetch 或换其他方式获取信息。'
            except Exception as e:
                return f"搜索出错: {str(e)}\n请稍后重试或更换关键词。"
        elif tool_name == "web_fetch":
            urls = []
            if args.get("url"): urls.append(args["url"])
            if args.get("urls"): urls.extend(args["urls"][:3])
            results = []
            for url in urls:
                try:
                    # 使用增强提取器(支持平台特定 + 通用HTML解析)
                    from engine.web_extract import web_extractor as _wex
                    import asyncio as _asyncio
                    bm = None
                    try:
                        from engine.browser import ensure_browser_connected
                        bm = _asyncio.run(ensure_browser_connected())
                    except Exception:
                        pass
                    ext_result = _asyncio.run(asyncio.wait_for(_wex.extract(url, _http_session, bm), timeout=15))
                    if ext_result and ext_result.content and len(ext_result.content) > 100:
                        text = ext_result.content
                    else:
                        # 回退到简单HTTP请求
                        r = _http_session.get(url, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"}, timeout=10)
                        raw = r.text
                        raw = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', raw)
                        raw = re.sub(r'<script[^>]*>.*?</script>', '', raw, flags=re.DOTALL)
                        raw = re.sub(r'<style[^>]*>.*?</style>', '', raw, flags=re.DOTALL)
                        raw = re.sub(r'<[^>]+>', ' ', raw)
                        raw = re.sub(r'\s+', ' ', raw).strip()
                        text = raw[:20000]
                    # ★ 内容质量检查: 如果提取的内容大部分是导航/菜单,标注为低质量
                    _nav_keywords = ['首页', '登录', '注册', '导航', '菜单', 'Copyright', '备案号', '网站地图']
                    _nav_count = sum(1 for k in _nav_keywords if k in text[:2000])
                    if _nav_count >= 3 and len(text) < 500:
                        results.append(f"[{url}]: ⚠️ 提取内容质量较低(可能为导航页面)。建议直接访问网站查看。\n原始片段: {text[:500]}")
                    else:
                        results.append(f"[{url}]: {text}")
                except Exception as e:
                    results.append(f"[{url}]: ⚠️ 抓取失败 - {str(e)[:100]}。建议更换URL或直接搜索。")
            return "\n\n".join(results) if results else "未提供URL"
        elif tool_name == "run_skill":
            skill_name = args.get("skill_name", "")
            skill_params = args.get("params", {})
            if not skill_name:
                return "错误: 缺少 skill_name 参数"
            try:
                from engine.skills import get_skill, render_prompt
                skill = get_skill(user_id, skill_name)
                if not skill:
                    return f"技能 '{skill_name}' 不存在。可用技能请查看系统提示词。"
                prompt = render_prompt(skill.get("prompt_template", ""), skill_params)
                if not prompt:
                    return "技能提示词模板为空"
                # 将渲染后的提示词作为当前对话的上下文返回
                return (
                    f"[技能: {skill.get('label', skill_name)}]\n"
                    f"已为以下任务生成定制提示词。请按照以下指示完成任务:\n\n"
                    f"---\n{prompt}\n---\n\n"
                    f"可用工具: {', '.join(skill.get('tools', []))}"
                )
            except Exception as e:
                return f"run_skill 错误: {str(e)}"
        elif tool_name == "platform_extract":
            url = args.get("url", "")
            if not url:
                return "错误: 缺少 url 参数"
            try:
                from engine.web_extract import web_extractor as _wex
                import asyncio as _asyncio
                bm = None
                try:
                    from engine.browser import ensure_browser_connected
                    bm = _asyncio.run(ensure_browser_connected())
                except Exception:
                    pass
                result = _asyncio.run(_wex.extract(url, _http_session, bm))
                if result and result.content:
                    output = result.content
                    if result.structured:
                        output += "\n\n[结构化数据]\n" + json.dumps(result.structured, ensure_ascii=False, indent=2)[:5000]
                    return output
                return f"无法从 {url} 提取内容"
            except Exception as e:
                return f"platform_extract 错误: {str(e)}"
        elif tool_name == "engine_push":
            import re, hashlib, shutil as _shutil
            msg = args.get("msg", "")
            file_path = args.get("file", "")

            # Step 1: clean Markdown pollution from URLs in msg
            def _clean_url(u):
                # strip leading/trailing markdown bold/italic/strikethrough markers
                u = re.sub(r'^[*_~`]+', '', u)
                u = re.sub(r'[*_~`]+$', '', u)
                u = re.sub(r'^[)\]]+', '', u)
                # extract URL from markdown link syntax [text](url)
                m = re.search(r'\]\(\s*(https?://\S+?)\s*\)', u)
                if m:
                    u = m.group(1)
                return u

            if msg:
                # ★ strip ** wrapped around URL: **https://...** → https://...
                msg = re.sub(r'\*{1,2}(https?://\S+?)\*{1,2}', r'\1', msg)
                # ★ strip trailing ** after URL: https://...** text → https://... text
                msg = re.sub(r'(https?://\S+?)\*{1,2}(\s|$)', r'\1\2', msg)
                # generic URL-level cleanup
                msg = re.sub(r'https?://\S+', lambda m: _clean_url(m.group(0)), msg)

            # Step 2: handle file parameter — copy to shared dir, generate download URL
            if file_path and os.path.isfile(file_path):
                import pwd, grp
                shared_dir = os.path.join(PROJECT_ROOT, 'uploads', 'shared')
                os.makedirs(shared_dir, exist_ok=True)
                fname = os.path.basename(file_path)
                # ★ include mtime in hash so re-pushes of same-named file get unique URLs
                h = hashlib.md5((fname + str(os.path.getmtime(file_path))).encode()).hexdigest()[:12]
                ext = os.path.splitext(fname)[1] or '.bin'
                shared_name = f"push_{h}{ext}"
                shared_path = os.path.join(shared_dir, shared_name)
                _shutil.copy2(file_path, shared_path)
                # ★ chown to www-data so nginx can serve the file
                try:
                    os.chown(shared_path, pwd.getpwnam('www-data').pw_uid, grp.getgrnam('www-data').gr_gid)
                except Exception:
                    pass  # non-fatal if running as non-root
                dl_url = f"https://naujtrats.xyz/oneapichat/uploads/shared/{shared_name}"
                if dl_url not in msg:
                    msg = (msg + "\n\n" + dl_url).strip()

            if msg:
                push_store = get_ns("heartbeat", user_id)
                data = push_store.get()
                pending = data.get("pending_messages", []) if data else []
                pending.append({"msg": msg, "time": datetime.now().isoformat()})
                if data is None:
                    data = {}
                data["pending_messages"] = pending
                push_store.set(data)
            return "消息已推送到用户"
        elif tool_name == "server_file_append":
            return _legacy_file_request("append", args)
        elif tool_name == "server_file_write":
            return _legacy_file_request("write", args)
        elif tool_name == "server_file_read":
            return _legacy_file_request("read", args)
            try:
                import asyncio
                from engine.browser import ensure_browser_connected
                bm = asyncio.run(ensure_browser_connected())
                result = asyncio.run(bm.navigate(args.get("url", "")))
                return json.dumps(result, ensure_ascii=False)
            except Exception as e:
                return f"浏览器导航失败: {str(e)}"
        elif tool_name == "browser_screenshot":
            try:
                import asyncio
                from engine.browser import ensure_browser_connected
                bm = asyncio.run(ensure_browser_connected())
                result = asyncio.run(bm.screenshot())
                return json.dumps(result, ensure_ascii=False)
            except Exception as e:
                return f"浏览器截图失败: {str(e)}"
        elif tool_name == "browser_click":
            try:
                import asyncio
                from engine.browser import ensure_browser_connected
                bm = asyncio.run(ensure_browser_connected())
                result = asyncio.run(bm.click(args.get("selector", "")))
                return json.dumps(result, ensure_ascii=False)
            except Exception as e:
                return f"浏览器点击失败: {str(e)}"
        elif tool_name == "browser_type":
            try:
                import asyncio
                from engine.browser import ensure_browser_connected
                bm = asyncio.run(ensure_browser_connected())
                result = asyncio.run(bm.type_text(args.get("selector", ""), args.get("text", "")))
                return json.dumps(result, ensure_ascii=False)
            except Exception as e:
                return f"浏览器输入失败: {str(e)}"
        elif tool_name == "browser_get_content":
            try:
                import asyncio
                from engine.browser import ensure_browser_connected
                bm = asyncio.run(ensure_browser_connected())
                result = asyncio.run(bm.get_content())
                return json.dumps(result, ensure_ascii=False)
            except Exception as e:
                return f"浏览器获取内容失败: {str(e)}"
        elif tool_name == "browser_get_snapshot":
            try:
                import asyncio
                from engine.browser import ensure_browser_connected
                bm = asyncio.run(ensure_browser_connected())
                result = asyncio.run(bm.get_snapshot())
                return json.dumps(result, ensure_ascii=False)
            except Exception as e:
                return f"浏览器获取结构失败: {str(e)}"
        elif tool_name == "video_edit":
            try:
                import json as _json
                action = args.get("action", "")
                params = args.get("params", {})
                input_path = args.get("input_path", "")
                output_path = args.get("output_path", "/tmp/video_output.mp4")
                if not input_path:
                    return "错误: 未提供输入视频路径"
                if not os.path.exists(input_path):
                    return f"错误: 输入文件不存在: {input_path}"
                if action == "info":
                    cmd = ["ffprobe", "-v", "quiet", "-print_format", "json", "-show_format", "-show_streams", input_path]
                    r = subprocess.run(cmd, capture_output=True, text=True, timeout=30)
                    return r.stdout
                elif action == "trim":
                    start = params.get("start", 0)
                    end = params.get("end", None)
                    from moviepy import VideoFileClip
                    clip = VideoFileClip(input_path)
                    if end:
                        clip = clip.subclipped(start, end)
                    else:
                        clip = clip.subclipped(start)
                    clip.write_videofile(output_path, codec="libx264", audio_codec="aac")
                    clip.close()
                    return f"裁剪完成: {output_path}"
                elif action == "speed":
                    factor = float(params.get("factor", 1.0))
                    from moviepy import VideoFileClip
                    from moviepy import vfx
                    clip = VideoFileClip(input_path)
                    clip = clip.with_effects([vfx.MultiplySpeed(factor)])
                    clip.write_videofile(output_path, codec="libx264", audio_codec="aac")
                    clip.close()
                    return f"调速完成 (x{factor}): {output_path}"
                elif action == "resize":
                    width = params.get("width", 0)
                    height = params.get("height", 0)
                    from moviepy import VideoFileClip
                    from moviepy import vfx
                    clip = VideoFileClip(input_path)
                    if width and height:
                        clip = clip.resized((width, height))
                    elif width:
                        clip = clip.resized(width=width)
                    elif height:
                        clip = clip.resized(height=height)
                    clip.write_videofile(output_path, codec="libx264", audio_codec="aac")
                    clip.close()
                    return f"缩放完成: {output_path}"
                elif action == "audio":
                    from moviepy import VideoFileClip
                    clip = VideoFileClip(input_path)
                    audio_output = output_path or input_path + ".mp3"
                    if clip.audio:
                        clip.audio.write_audiofile(audio_output)
                    clip.close()
                    return f"音频提取完成: {audio_output}"
                elif action == "concat":
                    files = params.get("files", [])
                    if not files:
                        return "错误: concat 需要 files 数组参数"
                    from moviepy import concatenate_videoclips
                    clips = [VideoFileClip(f) for f in files]
                    final = concatenate_videoclips(clips, method="compose")
                    final.write_videofile(output_path, codec="libx264", audio_codec="aac")
                    for c in clips: c.close()
                    final.close()
                    return f"拼接完成 ({len(files)}个视频): {output_path}"
                elif action == "overlay":
                    overlay_path = params.get("overlay_path", "")
                    if not overlay_path or not os.path.exists(overlay_path):
                        return "错误: overlay 需要 overlay_path 参数指向存在的文件"
                    x = params.get("x", 10)
                    y = params.get("y", 10)
                    scale = params.get("scale", 0.3)
                    from moviepy import CompositeVideoClip
                    clip = VideoFileClip(input_path)
                    ov = VideoFileClip(overlay_path).resized(scale)
                    ov = ov.with_position((x, y))
                    final = CompositeVideoClip([clip, ov])
                    final.write_videofile(output_path, codec="libx264", audio_codec="aac")
                    clip.close(); ov.close(); final.close()
                    return f"画中画完成: {output_path}"
                elif action == "text":
                    return _apply_subtitle(input_path, output_path, params)
                elif action == "rotate":
                    angle = float(params.get("angle", 90))
                    from moviepy import vfx
                    clip = VideoFileClip(input_path)
                    clip = clip.with_effects([vfx.Rotate(angle)])
                    clip.write_videofile(output_path, codec="libx264", audio_codec="aac")
                    clip.close()
                    return f"旋转完成 ({angle}°): {output_path}"
                elif action in ("filter", "video_filter"):
                    return _apply_ffmpeg_filter(input_path, output_path, params)
                elif action in ("transition", "video_transition"):
                    return _apply_ffmpeg_transition(input_path, output_path, params)
                elif action == "crop":
                    return _apply_crop(input_path, output_path, params)
                elif action == "reverse":
                    return _apply_reverse(input_path, output_path, params)
                elif action == "mute":
                    return _apply_mute(input_path, output_path, params)
                elif action == "bgm":
                    return _apply_bgm(input_path, output_path, params)
                elif action == "enhance":
                    return _apply_enhance(input_path, output_path, params)
                elif action == "gif":
                    return _apply_gif(input_path, output_path, params)
                elif action == "silent_cut":
                    return _apply_silent_cut(input_path, output_path, params)
                elif action == "stt":
                    return _apply_stt(input_path, output_path, params)
                elif action == "stt_to_timeline":
                    return _apply_stt_to_timeline(input_path, output_path, params)
                elif action == "style":
                    return _apply_subtitle_style(input_path, output_path, params)
                elif action == "tts":
                    return _apply_tts(params)
                elif action == "voice":
                    # voice: 将 TTS 生成的音频混入视频
                    audio_path = params.get("audio_path", "")
                    if not audio_path:
                        # 如果没有 audio_path,先调 TTS 生成
                        tts_result = _apply_tts(params)
                        if "失败" in tts_result or "异常" in tts_result:
                            return tts_result
                        audio_path = tts_result.split(": ")[1].split(" ")[0] if ": " in tts_result else "/tmp/tts_output.mp3"
                    return _apply_voice_to_video(input_path, audio_path, output_path, params)
                elif action == "compose":
                    return _apply_compose(input_path, output_path, params)
                else:
                    return f"未知操作: {action}, 支持: compose/crop/reverse/mute/bgm/enhance/gif/silent_cut/stt/stt_to_timeline/style/trim/concat/speed/resize/overlay/text/rotate/audio/filter/video_filter/transition/video_transition/tts/voice/frames/info"
            except ImportError as _e:
                return f"缺少依赖: {str(_e)}, 请先安装: pip install moviepy --break-system-packages"
            except Exception as _e:
                return f"视频剪辑失败: {str(_e)}"

        elif tool_name == "generate_ppt":
            try:
                title = args.get("title", "Presentation")
                pages = args.get("pages", [])
                theme_name = args.get("theme", "default")
                filename = args.get("filename", "")
                if not pages:
                    return "错误: 需要 pages 参数"
                from ppt_engine.build import build_pptx
                import shutil
                safe_name = (filename or 'output').replace('/', '_').replace('\\', '_')
                tmp_path = f"/tmp/ppt_{safe_name}.pptx"
                output = build_pptx(tmp_path, title, pages, theme_name)
                # Copy to web-accessible uploads dir
                web_dir = os.path.join(PROJECT_ROOT, 'uploads', 'shared')
                os.makedirs(web_dir, exist_ok=True)
                web_path = os.path.join(web_dir, f"ppt_{safe_name}.pptx")
                shutil.copy2(output, web_path)
                file_size_kb = os.path.getsize(web_path) // 1024
                web_url = f"https://naujtrats.xyz/oneapichat/uploads/shared/ppt_{safe_name}.pptx"
                # ★ 云盘全面结合: 同步到用户 OneAPIChat/generated
                try:
                    if user_id:
                        _import_to_cloudreve(user_id, web_path, "generated")
                except Exception:
                    pass
                return (
                    f"✅ PPT已生成\n"
                    f"📥 下载链接: {web_url}\n"
                    f"📦 大小: {file_size_kb} KB | 页数: {len(pages)}\n"
                    f"⚠️ 请直接复制上面的完整链接给用户，不要截断或省略。"
                )
            except ImportError as _e:
                return f"PPT生成缺少依赖: {str(_e)}"
            except Exception as _e:
                return f"PPT生成失败: {str(_e)}"

        # ★ 项目自描述：按需返回有界、可查询的事实，不把整份项目文档注入每轮
        elif tool_name == "project_self_describe":
            try:
                from engine.self_description import admit_self_context
                _self = admit_self_context(PROJECT_ROOT, query=str(args.get("query", "")), budget=min(int(args.get("budget", 14000) or 14000), 50000))
                return json.dumps(_self, ensure_ascii=False)
            except Exception as _e:
                return json.dumps({"ok": False, "error": "self-description unavailable", "detail": type(_e).__name__}, ensure_ascii=False)
        # ★ 通用转发: 子代理调用未知工具时自动转发到主引擎 API
        elif tool_name.startswith("server_") or tool_name == "engine_cron_list" or tool_name == "engine_cron_create" or tool_name == "engine_cron_delete":
            try:
                _engine_url = "http://127.0.0.1:8766/engine/" + {
                    "server_exec": "exec", "server_python": "python", "server_file_read": "file/read",
                    "server_file_write": "file/write", "server_file_write_chunked": "file/write_chunked",
                    "server_file_search": "file_search", "server_file_grep": "file_grep",
                    "server_file_edit": "file_edit",
                    "server_sys_info": "sys/info", "server_ps": "ps", "server_disk": "disk",
                    "server_network": "network", "server_docker": "docker", "server_db_query": "db_query",
                    "server_file_op": "file_op", "server_file_append": "file/write",
                    "engine_push": "agent/heartbeat"
                }.get(tool_name, tool_name)
                _params = {};
                for _k, _v in args.items():
                    if _v is None: continue
                    _params[_k] = str(_v);
                # ★ 复杂命令(含引号/特殊字符)走 POST body,避免 GET URL 转义/长度限制; 兼容 command/code 别名
                if tool_name == "server_exec":
                    _cmd = _params.get("cmd") or _params.get("command") or _params.get("query") or ""
                    _ep = {k: v for k, v in _params.items() if k in ("timeout", "cwd", "max_output", "user_id")}
                    _r = _http_session_no_proxy.post(_engine_url, params=_ep, data=_cmd.encode('utf-8') if isinstance(_cmd, str) else str(_cmd), headers={"Content-Type": "text/plain"}, timeout=max(int(_ep.get("timeout", 60) or 60) + 10, 5))
                elif tool_name == "server_python":
                    _script = _params.get("script") or _params.get("code") or _params.get("cmd") or ""
                    _ep = {k: v for k, v in _params.items() if k in ("timeout", "user_id")}
                    _r = _http_session_no_proxy.post(_engine_url, params=_ep, data=_script.encode('utf-8') if isinstance(_script, str) else str(_script), headers={"Content-Type": "text/plain"}, timeout=max(int(_ep.get("timeout", 30) or 30) + 10, 5))
                elif tool_name == "server_file_edit":
                    # ★ file_edit 端点要求 POST + JSON body (old_string/new_string), path 走 query
                    _ep = {k: v for k, v in _params.items() if k in ("path", "replace_all")}
                    _body = json.dumps({k: args[k] for k in ("old_string", "new_string") if args.get(k) is not None}, ensure_ascii=False)
                    _r = _http_session_no_proxy.post(_engine_url, params=_ep, data=_body, headers={"Content-Type": "application/json"}, timeout=30)
                else:
                    _r = _http_session_no_proxy.get(_engine_url, params=_params, timeout=30);
                _d = _r.json();
                return json.dumps(_d, ensure_ascii=False)
            except Exception as _e:
                return f"工具执行失败: {str(_e)}"
        # ★ P0: get_current_time 本地处理 (增强全球金融交易时区与市场开闭状态)
        elif tool_name == "get_current_time":
            from datetime import datetime as _dt, timezone as _tz, timedelta as _td
            _now_utc = _dt.now(_tz.utc)
            _bj_time = _now_utc + _td(hours=8)
            # 美东夏令时判定 (3-11月夏令时 UTC-4, 冬令时 UTC-5)
            _is_dst = 3 <= _now_utc.month <= 10 or (_now_utc.month == 11 and _now_utc.day < 7)
            _us_offset = 4 if _is_dst else 5
            _us_eastern = _now_utc - _td(hours=_us_offset)
            _tz_name = "EDT (UTC-4, 夏令时)" if _is_dst else "EST (UTC-5, 冬令时)"
            _london_offset = 1 if (3 <= _now_utc.month <= 10) else 0
            _london = _now_utc + _td(hours=_london_offset)
            _tokyo = _now_utc + _td(hours=9)
            _wk_zh = ["周一","周二","周三","周四","周五","周六","周日"][_bj_time.weekday()]
            _wk_en = ["Monday","Tuesday","Wednesday","Thursday","Friday","Saturday","Sunday"][_us_eastern.weekday()]

            def _us_m_st(dt):
                if dt.weekday() >= 5: return "🔴 周末休市 (Closed)"
                m = dt.hour * 60 + dt.minute
                if 9*60 + 30 <= m < 16*60: return "🟢 常规盘中交易中 (Regular Trading, 09:30~16:00)"
                elif 4*60 <= m < 9*60 + 30: return "🟡 盘前交易中 (Pre-market, 04:00~09:30)"
                elif 16*60 <= m < 20*60: return "🟡 盘后交易中 (After-hours, 16:00~20:00)"
                else: return "🔴 夜间闭市 (Closed)"

            def _cn_m_st(dt):
                if dt.weekday() >= 5: return "🔴 周末休市 (Closed)"
                m = dt.hour * 60 + dt.minute
                if (9*60 + 30 <= m < 11*60 + 30) or (13*60 <= m < 15*60): return "🟢 盘中交易中"
                elif 11*60 + 30 <= m < 13*60: return "🟡 午间休市"
                else: return "🔴 已收盘"

            def _hk_m_st(dt):
                if dt.weekday() >= 5: return "🔴 周末休市 (Closed)"
                m = dt.hour * 60 + dt.minute
                if (9*60 + 30 <= m < 12*60) or (13*60 <= m < 16*60): return "🟢 盘中交易中"
                elif 12*60 <= m < 13*60: return "🟡 午间休市"
                else: return "🔴 已收盘"

            return json.dumps({
                "beijing_time": f"{_bj_time.strftime('%Y-%m-%d %H:%M:%S')} {_wk_zh} (UTC+8)",
                "us_eastern_time": f"{_us_eastern.strftime('%Y-%m-%d %H:%M:%S')} {_wk_en} {_tz_name}",
                "london_time": f"{_london.strftime('%Y-%m-%d %H:%M:%S')} {'BST (UTC+1)' if _london_offset else 'GMT (UTC+0)'}",
                "tokyo_time": f"{_tokyo.strftime('%Y-%m-%d %H:%M:%S')} JST (UTC+9)",
                "iso_utc": _now_utc.strftime("%Y-%m-%dT%H:%M:%SZ"),
                "global_market_status": {
                    "us_stock_market": _us_m_st(_us_eastern),
                    "cn_a_stock_market": _cn_m_st(_bj_time),
                    "hk_stock_market": _hk_m_st(_bj_time),
                    "crypto_market": "🟢 全天候 24/7 交易中"
                },
                "time_zone_reminder": f"★ 重要时差提醒：美东时间比北京时间慢 {12 if _is_dst else 13} 小时。北京时间深夜/凌晨（21:30~次日04:00）正好是美股白天的常规盘中交易时间，绝非休市！分析美股与全球金融行情时请务必对齐美东交易日与实时状态。"
            }, ensure_ascii=False)

        # ★ 生图工具: 本地多提供商分发 (读 DB 配置, 支持 MiniMax/OpenAI/xAI/Custom)
        if tool_name in ('generate_image', 'generate_image_i2i'):
            return _exec_generate_image_local(user_id, tool_name, args)

        # ★ P0: MCP工具直连Node.js MCP服务器(port 18788),绕过PHP认证
        _mcp_prefixes = ('bilibili_', 'cr_', 'src_', 'chaoxing_', 'win_')
        if tool_name.startswith(_mcp_prefixes):
            try:
                _mcp_endpoint = "/bilibili/tools/call" if tool_name.startswith("bilibili_") else "/api/tools/call"
                _mcp_url = f"http://127.0.0.1:18788{_mcp_endpoint}"
                _resp = _http_session.post(_mcp_url, json={"name": tool_name, "arguments": args}, timeout=30)
                if _resp.ok:
                    _text = _resp.text
                    try:
                        _d = json.loads(_text)
                        if isinstance(_d, dict):
                            return json.dumps(_d.get("result", _d), ensure_ascii=False)
                    except Exception:
                        pass
                    return _text[:4000]
                return f"MCP工具 {tool_name} 返回HTTP {_resp.status_code}"
            except Exception as _e:
                return f"MCP工具 {tool_name} 转发异常: {str(_e)}"
        return f"[警告: 当前环境不支持 {tool_name} 工具] 跳过此操作,请用 web_search/web_fetch 替代"

    # ★ P0+P1: 加载子代理系统上下文(Skills + 人格/记忆/项目)
    _agent_system_parts = []

    # 1. 可用技能列表
    try:
        from engine.skills import list_skills
        _skills = list_skills(user_id)
        if _skills:
            _skill_lines = ["## 可用技能\n你可以调用 `run_skill` 工具执行以下预定义技能："]
            for _s in _skills:
                if _s.get("enabled", True):
                    _skill_lines.append(f"- **{_s['name']}**: {_s.get('description', '无描述')}")
            _agent_system_parts.append("\n".join(_skill_lines))
    except Exception:
        pass

    # 2. P1: 用户人格/记忆
    _mem_dir = Path(PROJECT_ROOT) / ".engine" / "memory"
    try:
        _mem_dir.mkdir(parents=True, exist_ok=True)
        _persona_file = _mem_dir / f"user_{user_id}_agent_persona.json"
        if _persona_file.exists():
            with open(_persona_file) as f:
                _p = json.load(f)
            if _p.get("name") or _p.get("style"):
                _agent_system_parts.append(f"## 用户人格\n- 名称: {_p.get('name', '未设置')}\n- 风格: {_p.get('style', '未设置')}")
    except Exception:
        pass

    # 3. P1: 最近记忆(最多5条)
    try:
        _mem_file = _mem_dir / f"user_{user_id}_agent_memory.json"
        if _mem_file.exists():
            with open(_mem_file) as f:
                _m = json.load(f)
            _entries = _m.get("entries", [])[-5:]
            if _entries:
                _lines = ["## 用户记忆"]
                for _e in _entries:
                    _lines.append(f"- {_e.get('key', '')}: {_e.get('content', '')[:200]}")
                _agent_system_parts.append("\n".join(_lines))
    except Exception:
        pass

    # 4. 有界项目自描述：按当前任务关键词选择，不把整份 CLAUDE.md 无条件塞入上下文
    try:
        _self_context = admit_self_context(PROJECT_ROOT, query=_run_prompt, budget=12000)
        if _self_context.get("text"):
            _agent_system_parts.append("## 项目自描述与运行规则\n" + _self_context["text"])
    except Exception:
        pass
    # 兼容已有项目规则，但只取稳定的操作约束摘要
    try:
        _claude_path = Path(PROJECT_ROOT) / "CLAUDE.md"
        if _claude_path.exists():
            _claude = _claude_path.read_text(encoding="utf-8")[:1200]
            _agent_system_parts.append(f"## 兼容项目规则摘要\n{_claude}")
    except Exception:
        pass

    # ★ 重要: 告诉子代理在获取足够数据后停止工具调用,返回最终答案
    _agent_system_parts.append(
        "## ⚠️ 停止规则\n"
        "1. 当你已经获取了足够的信息来回答用户的问题时,停止使用工具\n"
        "2. 最后一轮必须只输出最终答案(不再调用任何工具)\n"
        "3. 最终答案要综合所有已获取的信息,用中文条理清晰地输出\n"
        "4. 不要无限制地搜索,通常 2-3 轮工具调用后就应该总结输出"
    )
    _agent_system_context = "\n\n".join(_agent_system_parts) if _agent_system_parts else ""

    def _run_body():
        _lock = _get_agent_store_lock(user_id)
        _lock.acquire()
        try:
            current_agents = store.get()
            if name not in current_agents:
                current_agents[name] = agent
            current_agents[name]["status"] = "running"
            current_agents[name]["result"] = ""
            current_agents[name]["_started_at"] = time.time()
            store.set(current_agents)
        finally:
            _lock.release()

        MAX_EXECUTION_SECONDS = 600  # 10分钟强制超时
        _child_session_id = _runtime_subagent["child_session_id"]
        _turn_id = f"turn_{uuid.uuid4().hex}"
        try:
            agent_runtime.update_subagent(_runtime_subagent["agent_id"], "running", user_id=user_id)
            agent_runtime.append_event(_child_session_id, "turn/start", {"turn_id": _turn_id, "agent_id": _runtime_subagent["agent_id"]})
            agent_runtime.append_event(_child_session_id, "user/message", {"content": _run_prompt, "continuation": _is_followup})
            messages = []
            if _agent_system_context:
                messages.append({"role": "system", "content": _agent_system_context})
            if _is_followup and _previous_result:
                messages.extend([
                    {"role": "user", "content": _base_prompt},
                    {"role": "assistant", "content": _previous_result},
                    {"role": "user", "content": _run_prompt},
                ])
            else:
                messages.append({"role": "user", "content": _run_prompt})
            max_rounds = max_agent_rounds
            result_parts = []
            start_time = time.time()
            # ★ 死循环检测: 小参数模型易陷入工具复读/振荡死循环烧 token (python/engine/loop_guard.py)
            from engine.loop_guard import LoopGuard
            _lg = LoopGuard()
            _force_summary = False
            _active_step_id = None

            for round_num in range(max_rounds):
                # 检查总执行时间
                if time.time() - start_time > MAX_EXECUTION_SECONDS:
                    raise TimeoutError(f"子代理执行超过{MAX_EXECUTION_SECONDS//60}分钟,自动终止")

                # ★ 最终轮: 移除工具强制模型总结输出(不再调用工具)
                _is_final_round = (round_num >= max_rounds - 1)
                _tools_for_call = None if _is_final_round else TOOLS

                if _cancel_event.is_set() or agent_runtime.get_subagent(
                    _runtime_subagent["agent_id"], user_id=user_id
                ).get("cancel_requested"):
                    raise AgentRuntimeError(ErrorCode.CANCELLED, "subagent cancelled", status=499)

                _active_step_id = f"step_{uuid.uuid4().hex}"
                agent_runtime.append_event(_child_session_id, "step/start", {
                    "step_id": _active_step_id, "round": round_num + 1,
                })
                _completion = None
                _last_provider_error = None
                for _provider_attempt in range(5):
                    try:
                        _completion = complete_chat({
                            "model": model,
                            "base_url": base_url,
                            "anthropic_format": bool(main_config.get("anthropic_format")),
                            "messages": messages,
                            "tools": _tools_for_call or [],
                            "tool_choice": "auto" if not _is_final_round else "none",
                            "temperature": 0.3,
                            "max_tokens": 4096 if _is_final_round else 2048,
                        }, api_key, http_client=_agent_http_client, timeout=120)
                        if _completion.content.strip() or _completion.tool_calls:
                            break
                    except AgentRuntimeError as _provider_error:
                        _last_provider_error = _provider_error
                        if not _provider_error.retryable or _provider_attempt >= 4:
                            raise
                    if _provider_attempt < 4:
                        _retry_after = 0.0
                        if _last_provider_error and isinstance(_last_provider_error.details, dict):
                            try:
                                _retry_after = float(_last_provider_error.details.get("retry_after") or 0)
                            except (TypeError, ValueError):
                                _retry_after = 0.0
                        time.sleep(max(_retry_after, min(8.0, 1.5 * (2 ** _provider_attempt))))
                if _completion is None or (not _completion.content.strip() and not _completion.tool_calls):
                    raise AgentRuntimeError(
                        ErrorCode.EMPTY_RESPONSE, "provider returned an empty response",
                        status=502, retryable=True, details={"provider": _provider_info.name},
                    )
                asst_msg = _completion.as_openai_message()
                agent_runtime.append_event(_child_session_id, "assistant/message", {
                    "content": _completion.content,
                    "reasoning": _completion.reasoning,
                    "tool_calls": _completion.tool_calls,
                    "usage": _completion.usage,
                    "provider": _completion.provider.name,
                    "model": model,
                })
                if _completion.content:
                    cleaned = _completion.content
                    if '<think>' in cleaned or '</think>' in cleaned:
                        cleaned = re.sub(r'<think>.*?</think>', '', cleaned, flags=re.DOTALL).strip()
                    result_parts.append(cleaned)
                if not _completion.tool_calls:
                    agent_runtime.append_event(_child_session_id, "step/end", {
                        "step_id": _active_step_id, "status": "completed",
                    })
                    _active_step_id = None
                    break

                # ★ 早停: 如果已经获取了足够多数据(>3000字符),强制进入总结轮
                _total_result_len = sum(len(str(p)) for p in result_parts)
                if _total_result_len > 3000 and round_num >= 2:
                    # 下一轮将是最终总结轮(无工具)
                    max_rounds = min(max_rounds, round_num + 2)

                # Canonical completion is already safe to feed back to every provider adapter.
                messages.append(asst_msg)

                for tc in _completion.tool_calls:
                    _tc_function = tc.get("function") or {}
                    _tc_id = str(tc.get("id") or f"call_{uuid.uuid4().hex}")
                    tool_name = str(_tc_function.get("name") or "")
                    _raw_tool_args = _tc_function.get("arguments") or "{}"
                    try:
                        tool_args = _raw_tool_args if isinstance(_raw_tool_args, dict) else json.loads(_raw_tool_args)
                    except json.JSONDecodeError as _je:
                        # ★ 容错: server_exec/server_python 参数含引号导致JSON非法时尝试恢复
                        tool_args = _tolerant_tool_args(tool_name, str(_raw_tool_args))
                        if tool_args is None:
                            _invalid_args_error = AgentRuntimeError(
                                ErrorCode.INVALID_TOOL_ARGUMENTS,
                                f"invalid JSON arguments: {_je.msg}", status=400,
                                details={"line": _je.lineno, "column": _je.colno},
                            )
                            agent_runtime.append_event(_child_session_id, "tool/call", {
                                "call_id": _tc_id, "name": tool_name,
                                "arguments": json.dumps({"invalid_json": True, "length": len(str(_raw_tool_args))}),
                            })
                            agent_runtime.append_event(_child_session_id, "tool/result", {
                                "call_id": _tc_id, "name": tool_name, "error": _invalid_args_error.to_dict(),
                            })
                            result_parts.append(f"[工具: {tool_name}] 参数解析失败: 跳过")
                            messages.append({"role": "tool", "tool_call_id": _tc_id, "content": f"[错误] {tool_name} 参数解析失败: {str(_je)}, 请检查参数格式后重试"})
                            continue
                        result_parts.append(f"[工具: {tool_name}] 参数JSON非法,已容错恢复: {str(tool_args)[:80]}")
                    if not isinstance(tool_args, dict):
                        tool_args = {}
                    # ★ 参数别名兼容: command→cmd, code→script
                    if tool_name == "server_exec" and not tool_args.get("cmd"):
                        tool_args["cmd"] = tool_args.get("command") or tool_args.get("query") or ""
                    elif tool_name == "server_python" and not tool_args.get("script"):
                        tool_args["script"] = tool_args.get("code") or tool_args.get("cmd") or ""
                    # ★ 死循环检测: 重复/振荡调用仍经过统一管道记录，但不执行副作用。
                    _skip_tool = _lg.record_tool_call(tool_name, tool_args) == "skip"
                    if _skip_tool:
                        _guard_msg = f"【系统提示】系统检测到{_lg.last_reason},该调用未执行。请立即停止调用工具,直接根据已有信息输出最终总结。"
                        _force_summary = True
                        _legacy_handler = lambda _args, _context: _guard_msg
                    else:
                        _legacy_handler = lambda _args, _context, _name=tool_name: _execute_tool(_name, _args)
                    _tool_result = tool_pipeline.execute(
                        tool_name, tool_args,
                        ToolExecutionContext(
                            user_id=user_id, session_id=_child_session_id, call_id=_tc_id,
                            agent_id=_runtime_subagent["agent_id"], mode="legacy-subagent",
                            approved=True, cancel_event=_cancel_event,
                        ),
                        handler=_legacy_handler,
                    )
                    if _tool_result.ok:
                        result = _tool_result.value
                    else:
                        _tool_error = _tool_result.error or {}
                        result = f"[工具执行异常] {tool_name}: {_tool_error.get('message', 'unknown error')}"
                    # Log metadata only; tool output may contain credentials or private file contents.
                    _result_size = len(str(result))
                    print(f"[子代理:{name}] 工具调用完成: tool={tool_name} ok={_tool_result.ok} size={_result_size}", flush=True)
                    # ★ 全局净化：移除所有控制字符和 unicode surrogate
                    if isinstance(result, str):
                        result = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', result)
                        if len(result) > 8000:
                            result = result[:8000] + '...(截断)'
                    # ★ 工具结果保存到 result_parts(用于最终 agent 结果),截断到 2000 字符供 AI 分析
                    result_parts.append(f"[工具: {tool_name}] {str(result)[:2000]}")
                    # ★ put 结果时再做一次安全包装
                    safe_content = str(result) if result else '(empty)'
                    messages.append({"role": "tool", "tool_call_id": _tc_id, "content": safe_content})
                    # ★ P1: 广播子代理步骤进度
                    try:
                        _broadcast_to_user(user_id, 'agent:step', {
                            'agent': name,
                            'tool': tool_name,
                            'step': round_num + 1,
                            'max_steps': max_rounds,
                            'result_preview': str(result)[:100]
                        })
                    except Exception:
                        pass
                    # ★ 实时写入 partial result(带锁+重读,防止覆盖其他代理)
                    _lock.acquire()
                    try:
                        current = store.get()
                        current[name] = current.get(name, {})
                        current[name]["result"] = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', "\n".join(result_parts))
                        current[name]["status"] = "running"
                        current[name]["_started_at"] = current.get(name, {}).get("_started_at", time.time())
                        store.set(current)
                    finally:
                        _lock.release()

                agent_runtime.append_event(_child_session_id, "step/end", {
                    "step_id": _active_step_id, "status": "completed",
                    "tool_count": len(_completion.tool_calls),
                })
                _active_step_id = None

                # ★ 死循环检测(轮末): 连续纯工具轮 → 终止; 复读/无进展 → 强制总结轮
                _round_action = _lg.record_round(
                    had_content=bool(_completion.content), tool_count=len(_completion.tool_calls)
                )
                if _round_action == "abort":
                    result_parts.append(f"[系统检测到死循环: {_lg.last_reason},已终止本轮工作]")
                    try:
                        _broadcast_to_user(user_id, 'agent:warning', {'agent': name, 'msg': f'检测到死循环({_lg.last_reason}),已终止'})
                    except Exception:
                        pass
                    break
                _rep = _lg.feed_text("\n".join(result_parts))
                if _rep:
                    result_parts.append(f"[系统检测到死循环: {_rep}]")
                    max_rounds = min(max_rounds, round_num + 1)  # 复用早停: 下一轮为无工具总结轮
                    try:
                        _broadcast_to_user(user_id, 'agent:warning', {'agent': name, 'msg': f'检测到死循环({_rep}),已注入强制总结'})
                    except Exception:
                        pass
                if _force_summary:
                    max_rounds = min(max_rounds, round_num + 1)  # 强制进入最终总结轮(无工具)

            # Stable structured projection generated locally; the full raw result remains authoritative.
            _raw_content = "\n".join(result_parts)
            _error_lines = [line[:500] for line in result_parts if "错误" in line or "异常" in line]
            _action_lines = [line[:500] for line in result_parts if line.startswith("[工具:")]
            _summary_source = next((part for part in reversed(result_parts) if part and not part.startswith("[工具:")), _raw_content)
            _structured_output = {
                "summary": str(_summary_source or "")[:1000],
                "findings": [],
                "actions_taken": _action_lines[-100:],
                "errors": _error_lines[-100:],
                "raw_output": _raw_content[:20000],
            }

            # ★ 最终保存(带锁+重读)
            _lock.acquire()
            try:
                current = store.get()
                current[name] = current.get(name, {})
                final_result = "\n".join(result_parts)
                final_result = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', final_result)
                current[name]["result"] = final_result
                current[name].pop("error", None)
                if _structured_output:
                    current[name]["_structured"] = json.dumps(_structured_output, ensure_ascii=False)
                _conversation = current[name].get("_conversation")
                if not isinstance(_conversation, list):
                    _conversation = []
                _conversation.extend([
                    {"role": "user", "content": _run_prompt[:20000]},
                    {"role": "assistant", "content": final_result[:50000]},
                ])
                current[name]["_conversation"] = _conversation[-20:]
                current[name]["status"] = "completed"
                store.set(current)
            finally:
                _lock.release()
            agent_runtime.update_subagent(
                _runtime_subagent["agent_id"], "completed", user_id=user_id,
                result={"text": final_result, "structured": _structured_output},
            )
            agent_runtime.append_event(_child_session_id, "turn/end", {
                "turn_id": _turn_id, "status": "completed",
            })
        except Exception as e:
            _runtime_error = e if isinstance(e, AgentRuntimeError) else AgentRuntimeError(
                ErrorCode.INTERNAL, str(e), status=500
            )
            _cancelled = _runtime_error.code == ErrorCode.CANCELLED
            if _active_step_id:
                try:
                    agent_runtime.append_event(_child_session_id, "step/end", {
                        "step_id": _active_step_id, "status": "cancelled" if _cancelled else "failed",
                        "error": _runtime_error.to_dict(),
                    })
                except Exception:
                    pass
            try:
                agent_runtime.update_subagent(
                    _runtime_subagent["agent_id"], "cancelled" if _cancelled else "failed",
                    user_id=user_id, error=_runtime_error.to_dict(),
                )
                agent_runtime.append_event(_child_session_id, "turn/end", {
                    "turn_id": _turn_id, "status": "cancelled" if _cancelled else "failed",
                    "error": _runtime_error.to_dict(),
                })
            except Exception:
                pass
            _lock.acquire()
            try:
                current = store.get()
                current[name] = current.get(name, {})
                current[name]["error"] = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', _runtime_error.message)
                current[name]["status"] = "stopped" if _cancelled else "failed"
                store.set(current)
            finally:
                _lock.release()
        # ★ 通知引擎:此代理已完成,需要主代理处理
        # ★ 从 store 重新读取最新状态(不能用 agents 变量, _run 只写 store 不写 agents)
        _latest = store.get()
        _latest_agent = _latest.get(name, agents.get(name, {}))
        notify_store = get_ns("agent_notifications", user_id)
        notifs = notify_store.get()
        if not isinstance(notifs, list):
            notifs = []
        notifs.append({
            "agent": name,
            "status": _latest_agent.get("status", "unknown"),
            "result": _latest_agent.get("result", ""),
            "error": _latest_agent.get("error", ""),
            "structured": _latest_agent.get("_structured", ""),
            "time": datetime.now().isoformat(),
            "processed": False
        })
        # 裁剪超过50条的历史通知(防止内存泄漏)
        if len(notifs) > 50:
            notifs = notifs[-50:]
        notify_store.set(notifs)
        # Broadcast agent status change via SSE
        _broadcast_to_user(user_id, 'agent:status', {
            'agent': name,
            'status': _latest_agent.get('status', 'unknown'),
            'result_preview': (_latest_agent.get('result', '') or '')[:200],
            'time': notifs[-1]['time']
        })
        # ★ P0: 广播完整结果(含结构化数据),前端即时消费无需轮询
        _full_result = _latest_agent.get('result', '') or ''
        _structured_raw = _latest_agent.get('_structured', '')
        _structured_data = json.loads(_structured_raw) if _structured_raw else None
        # 服务端原子落盘子代理会话，无论是否有客户端在线，任务结果永不丢失
        try:
            chat_projection_store.commit_subagent(
                user_id=user_id,
                agent_name=name,
                prompt=_run_prompt,
                result=_full_result,
                error=_latest_agent.get('error', ''),
                status=_latest_agent.get('status', 'completed'),
            )
            _broadcast_to_user(user_id, 'chat:updated', {
                'chat_id': f'_agent_sub_{name}',
                'finished': True,
                'ts': time.time(),
            })
        except Exception as _sub_proj_err:
            print(f"[subagent] commit_subagent warning: {_sub_proj_err}")
        _broadcast_to_user(user_id, 'agent:result', {
            'agent': name,
            'status': _latest_agent.get('status', 'unknown'),
            'result': _full_result,
            'structured': _structured_data,
            'error': _latest_agent.get('error', ''),
            'time': notifs[-1]['time']
        })

    def _run():
        try:
            _run_body()
        finally:
            if _agent_http_client is not None:
                try:
                    _agent_http_client.close()
                except Exception:
                    pass

    t = threading.Thread(target=_run, name=f"agent_{user_id}_{name}", daemon=True)
    t.start()
    return {
        "ok": True, "agent": name, "status": "running",
        "runtime_agent_id": _runtime_subagent["agent_id"],
        "runtime_session_id": _runtime_subagent["child_session_id"],
    }

@app.post("/engine/agent/run")
async def agent_run_post(request: Request, user_id: str = Query("")):
    body = await request.json()
    return agent_run(
        name=str(body.get("name") or ""), user_id=user_id,
        message=str(body.get("message") or ""),
        from_ask="1" if body.get("from_ask") in (True, 1, "1") else "",
    )

@app.get("/engine/agent/status")
def agent_status(name: str = Query(...), user_id: str = Query("")):
    store = get_ns("agents", user_id)
    agents = store.get()
    agent = agents.get(name)
    if not agent:
        raise HTTPException(404, f"Agent {name} not found")
    return agent

# ==================== 主聊配置读取(所有 agent 同步主聊)====================
import base64

def _load_encryption_key() -> str:
    return load_encryption_key(PROJECT_ROOT)
ENCRYPTION_KEY = _load_encryption_key()  # ★ 必须在 _get_aes_key/_decrypt_xor 前初始化
def _get_aes_key() -> bytes:
    return get_aes_key(ENCRYPTION_KEY)
def _decrypt_xor(encoded: str) -> str:
    _aes = _get_aes_key() if (encoded and encoded.startswith("v2:")) else None
    return decrypt_xor(encoded, ENCRYPTION_KEY, _aes)
def _get_main_chat_config(user_id: str) -> dict:
    """Load the authenticated user's encrypted chat configuration locally.

    The engine already has an owner-bound user_id. A former HTTPS self-request used a
    fabricated session token, always failed authentication, disabled TLS verification,
    and then fell back to this same file. Read the per-user source directly instead.
    """
    result = {"api_key": "", "base_url": "", "model": "", "provider": "", "anthropic_format": False}
    if not user_id:
        return result
    config_path = os.path.join(PROJECT_ROOT, f"chat_data/config_user_{user_id}.json")
    try:
        with open(config_path, encoding="utf-8") as config_file:
            cfg = json.load(config_file)
    except Exception:
        return result

    # ★ Provider 映射(与前端 API_PROVIDERS 一致)
    PROVIDER_MAP = {
        "openai":    {"key_field": "apiKeyOpenAI",    "base": "https://api.openai.com/v1",         "default_model": "gpt-4o"},
        "deepseek":  {"key_field": "apiKeyDeepseek",  "base": "https://api.deepseek.com/v1",       "default_model": "deepseek-v4-flash-vision-exp", "vision_model": "deepseek-v4-flash-vision-exp"},
        "longcat":   {"key_field": "apiKeyLongCat",   "base": "https://api.longcat.chat/openai/v1", "default_model": "LongCat-2.0"},
        "minimax":   {"key_field": "apiKeyMiniMax",   "base": "https://api.minimaxi.com/v1",       "default_model": "MiniMax-M2.7"},
        "anthropic": {"key_field": "apiKeyAnth",      "base": "https://api.anthropic.com/v1",      "default_model": "claude-sonnet-4-20250514"},
        "gemini":    {"key_field": "apiKeyGemini",    "base": "https://generativelanguage.googleapis.com/v1beta/openai", "default_model": "gemini-2.0-flash"},
        "moonshot":  {"key_field": "apiKeyMoonshot",  "base": "https://api.moonshot.cn/v1",        "default_model": "moonshot-v1-8k"},
        "xai":       {"key_field": "apiKeyXAI",       "base": "https://api.x.ai/v1",                "default_model": "grok-4-latest"},
        "zhipu":     {"key_field": "apiKeyZhipu",     "base": "https://open.bigmodel.cn/api/paas/v4", "default_model": "glm-4-flash"},
        "opencode":  {"key_field": "apiKeyOpenCode",  "base": "https://opencode.ai/v1",            "default_model": "claude-sonnet-4-5"},
        "nvidia":    {"key_field": "apiKeyNvidia",    "base": "https://integrate.api.nvidia.com/v1", "default_model": "deepseek-ai/deepseek-v4-flash"},
        "openrouter":{"key_field": "apiKeyOpenRouter","base": "https://openrouter.ai/api/v1",      "default_model": "openai/gpt-4o"},
        "mimo":      {"key_field": "apiKeyMiMo",      "base": "https://api.mimo.ai/v1",             "default_model": "mimo-v4-flash"},
        "custom":    {"key_field": "apiKeyCustom",    "base": "",                                    "default_model": ""},
    }

    # 1. Determine the active provider. Older browser settings can retain a stale
    # baseUrlProvider after the endpoint has been switched (for example LongCat →
    # DeepSeek); let a known endpoint correct a known, conflicting provider.
    provider = cfg.get("baseUrlProvider", "") or ""
    raw_base_url = str(cfg.get("baseUrl", "") or "").strip()
    endpoint_provider = ""
    if raw_base_url:
        try:
            from urllib.parse import urlparse
            raw_host = (urlparse(raw_base_url).hostname or "").lower()
            for candidate, candidate_config in PROVIDER_MAP.items():
                candidate_host = (urlparse(candidate_config.get("base", "")).hostname or "").lower()
                if raw_host and candidate_host and raw_host == candidate_host:
                    endpoint_provider = candidate
                    break
        except Exception:
            pass
    if endpoint_provider and (not provider or (provider != "custom" and provider != endpoint_provider)):
        provider = endpoint_provider
    if not provider:
        model_select = (cfg.get("modelSelect", "") or "").lower()
        if "deepseek" in model_select: provider = "deepseek"
        elif "gpt" in model_select: provider = "openai"

    pm = PROVIDER_MAP.get(provider, {})
    result["provider"] = provider or "custom"
    result["anthropic_format"] = provider == "anthropic"

    # ★ 2. 读 key — 优先读取当前 provider 的专属 Key 字段，回退通用 apiKey，再次回退 apiKeyCustom
    key_field = pm.get("key_field", "")
    stored_key = ""
    if key_field and cfg.get(key_field):
        stored_key = cfg.get(key_field)
    elif provider == "custom" and cfg.get("apiKeyCustom"):
        stored_key = cfg.get("apiKeyCustom")
    elif cfg.get("apiKey"):
        stored_key = cfg.get("apiKey")
    elif cfg.get("apiKeyCustom"):
        stored_key = cfg.get("apiKeyCustom")
    elif cfg.get("apiKeyDeepseek"):
        stored_key = cfg.get("apiKeyDeepseek")

    if stored_key:
        decrypted = _decrypt_xor(stored_key)
        result["api_key"] = decrypted if decrypted else stored_key

    # 3. base_url
    custom_base = cfg.get("baseUrlCustom", "") or ""
    if custom_base and provider == "custom":
        result["base_url"] = custom_base
    elif pm.get("base"):
        result["base_url"] = pm["base"]
    else:
        result["base_url"] = cfg.get("baseUrl", "") or "https://api.deepseek.com/v1"

    # 4. model — 优先 provider-specific, 回退 provider 默认
    provider_model = cfg.get(f"model_{provider}", "") or ""
    generic_model = cfg.get("model", "") or ""
    result["model"] = provider_model or generic_model or pm.get("default_model", "gpt-4o")

    if not result["model"] or result["model"].startswith("加载中") or len(result["model"]) < 3:
        result["model"] = pm.get("default_model", "gpt-4o")

    print(f"[引擎] 主聊配置: provider={provider} model={result['model']} credential_configured={bool(result['api_key'])}", flush=True)
    return result


def _goal_create_agent(name: str, prompt: str, user_id: str):
    return agent_create(
        name=name, prompt=prompt, role="general", model="", base_url="",
        user_id=user_id, proxy_url="", proxy_enabled="",
    )

def _goal_run_agent(name: str, user_id: str, message: str, followup: bool):
    return agent_run(name=name, user_id=user_id, message=message, from_ask="1" if followup else "")

def _goal_get_agent(name: str, user_id: str):
    return get_ns("agents", user_id).get().get(name)

_goal_runner = DurableGoalRunner(
    agent_runtime, create_agent=_goal_create_agent, run_agent=_goal_run_agent,
    get_agent=_goal_get_agent, stop_agent=lambda name, user_id: agent_stop(name=name, user_id=user_id),
)

@app.on_event("startup")
def _resume_durable_goals_on_startup():
    resumed = _goal_runner.resume_active()
    if resumed:
        print(f"[runtime] resumed {resumed} durable goal(s)", flush=True)

_runtime_maintenance_task = None

async def _runtime_cache_maintenance_loop():
    while True:
        await asyncio.sleep(900)
        _prune_user_runtime_caches(3600)
        await prune_browser_managers(3600)

@app.on_event("startup")
async def _start_runtime_cache_maintenance():
    global _runtime_maintenance_task
    _runtime_maintenance_task = asyncio.create_task(_runtime_cache_maintenance_loop())

@app.on_event("shutdown")
async def _stop_runtime_cache_maintenance():
    global _runtime_maintenance_task
    if _runtime_maintenance_task is not None:
        _runtime_maintenance_task.cancel()
        try:
            await _runtime_maintenance_task
        except asyncio.CancelledError:
            pass
        _runtime_maintenance_task = None
    _prune_user_runtime_caches(0)
    await prune_browser_managers(0)


# ★ 服务器操控工具 → engine/server_tools.py
register_server_tools(app)
@app.get("/engine/agent/stop")
def agent_stop(name: str = Query(...), user_id: str = Query("")):
    """停止子代理(标记为 stopped)"""
    store = get_ns("agents", user_id)
    agents = store.get()
    agent = agents.get(name)
    if not agent:
        return {"ok": False, "error": "Agent not found"}
    _get_agent_cancel_event(user_id, name).set()
    _runtime_agent_id = str(agent.get("runtime_agent_id") or "")
    if _runtime_agent_id:
        try:
            agent_runtime.request_subagent_cancel(_runtime_agent_id, user_id, "stopped by user")
        except AgentRuntimeError:
            pass
    agents[name]["status"] = "stopping" if agent.get("status") == "running" else "stopped"
    store.set(agents)
    return {"ok": True, "agent": name, "status": agents[name]["status"]}

@app.get("/engine/agent/delete")
def agent_delete(name: str = Query(...), user_id: str = Query("")):
    """删除子代理(从列表中移除)"""
    store = get_ns("agents", user_id)
    agents = store.get()
    if name not in agents:
        return {"ok": False, "error": "Agent not found"}
    del agents[name]
    store.set(agents)
    return {"ok": True, "agent": name, "deleted": True}

@app.get("/engine/agent/notifications")
def agent_notifications(user_id: str = Query("")):
    """获取未处理的子代理完成通知(主代理调用)"""
    store = get_ns("agent_notifications", user_id)
    notifs = store.get()
    if not isinstance(notifs, list):
        store.set([])
        notifs = []
    unprocessed = [n for n in notifs if not n.get("processed", False)]
    return {"notifications": unprocessed, "count": len(unprocessed)}

@app.get("/engine/agent/notifications/mark")
def agent_notifications_mark(user_id: str = Query("")):
    """标记所有通知为已处理"""
    store = get_ns("agent_notifications", user_id)
    notifs = store.get()
    if not isinstance(notifs, list):
        store.set([])
        return {"ok": True}
    for n in notifs:
        n["processed"] = True
    store.set(notifs)
    return {"ok": True}
    return {"ok": True, "workflow": name, "steps": len(parsed_steps)}

    return {"ok": True, "workflow": name, "status": "running"}

    return wf_store.get()

    return wf

    return {"ok": True}

    return {"roles": [{"id": k, "label": v["label"], "desc": v["desc"]} for k, v in AGENT_ROLES.items()]}



# ==================== 工作流引擎 (核心逻辑→engine.workflow) ====================

@app.get("/engine/workflow/create")
def workflow_create(name: str = Query(...), steps: str = Query(...), user_id: str = Query("")):
    return create_workflow(name, steps, user_id, get_ns)

@app.get("/engine/workflow/run")
def workflow_run(name: str = Query(...), user_id: str = Query("")):
    return run_workflow(name, user_id, get_ns, _get_main_chat_config, AGENT_ROLES, _filter_tools_by_role)

@app.get("/engine/workflow/list")
def workflow_list(user_id: str = Query("")):
    return list_workflows(user_id, get_ns)

@app.get("/engine/workflow/status")
def workflow_status(name: str = Query(...), user_id: str = Query("")):
    return status_workflow(name, user_id, get_ns)

@app.get("/engine/workflow/delete")
def workflow_delete(name: str = Query(...), user_id: str = Query("")):
    return delete_workflow(name, user_id, get_ns)

@app.get("/engine/workflow/roles")
def workflow_roles(user_id: str = Query("")):
    return _wf_get_roles(AGENT_ROLES)

# ==================== 引擎层 API ====================

@app.get("/engine/v2/exec-policy/evaluate")
def exec_policy_evaluate(
    domain: str = Query("exec"),
    target: str = Query(...),
    user_id: str = Query("")
):
    """评估一个操作是否需要审批"""
    policy = _get_exec_policy(user_id)
    decision = policy.evaluate(domain, target)
    return {
        "ok": True,
        "domain": domain,
        "target": target,
        "decision": decision.kind,
        "reason": decision.reason,
        "matched_rule": decision.matched_rule,
        "matched_priority": decision.matched_priority,
    }


@app.get("/engine/v2/exec-policy/rules")
def exec_policy_rules(
    domain: str = Query(""),
    user_id: str = Query("")
):
    """获取策略规则列表"""
    policy = _get_exec_policy(user_id)
    return {"ok": True, "rules": policy.list_rules(domain), "count": len(policy.rules)}


@app.get("/engine/v2/exec-policy/add")
def exec_policy_add(
    domain: str = Query("exec"),
    pattern: str = Query(...),
    decision_kind: str = Query("skip"),
    reason: str = Query(""),
    priority: int = Query(2),
    description: str = Query(""),
    user_id: str = Query("")
):
    """添加策略规则"""
    if decision_kind == "skip":
        decision = ExecDecision.skip()
    elif decision_kind == "forbidden":
        decision = ExecDecision.forbidden(reason or "禁止操作")
    else:
        decision = ExecDecision.needs_approval(reason or "需要审批")
    rule = _get_exec_policy(user_id).add_rule(domain, pattern, decision, priority=priority, description=description)
    return {"ok": True, "rule": rule.to_dict()}


@app.get("/engine/v2/exec-policy/remove")
def exec_policy_remove(
    domain: str = Query("exec"),
    pattern: str = Query(...),
    priority: int = Query(-1),
    user_id: str = Query("")
):
    """移除策略规则"""
    p = Priority(priority) if priority >= 0 else None
    removed = _get_exec_policy(user_id).remove_rule(domain, pattern, p)
    return {"ok": removed}


@app.get("/engine/v2/exec-policy/reset")
def exec_policy_reset(user_id: str = Query("")):
    """重置为默认规则"""
    policy = _get_exec_policy(user_id)
    policy.reset_to_defaults()
    return {"ok": True, "rules": len(policy.rules)}


# ── 推测执行 API ─────────────────────────────────────

@app.get("/engine/v2/speculate")
def speculate(
    prompt: str = Query(...),
    user_id: str = Query("")
):
    """推测指令需要的工具调用"""
    result = _get_speculation_engine(user_id).predict(prompt)
    return {
        "ok": True,
        "suggested_tools": [
            {"tool_name": t.tool_name, "confidence": t.confidence,
             "estimated_duration_ms": t.estimated_duration_ms}
            for t in result.suggested_tools
        ],
        "estimated_savings_ms": result.estimated_savings_ms,
    }


@app.get("/engine/v2/speculate/confirm")
def speculate_confirm(user_id: str = Query("")):
    """确认推测结果（命中）"""
    engine = _get_speculation_engine(user_id)
    engine.confirm()
    return {"ok": True, "state": engine.state.value}


@app.get("/engine/v2/speculate/abort")
def speculate_abort(
    reason: str = Query("用户中止"),
    user_id: str = Query("")
):
    """中止推测"""
    engine = _get_speculation_engine(user_id)
    engine.abort(reason=reason)
    return {"ok": True, "state": engine.state.value}


@app.get("/engine/v2/speculate/status")
def speculate_status(user_id: str = Query("")):
    """推测引擎状态"""
    return {"ok": True, **_get_speculation_engine(user_id).summary()}


@app.get("/engine/v2/speculate/toggle")
def speculate_toggle(
    enabled: bool = Query(True),
    yolo: bool = Query(False),
    user_id: str = Query("")
):
    """切换推测引擎"""
    engine = _get_speculation_engine(user_id)
    if enabled:
        engine.enable(yolo_mode=yolo)
    else:
        engine.disable()
    return {"ok": True, "enabled": enabled, "yolo_mode": yolo}


# ── 重试机制 API ─────────────────────────────────────

@app.get("/engine/v2/retry/status")
def retry_status(
    task_id: str = Query(""),
    user_id: str = Query("")
):
    """查询重试任务状态"""
    engine = _get_retry_engine(user_id)
    if task_id:
        meta = engine.get_status(task_id)
        if not meta:
            return {"ok": False, "error": "Task not found (may have completed)"}
        return {"ok": True, "task": meta.to_dict()}
    return {"ok": True, **engine.summary()}


@app.get("/engine/v2/retry/list")
def retry_list(
    status: str = Query(""),
    user_id: str = Query("")
):
    """列出重试任务"""
    engine = _get_retry_engine(user_id)
    if status:
        try:
            s = RetryStatus(status)
            tasks = engine.list_tasks(s)
        except ValueError:
            tasks = engine.list_active()
    else:
        tasks = engine.list_active()
    return {"ok": True, "tasks": [t.to_dict() for t in tasks], "count": len(tasks)}


@app.get("/engine/v2/retry/config")
def retry_config(
    max_attempts: int = Query(5),
    backoff_base_ms: int = Query(500),
    user_id: str = Query("")
):
    """配置重试参数"""
    engine = _get_retry_engine(user_id)
    engine.max_attempts = max(1, min(int(max_attempts), 10))
    engine._default_backoff_base_ms = max(50, min(int(backoff_base_ms), 60000))
    return {"ok": True, "max_attempts": engine.max_attempts, "backoff_base_ms": engine._default_backoff_base_ms}


# ── 工具注册表 API ───────────────────────────────────

@app.get("/engine/v2/tools/list")
def tools_list(
    capability: str = Query(""),
    approval: str = Query(""),
    tag: str = Query(""),
    role: str = Query(""),
    user_id: str = Query("")
):
    """列出工具（支持按能力/审批要求/标签/角色过滤）"""
    if role:
        tools = tool_registry.to_openai_tools(role=role)
        return {"ok": True, "tools": tools, "count": len(tools), "format": "openai"}

    filters = {}
    if capability:
        try:
            filters["capabilities"] = [Capability[capability]]
        except KeyError:
            pass
    if approval:
        try:
            filters["approval"] = ApprovalKind(approval)
        except ValueError:
            pass
    if tag:
        tools = tool_registry.list_by_tag(tag)
    elif filters:
        tools = tool_registry.filter(**filters)
    else:
        tools = tool_registry.list_enabled()
    return {"ok": True, "tools": [t.to_dict() for t in tools], "count": len(tools)}


@app.get("/engine/v2/tools/openai")
def tools_openai(
    role: str = Query(""),
    user_id: str = Query("")
):
    """导出工具为 OpenAI tool format"""
    return {"ok": True, "tools": tool_registry.to_openai_tools(role=role)}


@app.get("/engine/v2/tools/summary")
def tools_summary(user_id: str = Query("")):
    """工具注册表摘要"""
    return {"ok": True, **tool_registry.summary()}


# ── 事件帧 API ───────────────────────────────────────

_session_flows: dict = {}


@app.get("/engine/v2/events/create")
def events_create(
    session_id: str = Query(""),
    user_id: str = Query("")
):
    """创建新的事件流会话"""
    builder = EventFlowBuilder(session_id=session_id)
    _session_flows[builder.session_id] = builder
    return {"ok": True, "session_id": builder.session_id}


@app.get("/engine/v2/events/emit")
def events_emit(
    event_type: str = Query(...),
    data: str = Query("{}"),
    session_id: str = Query(""),
    user_id: str = Query("")
):
    """发送一个事件帧"""
    try:
        etype = EventType(event_type)
        parsed = json.loads(data)
    except (ValueError, json.JSONDecodeError) as e:
        return {"ok": False, "error": str(e)}

    builder = _session_flows.get(session_id)
    if not builder:
        builder = EventFlowBuilder(session_id=session_id)
        _session_flows[session_id] = builder

    frame = builder.emit(etype, parsed)
    event_log.record(frame)
    return {"ok": True, "event_id": frame.event_id, "sequence": frame.sequence}


@app.get("/engine/v2/events/stream")
def events_stream(
    session_id: str = Query(""),
    user_id: str = Query("")
):
    """获取事件流（JSON Lines）"""
    builder = _session_flows.get(session_id)
    if not builder:
        return {"ok": False, "error": "Session not found"}
    return {"ok": True, "events": builder.to_events_list(), "summary": builder.summary()}


@app.get("/engine/v2/events/log")
def events_log(
    event_type: str = Query(""),
    session_id: str = Query(""),
    limit: int = Query(50),
    user_id: str = Query("")
):
    """查询事件日志"""
    etype = EventType(event_type) if event_type else None
    results = event_log.query(event_type=etype, session_id=session_id, limit=limit)
    return {"ok": True, "events": [e.to_dict() for e in results], "count": len(results)}


# ==================== 启动时恢复Cron + 修复Stuck代理 ====================
@app.on_event("startup")
async def startup():
    # 恢复全局 cron (无user_id)
    jobs = cron_store.get()
    for name, job in jobs.items():
        if job.get("enabled"):
            _cron_start(name, "", get_ns)
            print(f"[引擎] Cron 已恢复(全局): {name}")
    # 恢复各用户的 cron
    for f in ENGINE_DIR.glob("user_*_cron.json"):
        try:
            uid = f.stem.split("_", 1)[1].rsplit("_", 1)[0]
            user_jobs = json.loads(f.read_text(encoding="utf8"))
            for name, job in user_jobs.items():
                if job.get("enabled"):
                    _cron_start(name, uid, get_ns)
                    print(f"[引擎] Cron 已恢复(用户{uid}): {name}")
        except Exception:
            pass

    # ★ 修复引擎重启后遗留的 "running" 状态子代理
    from pathlib import Path as _Path
    for f in ENGINE_DIR.glob("*_agents.json"):
        try:
            data = json.loads(f.read_text(encoding="utf8"))
            changed = False
            for name, agent in data.items():
                if agent.get("status") == "running":
                    agent["status"] = "failed"
                    agent["error"] = "引擎重启,正在运行的子代理已终止"
                    changed = True
                    print(f"[引擎] 修复stuck代理: {f.stem}/{name} (running→failed)")
            if changed:
                f.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf8")
        except Exception:
            pass

    # ★ 启动定期清理任务(每5分钟检查stuck代理)
    def _periodic_cleanup():
        _compact_tick = 0
        while True:
            time.sleep(300)  # 5分钟
            _compact_tick += 1
            try:
                # ★ 每 30 分钟压缩/清理一次流快照：30 天过期删除 + 大文件压缩。
                #   此前只在引擎启动时跑一次，运行中 STREAM_DIR 会持续膨胀。
                if _compact_tick % 6 == 1:
                    _mres = compact_stream_files(STREAM_DIR, retention_days=30, max_chunks=2048)
                    # 上传文件按分层保留：shared交付物30天，用户附件90天。
                    # 不扫描聊天内容做激进删除；调用方可传入引用路径作为保护清单。
                    try:
                        _upload_roots = [PROJECT_ROOT / 'uploads']
                        # 生产部署的历史附件目录可能位于项目同级 /var/www/html/uploads。
                        # 只加入真实存在的目录，避免维护线程因目录布局差异报错。
                        _legacy_upload_root = PROJECT_ROOT.parent / 'uploads'
                        if _legacy_upload_root.is_dir() and _legacy_upload_root not in _upload_roots:
                            _upload_roots.append(_legacy_upload_root)
                        for _upload_root in _upload_roots:
                            _ures = cleanup_uploads(_upload_root, user_retention_days=90, shared_retention_days=30)
                            if _ures.get('deleted') or _ures.get('errors'):
                                print(f"[Engine] upload maintenance ({_upload_root}): {_ures}")
                    except Exception as _uex:
                        print(f"[Engine] upload maintenance error: {_uex}")
                    if _mres.get('deleted') or _mres.get('compacted') or _mres.get('errors'):
                        print(f"[Engine] stream maintenance: {_mres}")
                    # ★ 运行时事件/任务/子代理 prune：此前从未被调用，agent_runtime.db 会无上限增长。
                    try:
                        _pres = agent_runtime.prune(retention_days=90)
                        if any(_pres.values()):
                            print(f"[Engine] runtime prune: {_pres}")
                    except Exception as _pex:
                        print(f"[Engine] runtime prune error: {_pex}")
            except Exception as _mex:
                print(f"[Engine] stream maintenance error: {_mex}")
            try:
                now = time.time()
                for f in ENGINE_DIR.glob("*_agents.json"):
                    try:
                        data = json.loads(f.read_text(encoding="utf8"))
                        changed = False
                        for name, agent in list(data.items()):
                            started = agent.get("_started_at", 0)
                            status = agent.get("status", "")
                            # running超30分钟 → failed
                            if status == "running" and started and (now - started) > 1800:
                                agent["status"] = "failed"
                                agent["error"] = "子代理执行超时(超过30分钟)"
                                changed = True
                                print(f"[引擎] 超时清理: {f.stem}/{name}")
                            # completed/failed超24小时 → 删除
                            if status in ("completed", "failed"):
                                created_str = agent.get("created", "")
                                if not created_str:
                                    continue
                                try:
                                    created = datetime.fromisoformat(created_str).timestamp()
                                    if (now - created) > 86400:  # 24小时
                                        if "result" in agent and len(agent.get("result", "")) > 10000:
                                            # 结果很大的,只保留摘要
                                            agent["result"] = agent["result"][:500] + f"\n\n[自动截断: 原结果共{len(agent['result'])}字符]"
                                            changed = True
                                        else:
                                            del data[name]
                                            changed = True
                                            print(f"[引擎] 自动删除过期代理: {f.stem}/{name}")
                                except Exception:
                                    pass
                        if changed:
                            f.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf8")
                    except Exception:
                        pass
            except Exception:
                pass

    t = threading.Thread(target=_periodic_cleanup, name="periodic_cleanup", daemon=True)
    t.start()
    print("[引擎] 定期清理线程已启动(每5分钟)")

# ==================== 前端心跳注入 ====================
_heartbeat_html = """
<script>
// OneAPIChat Engine 心跳
(function(){
    var ENGINE_URL = window.location.origin + '/oneapichat/';
    var HEARTBEAT_INTERVAL = 15000; // 15秒
    var CUSTOM_PROMPT = '';
    var token = (typeof localStorage !== 'undefined') ? localStorage.getItem('authToken') : '';

    setInterval(function(){
        var token = (typeof localStorage !== 'undefined') ? localStorage.getItem('authToken') : '';
        var authSuffix = token ? '&auth_token=' + encodeURIComponent(token) : '';
        fetch(ENGINE_URL + 'engine_api.php?action=heartbeat' + authSuffix)
            .then(function(r){ return r.json(); })
            .then(function(d){
                if(d.pending && d.pending.length > 0){
                    for(var i=0; i<d.pending.length; i++){
                        var msg = d.pending[i].msg || d.pending[i];
                        // 插入为system消息
                        if(window.chatHistory && window.currentChatId){
                            window.chatHistory[window.currentChatId].push({
                                role: 'system',
                                content: '【引擎通知】' + msg
                            });
                        }
                    }
                }
            })
            .catch(function(){});
    }, HEARTBEAT_INTERVAL);
})();
</script>
"""

# ==================== 启动 ====================
# ==================== 流式聊天后端 (SSE) ====================
import threading
import queue

def _normalize_openai_tool_turns(messages, source=''):
    """Repair a persisted transcript before any OpenAI-compatible provider call.

    Providers require assistant.tool_calls followed immediately by matching tool
    messages. Refreshes and interrupted tool handoffs can leave only part of a
    turn, or leave tool messages detached. Keep only complete adjacent pairs in
    the wire copy so a stale browser transcript cannot produce a 400.
    """
    if not isinstance(messages, list):
        return []
    normalized = []
    repaired = 0
    i = 0
    while i < len(messages):
        message = messages[i]
        if not isinstance(message, dict):
            i += 1
            continue
        if message.get('role') != 'assistant' or not isinstance(message.get('tool_calls'), list) or not message.get('tool_calls'):
            # A detached tool result is invalid outside an assistant tool turn.
            if message.get('role') != 'tool':
                normalized.append(message)
            i += 1
            continue
        calls = []
        seen_ids = set()
        for call in message.get('tool_calls') or []:
            if not isinstance(call, dict):
                continue
            call_id = str(call.get('id') or '').strip()
            if not call_id or call_id in seen_ids:
                continue
            seen_ids.add(call_id)
            fn = call.get('function')
            if isinstance(fn, dict) and not isinstance(fn.get('arguments'), str):
                fn['arguments'] = json.dumps(fn.get('arguments') or {}, ensure_ascii=False)
            calls.append(call)
        j = i + 1
        results = {}
        while j < len(messages) and isinstance(messages[j], dict) and messages[j].get('role') == 'tool':
            result = messages[j]
            result_id = str(result.get('tool_call_id') or '').strip()
            if result_id and result_id not in results:
                result['content'] = result.get('content') if isinstance(result.get('content'), str) else json.dumps(result.get('content') or '', ensure_ascii=False)
                results[result_id] = result
            j += 1
        valid_calls = [call for call in calls if str(call.get('id')) in results]
        repaired += len(calls) - len(valid_calls)
        if valid_calls:
            message['tool_calls'] = valid_calls
            normalized.append(message)
            for call in valid_calls:
                normalized.append(results[str(call.get('id'))])
        else:
            message.pop('tool_calls', None)
            normalized.append(message)
            repaired += len(calls)
        i = j
    # ★ 防御: API (尤其 Gemini) 严格禁止请求以 assistant (model) 轮次结尾
    while normalized and normalized[-1].get('role') == 'assistant':
        last_norm = normalized[-1]
        c = str(last_norm.get('content') or '').strip()
        if not c or c == '(empty)':
            normalized.pop()
            repaired += 1
        else:
            normalized.append({'role': 'user', 'content': '请基于上述内容继续。'})
            break
    messages[:] = normalized
    if repaired and source:
        print(f'[{source}] normalized {repaired} incomplete tool calls', flush=True)
    return messages


def _tool_calls_are_complete(tool_call_map: dict) -> bool:
    """Return whether accumulated tool-call deltas form executable calls.

    Some gateways emit ``finish_reason=tool_calls`` before the final argument
    delta.  The resumable reader uses this predicate before deciding it is safe
    to drain/close the provider stream.
    """
    if not isinstance(tool_call_map, dict) or not tool_call_map:
        return False
    for call in tool_call_map.values():
        if not isinstance(call, dict):
            return False
        function = call.get('function') or {}
        if not str(function.get('name') or '').strip():
            return False
        raw_args = function.get('arguments')
        if not isinstance(raw_args, str) or not raw_args.strip():
            return False
        try:
            parsed = json.loads(raw_args)
        except (TypeError, ValueError):
            return False
        if not isinstance(parsed, dict):
            return False
    return True


def _stream_openai_to_sse(request_data: dict, chat_id: str, msg_id: str, user_id: str):
    """在后台线程中将 OpenAI 流式响应转为 SSE，实时保存进度到 SQLite"""
    from openai import OpenAI
    store = get_chat_store(user_id)
    store.init_progress(msg_id, chat_id, request_data.get('model', ''))
    full_text = ''
    reasoning_text = ''
    tool_calls = []
    usage = None
    error = ''
    stop_reason = ''
    seq = 0

    def sse_event(data_str: str, event_type: str = 'chunk'):
        return f"event: {event_type}\ndata: {data_str}\n\n"

    try:
        # ★ 代理配置: 请求级优先 → 全局回退
        _httpc = None
        _rp = request_data.get('proxy_url', '')
        if request_data.get('proxy_enabled') and _rp:
            import httpx; _httpc = httpx.Client(proxy=_rp)
        elif _requires_auto_proxy(request_data.get('base_url', '')):
            import httpx; _httpc = httpx.Client(proxy=_PROXY_URL)
        is_longcat = _is_longcat_request(request_data)
        client = OpenAI(api_key=request_data.get('api_key', ''),
                        base_url=request_data.get('base_url', '').strip().rstrip('/') or None,
                        http_client=_httpc,
                        timeout=600.0 if is_longcat else 300.0)
        model = 'LongCat-2.0' if is_longcat else request_data.get('model', 'deepseek-chat')
        messages = request_data.get('messages', [])
        if is_longcat:
            _sanitize_longcat_openai_messages(messages)
        # ★ 去重 tool_call_id: MiniMax/DeepSeek 拒绝同一请求中重复的 tool_call id
        seen_tc_ids_s2 = set()
        for m in messages:
            if isinstance(m, dict) and m.get('role') == 'assistant' and 'tool_calls' in m:
                unique_calls = []
                for tc in m['tool_calls']:
                    tc_id = tc.get('id', '') if isinstance(tc, dict) else ''
                    if tc_id and tc_id not in seen_tc_ids_s2:
                        seen_tc_ids_s2.add(tc_id)
                        unique_calls.append(tc)
                if unique_calls:
                    m['tool_calls'] = unique_calls
                else:
                    del m['tool_calls']
        _dup_tool_msgs_s2 = 0
        _seen_tool_ids_s2 = set()
        for m in messages:
            if isinstance(m, dict) and m.get('role') == 'tool' and m.get('tool_call_id'):
                _tid = m['tool_call_id']
                if _tid in _seen_tool_ids_s2:
                    m['_remove'] = True
                    _dup_tool_msgs_s2 += 1
                else:
                    _seen_tool_ids_s2.add(_tid)
        if _dup_tool_msgs_s2 > 0:
            messages[:] = [m for m in messages if not m.get('_remove')]
        # ★ 清理空 tool_calls:[] 数组 — DeepSeek API 拒绝 empty array
        for m in messages:
            if isinstance(m, dict) and m.get('role') == 'assistant' and 'tool_calls' in m:
                if not m['tool_calls'] or len(m['tool_calls']) == 0:
                    del m['tool_calls']
        _normalize_openai_tool_turns(messages, '_stream_openai_to_sse')
        tools = request_data.get('tools', None)
        stream_params = {'model': model, 'messages': messages, 'stream': True}
        if tools:
            stream_params['tools'] = tools
        if request_data.get('reasoning'):
            stream_params['reasoning'] = request_data.get('reasoning')
        if is_longcat:
            stream_params['extra_body'] = {'thinking': _longcat_thinking(request_data)}
        else:
            # ★ 转发客户端 extra_body (Gemini thinking_level/include_thoughts 等),
            #   但跳过 LongCat/MiniMax 风格的 thinking key,避免顶到不支持的上游
            _eb = request_data.get('extra_body')
            if isinstance(_eb, dict) and _eb and 'thinking' not in _eb:
                stream_params['extra_body'] = _eb
        if request_data.get('reasoning_effort'):
            stream_params['reasoning_effort'] = request_data['reasoning_effort']
        if request_data.get('thinking_level'):
            stream_params['thinking_level'] = request_data['thinking_level']
        # 发送初始事件
        yield sse_event(json.dumps({'type': 'start', 'msg_id': msg_id}))

        stream = client.chat.completions.create(**stream_params)
        # ★ MiniMax/DeepSeek 内联思考标签提取状态机（含跨chunk边界保护）
        _think_buf_s2 = ''
        _in_think_s2 = False
        _chunk_carry_s2 = ''
        _TAG_MAX2 = 10
        _TAGS2 = ('(think)', '(endthink)', '<think>', '</think>')
        for chunk in stream:
            # ★ 防御: 部分提供商(LongCat等)会返回空choices块(仅usage),直接取[0]会 IndexError
            if not chunk.choices:
                if hasattr(chunk, 'usage') and chunk.usage:
                    try: usage = chunk.usage.model_dump()
                    except Exception: pass
                continue
            choice = chunk.choices[0]
            if choice.finish_reason:
                stop_reason = choice.finish_reason
            delta = choice.delta
            seq += 1

            content_delta = delta.content or ''
            if content_delta:
                # ★ 拼接上一chunk尾部缓冲以检测跨边界标签
                _c2 = _chunk_carry_s2 + content_delta
                _chunk_carry_s2 = ''
                _display_c = ''
                _pos = 0
                while _pos < len(_c2):
                    if _in_think_s2:
                        _end_idx = _c2.find('(endthink)', _pos)
                        if _end_idx != -1:
                            _think_buf_s2 += _c2[_pos:_end_idx]
                            reasoning_text += _think_buf_s2
                            store.write_chunk(msg_id, 'reasoning', _think_buf_s2)
                            yield sse_event(json.dumps({'type': 'reasoning', 'delta': _think_buf_s2, 'seq': seq}))
                            _think_buf_s2 = ''
                            _in_think_s2 = False
                            _pos = _end_idx + 10
                        else:
                            _think_buf_s2 += _c2[_pos:]
                            reasoning_text += _c2[_pos:]
                            store.write_chunk(msg_id, 'reasoning', _c2[_pos:])
                            yield sse_event(json.dumps({'type': 'reasoning', 'delta': _c2[_pos:], 'seq': seq}))
                            _pos = len(_c2)
                    else:
                        _start_idx = _c2.find('(think)', _pos)
                        if _start_idx != -1:
                            _display_c += _c2[_pos:_start_idx]
                            _pos = _start_idx + 7
                            _in_think_s2 = True
                        else:
                            _display_c += _c2[_pos:]
                            _pos = len(_c2)
                # ★ 跨chunk边界保护: 尾部可能是不完整标签前缀
                if _display_c and not _in_think_s2:
                    _tail_start2 = max(0, len(_display_c) - _TAG_MAX2)
                    for _ti2 in range(_tail_start2, len(_display_c)):
                        _tail2 = _display_c[_ti2:]
                        if any(tag.startswith(_tail2) for tag in _TAGS2):
                            _chunk_carry_s2 = _tail2
                            _display_c = _display_c[:_ti2]
                            break
                if _display_c:
                    full_text += _display_c
                    store.write_chunk(msg_id, 'content', _display_c)
                    yield sse_event(json.dumps({'type': 'content', 'delta': _display_c, 'seq': seq}))

            # 思考增量（reasoning_content 字段 — DeepSeek 等）
            reasoning_delta = delta.reasoning_content or ''
            if reasoning_delta:
                reasoning_text += reasoning_delta
                store.write_chunk(msg_id, 'reasoning', reasoning_delta)
                yield sse_event(json.dumps({'type': 'reasoning', 'delta': reasoning_delta, 'seq': seq}))

            # Tool calls
            if delta.tool_calls:
                for tc in delta.tool_calls:
                    tc_dict = {'id': tc.id, 'type': tc.type,
                                'function': {'name': tc.function.name,
                                             'arguments': tc.function.arguments or ''}}
                    # ★ 保留 Gemini thought_signature (Google API 要求回传,否则 400)
                    if hasattr(tc, 'thought_signature') and tc.thought_signature:
                        tc_dict['thought_signature'] = tc.thought_signature
                    if hasattr(tc.function, 'thought_signature') and tc.function.thought_signature:
                        tc_dict['function']['thought_signature'] = tc.function.thought_signature
                    tool_calls.append(tc_dict)
                    yield sse_event(json.dumps({'type': 'tool_call', 'delta': tc_dict, 'seq': seq}))

            # Usage
            if chunk.usage:
                try:
                    usage = chunk.usage.model_dump()
                except Exception:
                    try:
                        usage = json.loads(chunk.usage.model_dump_json())
                    except Exception:
                        usage = dict(chunk.usage)

        # 流结束 — 刷新跨chunk缓冲 + 安全网清理内联标签
        if _chunk_carry_s2:
            full_text += _chunk_carry_s2
        if _in_think_s2 and _think_buf_s2:
            reasoning_text += _think_buf_s2
        import re as _re_tmp2
        _inlines2 = _re_tmp2.findall(r'\(think\)([\s\S]*?)\(endthink\)', full_text)
        if _inlines2:
            reasoning_text += ''.join(_inlines2)
            full_text = _re_tmp2.sub(r'', full_text)
        _open_m2 = _re_tmp2.search(r'\(think\)([\s\S]*?)$', full_text)
        if _open_m2 and len(_open_m2.group(1)) < 3000:
            reasoning_text += _open_m2.group(1)
            full_text = _re_tmp2.sub(r'', full_text)
        store.finish_stream(msg_id, full_text.strip(), reasoning_text.strip(), tool_calls, usage)
        yield sse_event(json.dumps({'type': 'done', 'full_text': full_text.strip(), 'reasoning_text': reasoning_text.strip(),
                                     'tool_calls': tool_calls, 'usage': usage,
                                     'stop_reason': stop_reason,
                                     'truncated': stop_reason in ('max_tokens', 'length')}))

    except Exception as e:
        error = str(e)
        print(f"[stream] error: {error}")
        store.finish_stream(msg_id, full_text, reasoning_text, tool_calls, usage, error)
        yield sse_event(json.dumps({'type': 'error', 'error': error}))

def _run_stream(request_data: dict, chat_id: str, msg_id: str, user_id: str, result_queue):
    """后台线程运行器,逐块转发SSE事件,不缓存"""
    try:
        for chunk in _stream_openai_to_sse(request_data, chat_id, msg_id, user_id):
            result_queue.put(('chunk', chunk))
        result_queue.put(('done', None))
    except Exception as e:
        result_queue.put(('error', str(e)))

@app.post("/engine/chat/stream")
async def chat_stream(request: Request, user_id: str = Query("")):
    """
    后端流式聊天端点:
    - 接收消息，转发给 OpenAI，流式返回 SSE
    - 实时将进度保存到 SQLite（刷新恢复）
    """
    try:
        body = await request.json()
    except Exception:
        return {"error": "invalid JSON body"}

    chat_id = body.get('chat_id') or ''
    msg_id = body.get('msg_id') or f"msg_{int(time.time()*1000)}"
    request_data = body.get('request', {})

    if not request_data.get('api_key'):
        return {"error": "api_key required"}

    # 启动后台线程执行流式请求（避免 FastAPI 线程阻塞）
    result_queue = queue.Queue()
    t = threading.Thread(target=_run_stream, args=(request_data, chat_id, msg_id, user_id, result_queue), daemon=True)
    t.start()

    async def event_generator():
        # 前端通过 EventSource 接收 SSE
        while True:
            try:
                status, data = result_queue.get(timeout=60)
                if status == 'error':
                    yield f"event: error\ndata: {json.dumps({'error': data})}\n\n"
                    break
                elif status == 'done':
                    break
                elif status == 'chunk':
                    yield data
                    await asyncio.sleep(0.001)
            except queue.Empty:
                yield f"event: timeout\ndata: {json.dumps({'error': 'stream timeout'})}\n\n"
                break

    return StreamingResponse(event_generator(), media_type="text/event-stream")

@app.get("/engine/chat/progress/{msg_id}")
async def chat_progress(msg_id: str, user_id: str = Query("")):
    """查询流式进度（用于刷新恢复）"""
    store = get_chat_store(user_id)
    return store.get_progress(msg_id)


# ═══════════════════════════════════════════════════════════════
# 可恢复流式系统 — 流在后端线程独立运行，前端断线重连无感续接
# ═══════════════════════════════════════════════════════════════

_resumable = {}           # {stream_id: {queue, chunks, finished, created}}
_resumable_lock = threading.Lock()
_RESUMABLE_TTL = 1800      # 30 分钟超时


def _is_cancelled(stream_id: str) -> bool:
    """检查流是否被用户取消(停止键)"""
    with _resumable_lock:
        entry = _resumable.get(stream_id)
        return bool(entry and entry.get('cancel'))


def _is_longcat_request(request_data: dict) -> bool:
    """LongCat 识别同时覆盖官方 base URL 和模型名。"""
    model = str(request_data.get('model', '') or '').lower()
    base_url = str(request_data.get('base_url', '') or '').lower()
    anthropic_url = str(request_data.get('anthropic_url', '') or '').lower()
    return 'longcat' in model or 'api.longcat.chat' in base_url or 'api.longcat.chat' in anthropic_url


def _longcat_thinking(request_data: dict) -> dict:
    """官方只接受 enabled/disabled；未配置时默认关闭以保证普通问答有正文。"""
    thinking = request_data.get('thinking')
    mode = thinking.get('type') if isinstance(thinking, dict) else ''
    return {'type': 'enabled' if mode == 'enabled' else 'disabled'}


def _longcat_text_content(content):
    """把 OpenAI 多模态块降级为 LongCat 当前支持的纯文本输入。"""
    if not isinstance(content, list):
        return content
    parts = []
    for block in content:
        if isinstance(block, str):
            parts.append(block)
        elif isinstance(block, dict):
            block_type = block.get('type', '')
            if block_type == 'text':
                parts.append(str(block.get('text', '') or ''))
            elif block_type in ('image', 'image_url'):
                parts.append('[图片]')
            elif block_type in ('video', 'video_url'):
                parts.append('[视频]')
            elif block_type == 'tool_result':
                value = block.get('content', '')
                parts.append(value if isinstance(value, str) else json.dumps(value, ensure_ascii=False))
    return '\n'.join(part for part in parts if part)


def _sanitize_longcat_openai_messages(messages: list) -> list:
    """移除 LongCat OpenAI 端点不支持的历史推理元数据。"""
    for message in messages:
        if not isinstance(message, dict):
            continue
        if isinstance(message.get('content'), list):
            message['content'] = _longcat_text_content(message['content'])
        for key in ('reasoning_content', 'reasoning', 'reasoning_details'):
            message.pop(key, None)
    return messages


# ═══════════════════════════════════════════════════════════════
# StreamBuffer — 磁盘持久化流缓冲（引擎重启不丢 chunks）
# ═══════════════════════════════════════════════════════════════

MAX_STREAM_REPLAY_CHUNKS = 2048
MAX_STREAM_REPLAY_BYTES = 4 * 1024 * 1024


def _trim_stream_replay(chunks: list) -> tuple[list, int, int]:
    values = list(chunks or [])
    sizes = [len(item.encode('utf-8', errors='ignore')) if isinstance(item, str) else len(json.dumps(item, ensure_ascii=False).encode('utf-8')) for item in values]
    total = sum(sizes)
    dropped = 0
    while values and (len(values) > MAX_STREAM_REPLAY_CHUNKS or total > MAX_STREAM_REPLAY_BYTES):
        values.pop(0)
        total -= sizes.pop(0)
        dropped += 1
    return values, dropped, max(0, total)


class StreamBuffer:
    """msg_id 粒度的流缓冲，并持久化可立即恢复的聚合快照。"""
    __slots__ = (
        'msg_id', 'path', 'chunks', 'content', 'reasoning', 'tool_calls',
        'usage', 'finished', 'error', 'stop_reason', 'truncated',
        'stream_id', 'chat_id', 'user_id', 'chunk_base_offset', 'total_events', 'replay_bytes',
        '_last_save', '_lock'
    )

    def __init__(self, msg_id: str):
        # 严格校验 msg_id 防止路径遍历或异常字符
        safe_msg_id = re.sub(r'[^a-zA-Z0-9_-]', '', str(msg_id or ''))[:128]
        if not safe_msg_id:
            safe_msg_id = f"msg_{int(time.time()*1000)}"
        self.msg_id = safe_msg_id
        self.path = (STREAM_DIR / f"{safe_msg_id}.json").resolve()
        # 确保路径严格位于 STREAM_DIR 内部
        try:
            is_inside = os.path.commonpath((str(self.path), str(STREAM_DIR.resolve()))) == str(STREAM_DIR.resolve())
        except (OSError, ValueError):
            is_inside = False
        if not is_inside:
            self.path = (STREAM_DIR / f"msg_{uuid.uuid4().hex[:12]}.json").resolve()
        self.chunks: list = []
        self.content: str = ''
        self.reasoning: str = ''
        self.tool_calls: list = []
        self.usage = None
        self.finished: bool = False
        self.error: str = ''
        self.stop_reason: str = ''
        self.truncated: bool = False
        self.stream_id: str = ''
        self.chat_id: str = ''
        self.user_id: str = ''
        self.chunk_base_offset: int = 0
        self.total_events: int = 0
        self.replay_bytes: int = 0
        self._last_save = 0.0
        self._lock = threading.RLock()
        self._load()

    def _load(self):
        if self.path.exists():
            try:
                data = json.loads(self.path.read_text(encoding='utf-8'))
                loaded_chunks = data.get("chunks", []) or []
                self.chunk_base_offset = int(data.get("chunk_base_offset", 0) or 0)
                self.chunks = loaded_chunks
                self.total_events = max(int(data.get("event_count", 0) or 0), self.chunk_base_offset + len(self.chunks))
                self.content = data.get("content", "")
                self.reasoning = data.get("reasoning", data.get("reasoning_text", ""))
                self.tool_calls = data.get("tool_calls", []) or []
                self.usage = data.get("usage")
                self.finished = bool(data.get("finished", False))
                self.error = data.get("error", "") or ""
                self.stop_reason = data.get("stop_reason", "") or ""
                self.truncated = bool(data.get("truncated", False))
                self.stream_id = data.get("stream_id", "") or ""
                self.chat_id = data.get("chat_id", "") or ""
                self.user_id = data.get("user_id", "") or ""
                self._last_save = float(data.get("ts", 0) or 0)
                # 升级前的缓冲文件只保存 chunks/finished。首次读取时从 SSE 历史
                # 重建聚合快照，避免部署升级后的第一次刷新显示为空白。
                if (not data.get("snapshot_version") and self.chunks and
                        not any(key in data for key in (
                            "content", "reasoning", "tool_calls", "usage", "error", "stop_reason"
                        ))):
                    self._rebuild_snapshot_from_chunks()
                self.chunks, dropped, self.replay_bytes = _trim_stream_replay(self.chunks)
                self.chunk_base_offset += dropped
            except Exception:
                pass

    def _rebuild_snapshot_from_chunks(self):
        for payload in self.chunks:
            if not isinstance(payload, str):
                continue
            event_type = ''
            event_data = None
            for raw_line in payload.splitlines():
                line = raw_line.strip()
                if line.startswith('event:'):
                    event_type = line[6:].strip()
                elif line.startswith('data:'):
                    try:
                        event_data = json.loads(line[5:].strip())
                    except Exception:
                        event_data = None
            if not isinstance(event_data, dict):
                continue
            event_type = event_type or str(event_data.get('type', '') or '')
            if event_type == 'content':
                self.content += str(event_data.get('delta', '') or '')
            elif event_type == 'reasoning':
                self.reasoning += str(event_data.get('delta', '') or '')
            elif event_type == 'tool_call':
                tools = event_data.get('tools')
                if isinstance(tools, list):
                    self.tool_calls = tools
                elif event_data.get('function'):
                    self.tool_calls.append(event_data)
            elif event_type == 'done' or 'full_text' in event_data:
                self.content = event_data.get('full_text', self.content) or self.content
                self.reasoning = event_data.get('reasoning_text', self.reasoning) or self.reasoning
                if isinstance(event_data.get('tool_calls'), list):
                    self.tool_calls = event_data['tool_calls']
                self.usage = event_data.get('usage', self.usage)
                self.stop_reason = event_data.get('stop_reason', self.stop_reason) or self.stop_reason
                self.truncated = bool(event_data.get('truncated', self.truncated))
                self.finished = True
                self.error = ''
            elif event_type == 'error' or event_data.get('error'):
                self.error = str(event_data.get('error', 'stream error') or 'stream error')
                self.finished = True

    def _save(self):
        with self._lock:
            try:
                now = time.time()
                payload = {
                    "snapshot_version": 1,
                    "chunks": self.chunks,
                    "chunk_base_offset": self.chunk_base_offset,
                    "event_count": self.total_events,
                    "replay_bytes": self.replay_bytes,
                    "content": self.content,
                    "reasoning": self.reasoning,
                    "tool_calls": self.tool_calls,
                    "usage": self.usage,
                    "finished": self.finished,
                    "error": self.error,
                    "stop_reason": self.stop_reason,
                    "truncated": self.truncated,
                    "stream_id": self.stream_id,
                    "chat_id": self.chat_id,
                    "user_id": self.user_id,
                    "ts": now,
                }
                # 原子替换，避免进程退出或并发读取时留下半截 JSON。
                tmp_path = self.path.with_suffix(self.path.suffix + '.tmp')
                tmp_path.write_text(json.dumps(payload, ensure_ascii=False), encoding='utf-8')
                os.replace(tmp_path, self.path)
                self._last_save = now
            except Exception as e:
                print(f"[StreamBuffer] save error={type(e).__name__}")

    def set_meta(self, stream_id: str = '', chat_id: str = '', user_id: str = ''):
        with self._lock:
            if stream_id:
                self.stream_id = stream_id
            if chat_id:
                self.chat_id = chat_id
            if user_id:
                self.user_id = user_id

    def _append_replay(self, sse_payload: str):
        self.chunks.append(sse_payload)
        self.total_events += 1
        self.replay_bytes += len(sse_payload.encode('utf-8', errors='ignore'))
        drop_count = max(0, len(self.chunks) - MAX_STREAM_REPLAY_CHUNKS)
        dropped_bytes = 0
        for index, payload in enumerate(self.chunks):
            if index < drop_count or self.replay_bytes - dropped_bytes > MAX_STREAM_REPLAY_BYTES:
                dropped_bytes += len(payload.encode('utf-8', errors='ignore')) if isinstance(payload, str) else len(json.dumps(payload, ensure_ascii=False).encode('utf-8'))
                drop_count = max(drop_count, index + 1)
            else:
                break
        if drop_count:
            del self.chunks[:drop_count]
            self.chunk_base_offset += drop_count
            self.replay_bytes = max(0, self.replay_bytes - dropped_bytes)

    def record(self, event_type: str, data: dict, sse_payload: str):
        """追加原始事件并同步聚合快照；工具和终态事件立即落盘。"""
        with self._lock:
            if event_type == 'tool_call' and data.get('partial'):
                # The aggregate tool_calls snapshot below is sufficient for reconnect; storing
                # every growing partial argument snapshot causes quadratic disk amplification.
                self.total_events += 1
            else:
                self._append_replay(sse_payload)
            if event_type == 'content':
                self.content += str(data.get('delta', '') or '')
            elif event_type == 'reasoning':
                self.reasoning += str(data.get('delta', '') or '')
            elif event_type == 'tool_call':
                tools = data.get('tools')
                if isinstance(tools, list):
                    self.tool_calls = tools
                elif data.get('function'):
                    self.tool_calls.append(data)
            elif event_type == 'done':
                self.content = data.get('full_text', self.content) or self.content
                self.reasoning = data.get('reasoning_text', self.reasoning) or self.reasoning
                if isinstance(data.get('tool_calls'), list):
                    self.tool_calls = data['tool_calls']
                self.usage = data.get('usage', self.usage)
                self.stop_reason = data.get('stop_reason', self.stop_reason) or self.stop_reason
                self.truncated = bool(data.get('truncated', self.truncated))
                self.finished = True
                self.error = ''
            elif event_type == 'error':
                self.error = str(data.get('error', 'stream error') or 'stream error')
                self.finished = True

            should_save = (
                event_type in ('tool_call', 'done', 'error') or
                len(self.chunks) % 5 == 0 or
                time.time() - self._last_save >= 1.0
            )
        if should_save:
            self._save()

    def append(self, sse_payload: str):
        # 兼容旧调用；新流应使用 record()，以同时维护聚合快照。
        with self._lock:
            self._append_replay(sse_payload)
            should_save = self.total_events % 5 == 0
        if should_save:
            self._save()

    def since(self, offset: int):
        with self._lock:
            absolute_offset = max(0, int(offset))
            if absolute_offset >= self.chunk_base_offset + len(self.chunks):
                return []
            start = max(0, absolute_offset - self.chunk_base_offset)
            return list(self.chunks[start:])

    def snapshot(self):
        """返回一致的首屏状态，客户端无需等待下一个 token 才能重绘。"""
        with self._lock:
            return {
                'msg_id': self.msg_id,
                'stream_id': self.stream_id,
                'chat_id': self.chat_id,
                'full_text': self.content,
                'reasoning_text': self.reasoning,
                'tool_calls': json.loads(json.dumps(self.tool_calls, ensure_ascii=False)),
                'usage': self.usage,
                'finished': self.finished,
                'error': self.error,
                'stop_reason': self.stop_reason,
                'truncated': self.truncated,
                'offset': self.chunk_base_offset + len(self.chunks),
                'event_count': self.total_events,
                'replay_base_offset': self.chunk_base_offset,
                'replay_bytes': self.replay_bytes,
                'updated_at': self._last_save,
            }

    def done(self):
        with self._lock:
            self.finished = True
        self._save()


_stream_buffers: dict = {}  # {msg_id: StreamBuffer}
_stream_buffers_lock = threading.Lock()


def _cleanup_stream_runtime_state(max_completed_age: int = 600):
    cutoff = time.time() - max(60, int(max_completed_age))
    with _stream_buffers_lock:
        stale_buffers = [
            key for key, value in _stream_buffers.items()
            if value.finished and value._last_save and value._last_save < cutoff
        ]
        for key in stale_buffers:
            _stream_buffers.pop(key, None)
    with _resumable_lock:
        stale_streams = [
            key for key, value in _resumable.items()
            if value.get('finished') and float(value.get('finished_at') or 0) < cutoff
        ]
        for key in stale_streams:
            _resumable.pop(key, None)


def _get_stream_buffer(msg_id: str) -> StreamBuffer:
    if len(_stream_buffers) > 512 or len(_resumable) > 512:
        _cleanup_stream_runtime_state()
    with _stream_buffers_lock:
        if msg_id not in _stream_buffers:
            _stream_buffers[msg_id] = StreamBuffer(msg_id)
        return _stream_buffers[msg_id]


def _record_runtime_stream_event(stream_id: str, event_type: str, data: dict) -> dict:
    """Mirror legacy SSE events into the durable DSH-style session log."""
    with _resumable_lock:
        entry = dict(_resumable.get(stream_id, {}))
    session_id = entry.get('runtime_session_id', '')
    if not session_id:
        return data
    try:
        if event_type == 'content':
            agent_runtime.append_stream_delta(session_id, content_delta=str(data.get('delta', '') or ''))
        elif event_type == 'reasoning':
            agent_runtime.append_stream_delta(session_id, reasoning_delta=str(data.get('delta', '') or ''))
        elif event_type == 'done':
            with _resumable_lock:
                live = _resumable.get(stream_id, {})
                if live.get('runtime_terminal'):
                    return data
                live['runtime_terminal'] = True
                live['finished_at'] = time.time()
            agent_runtime.flush_stream_delta(session_id)
            provider = detect_provider(
                entry.get('base_url', '') or entry.get('anthropic_url', ''),
                entry.get('model', ''),
                anthropic_format=bool(entry.get('anthropic_format')),
            )
            normalized_usage = normalize_usage(data.get('usage'), provider)
            if normalized_usage:
                data['usage_normalized'] = normalized_usage
            for call in data.get('tool_calls') or []:
                fn = call.get('function') or {} if isinstance(call, dict) else {}
                agent_runtime.append_event(session_id, 'tool/call', {
                    'call_id': call.get('id', '') if isinstance(call, dict) else '',
                    'name': fn.get('name', ''),
                    'arguments': fn.get('arguments', ''),
                })
            agent_runtime.append_event(session_id, 'assistant/message', {
                'message_id': entry.get('msg_id', ''),
                'content': data.get('full_text', ''),
                'reasoning': data.get('reasoning_text', ''),
                'tool_calls': data.get('tool_calls') or [],
                'usage': normalized_usage or data.get('usage') or {},
                'model': entry.get('model', ''),
            })
            reason_kind = 'max-tokens' if data.get('truncated') else 'completed'
            agent_runtime.append_event(session_id, 'step/end', {'reason': {'kind': reason_kind}})
            agent_runtime.append_event(session_id, 'turn/end', {'reason': {'kind': reason_kind}})
            if entry.get('runtime_job_id'):
                agent_runtime.update_job(entry['runtime_job_id'], 'completed', result={
                    'message_id': entry.get('msg_id', ''),
                    'stop_reason': data.get('stop_reason', ''),
                    'usage': normalized_usage or data.get('usage') or {},
                }, user_id=entry.get('user_id') or None)
        elif event_type == 'error':
            with _resumable_lock:
                live = _resumable.get(stream_id, {})
                if live.get('runtime_terminal'):
                    return data
                live['runtime_terminal'] = True
                live['finished_at'] = time.time()
            agent_runtime.flush_stream_delta(session_id)
            provider = detect_provider(
                entry.get('base_url', '') or entry.get('anthropic_url', ''),
                entry.get('model', ''),
                anthropic_format=bool(entry.get('anthropic_format')),
            )
            classified = classify_provider_error(Exception(str(data.get('error') or 'provider stream failed')), provider=provider.name)
            data.update({
                'code': classified.code.value,
                'retryable': classified.retryable,
                'details': classified.details,
            })
            agent_runtime.append_event(session_id, 'session/error', {'error': classified.to_dict()})
            agent_runtime.append_event(session_id, 'step/end', {'reason': {'kind': 'error'}})
            agent_runtime.append_event(session_id, 'turn/end', {'reason': {'kind': 'error'}})
            if entry.get('runtime_job_id'):
                agent_runtime.update_job(entry['runtime_job_id'], 'failed', error=classified.to_dict(), user_id=entry.get('user_id') or None)
    except Exception as runtime_error:
        print(f"[runtime] stream event mirror failed: {runtime_error}", flush=True)
    return data


def _finalize_cancelled_stream(stream_id: str, message: str = 'generation cancelled') -> bool:
    """Persist and publish one terminal cancellation event for reconnecting clients.

    A cancelled provider request can remain blocked inside an HTTP read for a while.  The
    DELETE endpoint therefore closes the logical stream immediately, while the worker uses
    the cancel flag to stop processing/retrying once control returns from the provider.
    """
    with _resumable_lock:
        entry = _resumable.get(stream_id)
        if not entry or entry.get('cancel_notified'):
            return False
        entry['cancel'] = True
        entry['cancel_notified'] = True
        entry['finished'] = True
        msg_id = entry.get('msg_id', stream_id)
        q = entry.get('queue')

    error = AgentRuntimeError(ErrorCode.CANCELLED, message, status=499, retryable=False).to_dict()
    data = _record_runtime_stream_event(stream_id, 'error', {'error': error})
    sse = f"event: error\ndata: {json.dumps(data, ensure_ascii=False)}\n\n"

    with _resumable_lock:
        live = _resumable.get(stream_id)
        if live is not None:
            chunks = live.get('chunks')
            if isinstance(chunks, list):
                chunks.append(sse)
                if len(chunks) > MAX_STREAM_REPLAY_CHUNKS:
                    del chunks[:len(chunks) - MAX_STREAM_REPLAY_CHUNKS]

    _get_stream_buffer(msg_id).record('error', data, sse)
    if q is not None:
        try:
            q.put(sse)
            q.put(None)
        except Exception:
            pass
    _complete_task_from_stream(stream_id, 'failed')
    return True


def _generate_resumable_anthropic(request: dict, stream_id: str, _emit):
    """Anthropic Messages API 流式分支 — 用 requests 直接调 /messages?stream=True"""
    import requests as _requests
    q = _resumable[stream_id]['queue']
    full = ''
    reasoning = ''
    tool_calls = []
    usage = None
    stop_reason = ''

    anthropic_url = request.get('anthropic_url', '')
    api_key = request.get('api_key', '')
    model = request.get('model', '')
    is_longcat = _is_longcat_request(request)
    if is_longcat:
        model = 'LongCat-2.0'
    # ★ 代理配置: 请求级优先 → 全局回退
    _req_proxy = request.get('proxy_url', '')
    _proxies = None
    if request.get('proxy_enabled') and _req_proxy:
        _proxies = {'http': _req_proxy, 'https': _req_proxy}
    elif _requires_auto_proxy(anthropic_url):
        _proxies = {'http': _PROXY_URL, 'https': _PROXY_URL}

    # ★ 提取 system 消息 (Anthropic 用顶级 system 字段)
    raw_system = request.get('system', '')
    if isinstance(raw_system, str):
        system_content = raw_system
    elif isinstance(raw_system, list):
        system_content = '\n\n'.join(
            str(block.get('text', '') or '') for block in raw_system
            if isinstance(block, dict) and block.get('type') == 'text'
        )
    else:
        system_content = ''
    anthropic_messages = []
    for m in request.get('messages', []):
        if isinstance(m, dict) and m.get('role') == 'system':
            system_content += (system_content and '\n\n') + (m.get('content', '') or '')
        else:
            anthropic_messages.append(m)

    payload = {
        'model': model,
        'messages': anthropic_messages,
        'max_tokens': request.get('max_tokens') or request.get('max_completion_tokens') or 4096,
        'stream': True,
    }
    if system_content:
        payload['system'] = system_content
    if request.get('tools'):
        payload['tools'] = request['tools']
    if request.get('temperature') is not None:
        payload['temperature'] = request['temperature']
    if is_longcat:
        payload['thinking'] = _longcat_thinking(request)
    elif isinstance(request.get('thinking'), dict):
        payload['thinking'] = request['thinking']

    # ★ 认证头: 原生Anthropic用 x-api-key, 其他(LongCat/DeepSeek)用 Bearer
    if 'api.anthropic.com' in anthropic_url:
        headers = {
            'x-api-key': api_key,
            'anthropic-version': '2023-06-01',
            'content-type': 'application/json',
        }
    else:
        headers = {
            'Authorization': f'Bearer {api_key}',
            'x-api-key': api_key,
            'anthropic-version': '2023-06-01',
            'content-type': 'application/json',
        }

    _tool_use_map = {}  # index → {id, name, input_json}
    _current_tool_idx = -1

    try:
        from urllib.parse import urlparse as _urlparse
        _anthropic_host = _urlparse(anthropic_url).hostname or 'unknown'
        print(f"[_generate_resumable_anthropic] Starting stream {stream_id} model={model} host={_anthropic_host}", flush=True)
        def _open_anthropic_stream():
            _last_error = None
            for _attempt in range(5):
                try:
                    _response = _requests.post(
                        anthropic_url, json=payload, headers=headers, proxies=_proxies,
                        stream=True, timeout=(30, 600),
                    )
                except Exception as _request_exc:
                    _last_error = classify_provider_error(_request_exc, provider='anthropic')
                else:
                    if _response.status_code == 200:
                        return _response
                    _last_error = classify_provider_error(
                        Exception(f'Anthropic HTTP {_response.status_code}'), status=_response.status_code,
                        response_body=_response.text[:4000],
                        request_id=_response.headers.get('x-request-id'), provider='anthropic',
                    )
                    _response.close()
                if not _last_error.retryable or _attempt >= 4:
                    raise _last_error
                _retry_after = 0.0
                if isinstance(_last_error.details, dict):
                    try:
                        _retry_after = float(_last_error.details.get('retry_after') or 0)
                    except (TypeError, ValueError):
                        _retry_after = 0.0
                time.sleep(max(_retry_after, min(10.0, 1.5 * (2 ** _attempt))))
            raise _last_error or AgentRuntimeError(ErrorCode.PROVIDER_UNAVAILABLE, 'Anthropic unavailable', status=503)

        with _open_anthropic_stream() as resp:
            for line in resp.iter_lines():
                if _is_cancelled(stream_id):
                    print(f"[_generate_resumable_anthropic] CANCELLED stream {stream_id}", flush=True)
                    break
                if not line:
                    continue
                try:
                    line_str = line.decode('utf-8') if isinstance(line, bytes) else line
                except Exception:
                    continue
                if not line_str.startswith('data:'):
                    continue
                try:
                    d = json.loads(line_str[5:].strip())
                except Exception:
                    continue
                etype = d.get('type', '')

                if etype == 'message_start':
                    mu = d.get('message', {}).get('usage')
                    if mu:
                        usage = {'prompt_tokens': mu.get('input_tokens', 0), 'completion_tokens': 0,
                                 'total_tokens': mu.get('input_tokens', 0)}
                elif etype == 'content_block_start':
                    cb = d.get('content_block', {}) or {}
                    if cb.get('type') == 'tool_use':
                        idx = d.get('index', 0)
                        _current_tool_idx = idx
                        _tool_use_map[idx] = {'id': cb.get('id', ''), 'name': cb.get('name', ''), 'input_json': ''}
                    elif cb.get('type') == 'text':
                        pass
                elif etype == 'content_block_delta':
                    delta = d.get('delta', {}) or {}
                    if delta.get('type') == 'text_delta':
                        txt = delta.get('text', '')
                        if txt:
                            full += txt
                            _emit('content', {'delta': txt})
                    elif delta.get('type') == 'thinking_delta':
                        rtxt = delta.get('thinking', '')
                        if rtxt:
                            reasoning += rtxt
                            _emit('reasoning', {'delta': rtxt})
                    elif delta.get('type') == 'input_json_delta':
                        if _current_tool_idx >= 0 and _current_tool_idx in _tool_use_map:
                            _tool_use_map[_current_tool_idx]['input_json'] += delta.get('partial_json', '')
                            # ★ 实时推送部分tool_call快照(前端展示)
                            _snap = []
                            for _idx in sorted(_tool_use_map.keys()):
                                _tu = _tool_use_map[_idx]
                                _snap.append({'id': _tu['id'], 'type': 'function',
                                              'function': {'name': _tu['name'], 'arguments': _tu['input_json']}})
                            _emit('tool_call', {'partial': True, 'tools': _snap})
                elif etype == 'content_block_stop':
                    if _current_tool_idx >= 0 and _current_tool_idx in _tool_use_map:
                        tu = _tool_use_map[_current_tool_idx]
                        inp = {}
                        try:
                            inp = json.loads(tu['input_json']) if tu['input_json'] else {}
                        except Exception:
                            # ★ 容错: 参数JSON非法时恢复 server_exec/server_python 的命令
                            inp = _tolerant_tool_args(tu['name'], tu['input_json']) or {}
                        tool_calls.append({'id': tu['id'], 'type': 'function',
                                           'function': {'name': tu['name'], 'arguments': json.dumps(inp)}})
                        _current_tool_idx = -1
                elif etype == 'message_delta':
                    delta_info = d.get('delta', {}) or {}
                    if delta_info.get('stop_reason'):
                        stop_reason = delta_info['stop_reason']
                    mu = d.get('usage')
                    if mu:
                        if usage:
                            usage['completion_tokens'] = mu.get('output_tokens', 0)
                            usage['total_tokens'] = usage.get('prompt_tokens', 0) + mu.get('output_tokens', 0)
                        else:
                            usage = {'prompt_tokens': 0, 'completion_tokens': mu.get('output_tokens', 0),
                                     'total_tokens': mu.get('output_tokens', 0)}
                # message_stop / ping → 忽略
    except Exception as e:
        _provider_error = e if isinstance(e, AgentRuntimeError) else classify_provider_error(e, provider='anthropic')
        _resumable[stream_id]['error'] = _provider_error.to_dict()
        _emit('error', {'error': _provider_error.to_dict()})
        _resumable[stream_id]['finished'] = True
        _complete_task_from_stream(stream_id, 'failed')
        try: q.put(None)
        except Exception: pass
        return

    if _is_cancelled(stream_id):
        _finalize_cancelled_stream(stream_id)
        with _resumable_lock:
            _resumable.pop(stream_id, None)
        return

    truncated = stop_reason in ('max_tokens', 'length')
    done_data = {'full_text': full.strip(), 'reasoning_text': reasoning.strip(),
                 'tool_calls': tool_calls, 'usage': usage,
                 'stop_reason': stop_reason, 'truncated': truncated}
    _emit('done', done_data)
    _resumable[stream_id]['finished'] = True
    _complete_task_from_stream(stream_id, 'completed')
    try: q.put(None)
    except Exception: pass


def _generate_resumable(request: dict, stream_id: str):
    """后台线程: 调用 OpenAI 流式 API，逐 chunk 写入缓存和队列"""
    from openai import OpenAI
    q = _resumable[stream_id]['queue']
    full = ''
    reasoning = ''
    tool_calls = []
    usage = None
    stop_reason = ''

    # 关联磁盘缓冲（从 stream_id 提取 msg_id，或使用 stream_id 本身）
    msg_id = _resumable[stream_id].get('msg_id', stream_id)
    buf = _get_stream_buffer(msg_id)

    def _emit(ev_type, data):
        data = _record_runtime_stream_event(stream_id, ev_type, data)
        sse = f"event: {ev_type}\ndata: {json.dumps(data, ensure_ascii=False)}\n\n"
        live_chunks = _resumable[stream_id]['chunks']
        if ev_type != 'tool_call' or not data.get('partial'):
            live_chunks.append(sse)
            if len(live_chunks) > MAX_STREAM_REPLAY_CHUNKS:
                del live_chunks[:len(live_chunks) - MAX_STREAM_REPLAY_CHUNKS]
        buf.record(ev_type, data, sse)
        q.put(sse)  # queue.Queue is thread-safe

    # ★ Anthropic Messages API 流式分支 (前端已转换消息/工具为Anthropic格式)
    if request.get('anthropic_format') and request.get('anthropic_url'):
        return _generate_resumable_anthropic(request, stream_id, _emit)

    try:
        from urllib.parse import urlparse as _urlparse
        _provider_host = _urlparse(request.get('base_url', '')).hostname or 'unknown'
        print(f"[_generate_resumable] Starting stream {stream_id} with model={request.get('model','?')} host={_provider_host}", flush=True)
        # ★ 代理配置: 请求级优先 → 全局回退
        _http_client = None
        _req_proxy = request.get('proxy_url', '')
        if request.get('proxy_enabled') and _req_proxy:
            import httpx
            _http_client = httpx.Client(proxy=_req_proxy)
        elif _requires_auto_proxy(request.get('base_url', '')):
            import httpx
            _http_client = httpx.Client(proxy=_PROXY_URL)
        is_longcat = _is_longcat_request(request)
        model = 'LongCat-2.0' if is_longcat else request.get('model', 'deepseek-chat')
        client = OpenAI(
            api_key=request.get('api_key', ''),
            base_url=request.get('base_url', '').strip().rstrip('/') or None,
            http_client=_http_client,
            timeout=600.0 if is_longcat else 300.0,
            # Retry ownership belongs to _iter_provider_chunks below. Keeping the SDK
            # default (2 retries) multiplies one logical 5-attempt budget into as many
            # as 15 upstream requests and continues work after a user cancellation.
            max_retries=0,
        )
        messages = request.get('messages', [])
        if is_longcat:
            _sanitize_longcat_openai_messages(messages)
        # ★ 去重 tool_call_id: MiniMax/DeepSeek 拒绝同一请求中重复的 tool_call id
        #  检查范围: 1) assistant 消息的 tool_calls[].id  2) tool 消息的 tool_call_id
        seen_tc_ids = set()
        for m in messages:
            if isinstance(m, dict) and m.get('role') == 'assistant' and 'tool_calls' in m:
                unique_calls = []
                for tc in m['tool_calls']:
                    tc_id = tc.get('id', '') if isinstance(tc, dict) else ''
                    if tc_id and tc_id not in seen_tc_ids:
                        seen_tc_ids.add(tc_id)
                        unique_calls.append(tc)
                    elif tc_id:
                        print(f"[_generate_resumable] 去重 assistant tool_call id: {tc_id}", flush=True)
                if unique_calls:
                    m['tool_calls'] = unique_calls
                else:
                    del m['tool_calls']
        # ★ 去重 tool 消息的 tool_call_id（MiniMax 也检查 tool 消息中的重复）
        _dup_tool_msgs = 0
        _seen_tool_ids = set()
        for m in messages:
            if isinstance(m, dict) and m.get('role') == 'tool' and m.get('tool_call_id'):
                _tid = m['tool_call_id']
                if _tid in _seen_tool_ids:
                    m['_remove'] = True
                    _dup_tool_msgs += 1
                    print(f"[_generate_resumable] 去重 tool 消息 tool_call_id: {_tid}", flush=True)
                else:
                    _seen_tool_ids.add(_tid)
        if _dup_tool_msgs > 0:
            messages[:] = [m for m in messages if not m.get('_remove')]
            print(f"[_generate_resumable] 移除 {_dup_tool_msgs} 条重复 tool 消息", flush=True)
        # ★ 清理空 tool_calls:[] 数组 — DeepSeek API 拒绝 empty array
        # ★ 确保 reasoning_content 传递给后续请求(DeepSeek thinking模式要求)
        _has_any_reasoning = not is_longcat and any(
            isinstance(m, dict) and m.get('role') == 'assistant' and (m.get('reasoning_content') or m.get('reasoning'))
            for m in messages
        )
        for m in messages:
            if isinstance(m, dict) and m.get('role') == 'assistant' and 'tool_calls' in m:
                if not m['tool_calls'] or len(m['tool_calls']) == 0:
                    del m['tool_calls']
            # DeepSeek: 若对话中有过reasoning, 所有assistant消息都需带reasoning_content
            if _has_any_reasoning and isinstance(m, dict) and m.get('role') == 'assistant':
                if 'reasoning_content' not in m:
                    m['reasoning_content'] = m.get('reasoning', '')
        _normalize_openai_tool_turns(messages, '_generate_resumable')
        # Provider-neutral parameter adaptation keeps chat/subagent behavior aligned.
        adapted_request = dict(request)
        adapted_request.update({'model': model, 'messages': messages, 'stream': True})
        _provider_info, params = prepare_openai_request(adapted_request)
        params.setdefault('temperature', request.get('temperature', 0.7))
        if 'max_tokens' not in params and 'max_completion_tokens' not in params:
            params['max_tokens'] = request.get('max_tokens', 4096)
        if is_longcat:
            params['extra_body'] = {'thinking': _longcat_thinking(request)}
        elif isinstance(request.get('extra_body'), dict) and request.get('extra_body'):
            params['extra_body'] = request['extra_body']
        # stream_options is forwarded only when explicitly requested; custom OpenAI-compatible
        # gateways frequently reject unknown fields even when official providers accept them.

        print(f"[_generate_resumable] Calling API...", flush=True)
        # ★ 诊断: 打印消息摘要检测重复 tool_call_id
        _tc_diag = {}
        for _mi, _mm in enumerate(messages):
            _tc_ids = []
            if isinstance(_mm, dict) and _mm.get('role') == 'assistant' and 'tool_calls' in _mm:
                _tc_ids = [tc.get('id','?') for tc in _mm['tool_calls'] if isinstance(tc, dict)]
            elif isinstance(_mm, dict) and _mm.get('role') == 'tool' and _mm.get('tool_call_id'):
                _tc_ids = [_mm['tool_call_id']]
            if _tc_ids:
                for _tid in _tc_ids:
                    if _tid in _tc_diag:
                        print(f"[_generate_resumable] ⚠️ DUPLICATE tool_call_id '{_tid}' at msg[{_mi}] (first at msg[{_tc_diag[_tid]}])", flush=True)
                    else:
                        _tc_diag[_tid] = _mi
            _role = _mm.get('role','?') if isinstance(_mm, dict) else '?'
            _content_size = len(str(_mm.get('content') or '')) if isinstance(_mm, dict) else 0
            print(f"[_generate_resumable] msg[{_mi}] role={_role} tool_call_count={len(_tc_ids)} content_size={_content_size}", flush=True)
        _tc_by_index = {}  # ★ 按index合并增量tool_call delta
        _tc_order = []     # 保持顺序
        # ★ MiniMax/DeepSeek 内联思考标签提取状态机（流式分块处理，含跨chunk边界保护）
        _think_buf = ''      # 暂存 (think)...(endthink) 跨chunk不完整块
        _in_think = False    # 是否在 (think) 块内
        _chunk_carry = ''    # ★ 跨chunk边界保护: 上一chunk尾部可能是不完整标签前缀
        _TAG_MAX = 10        # max(len('(endthink)'), len('</think>'), len('(think)'), len('<think>'))
        _TAGS = ('(think)', '(endthink)', '<think>', '</think>')

        def _iter_provider_chunks():
            _last_error = None
            for _attempt in range(5):
                if _is_cancelled(stream_id):
                    return
                try:
                    for _chunk in client.chat.completions.create(**params):
                        if _is_cancelled(stream_id):
                            return
                        yield _chunk
                    return
                except Exception as _stream_exc:
                    if _is_cancelled(stream_id):
                        return
                    _last_error = classify_provider_error(_stream_exc, provider=_provider_info.name)
                    # Never replay after user-visible output: doing so could duplicate text, tools, or billing.
                    if full or reasoning or _tc_by_index or not _last_error.retryable or _attempt >= 4:
                        raise _last_error
                    print(f"[_generate_resumable] provider retry {_attempt + 2}/5 code={_last_error.code.value}", flush=True)
                    _retry_after = 0.0
                    if isinstance(_last_error.details, dict):
                        try:
                            _retry_after = float(_last_error.details.get('retry_after') or 0)
                        except (TypeError, ValueError):
                            _retry_after = 0.0
                    _delay = max(_retry_after, min(10.0, 1.5 * (2 ** _attempt)))
                    _deadline = time.monotonic() + _delay
                    while time.monotonic() < _deadline:
                        if _is_cancelled(stream_id):
                            return
                        time.sleep(min(0.2, _deadline - time.monotonic()))
            if _last_error is not None:
                raise _last_error

        _saw_finish_reason = False
        _trailing_drain_chunks = 0
        _MAX_TRAILING_DRAIN = 50
        for chunk in _iter_provider_chunks():
            if _is_cancelled(stream_id):
                print(f"[_generate_resumable] CANCELLED stream {stream_id}", flush=True)
                break
            # ★ 防御: 部分提供商(LongCat等)会返回空choices块(仅usage),直接取[0]会 IndexError
            if not chunk.choices:
                if hasattr(chunk, 'usage') and chunk.usage:
                    try: usage = chunk.usage.model_dump()
                    except Exception: pass
                continue
            choice = chunk.choices[0]
            if choice.finish_reason:
                stop_reason = choice.finish_reason
                _saw_finish_reason = True
            delta = choice.delta
            c = (delta.content or '')
            r = (
                getattr(delta, 'reasoning_content', None)
                or getattr(delta, 'reasoning', None)
                or getattr(delta, 'reasoning_text', None)
                or getattr(delta, 'thinking', None)
                or ''
            )
            if not r and hasattr(delta, 'model_extra') and delta.model_extra:
                r = (
                    delta.model_extra.get('reasoning_content')
                    or delta.model_extra.get('reasoning')
                    or delta.model_extra.get('reasoning_text')
                    or delta.model_extra.get('thinking')
                    or ''
                )
            if c:
                # ★ 拼接上一chunk的尾部缓冲以检测跨边界标签
                c = _chunk_carry + c
                _chunk_carry = ''
                _display_c = ''
                _pos = 0
                while _pos < len(c):
                    if _in_think:
                        _end_idx = c.find('(endthink)', _pos)
                        if _end_idx != -1:
                            _think_buf += c[_pos:_end_idx]
                            reasoning += _think_buf
                            _emit('reasoning', {'delta': _think_buf})
                            _think_buf = ''
                            _in_think = False
                            _pos = _end_idx + 10  # len('(endthink)') = 10
                        else:
                            _think_buf += c[_pos:]
                            reasoning += c[_pos:]
                            _emit('reasoning', {'delta': c[_pos:]})
                            _pos = len(c)
                    else:
                        _start_idx = c.find('(think)', _pos)
                        if _start_idx != -1:
                            _display_c += c[_pos:_start_idx]  # 标签前的内容作为正文
                            _pos = _start_idx + 7  # len('(think)') = 7
                            _in_think = True
                        else:
                            _display_c += c[_pos:]
                            _pos = len(c)
                # ★ 跨chunk边界保护: 尾部可能是不完整标签前缀，截留到下一chunk
                if _display_c and not _in_think:
                    _tail_start = max(0, len(_display_c) - _TAG_MAX)
                    for _ti in range(_tail_start, len(_display_c)):
                        _tail = _display_c[_ti:]
                        if any(tag.startswith(_tail) for tag in _TAGS):
                            _chunk_carry = _tail
                            _display_c = _display_c[:_ti]
                            break
                if _display_c:
                    full += _display_c
                    _emit('content', {'delta': _display_c})
            if r:
                reasoning += r
                _emit('reasoning', {'delta': r})
            if delta.tool_calls:
                for tc in delta.tool_calls:
                    idx = getattr(tc, 'index', 0)
                    if idx not in _tc_by_index:
                        _tc = {
                            'id': tc.id or '',
                            'type': 'function',
                            'function': {'name': '', 'arguments': ''}
                        }
                        # ★ 保留 Gemini thought_signature (Google API 要求回传,否则 400)
                        if hasattr(tc, 'thought_signature') and tc.thought_signature:
                            _tc['thought_signature'] = tc.thought_signature
                        if hasattr(tc.function, 'thought_signature') and tc.function.thought_signature:
                            _tc['function']['thought_signature'] = tc.function.thought_signature
                        _tc_by_index[idx] = _tc
                        _tc_order.append(idx)
                    else:
                        _tc = _tc_by_index[idx]
                    # 合并: id取第一次的
                    if tc.id and not _tc['id']:
                        _tc['id'] = tc.id
                    # 合并: name取非空的
                    if tc.function and tc.function.name:
                        _tc['function']['name'] = tc.function.name
                    # 合并: arguments增量拼接
                    if tc.function and tc.function.arguments:
                        arg = tc.function.arguments
                        cur = _tc['function']['arguments']
                        if not cur:
                            _tc['function']['arguments'] = arg
                        elif arg != cur and not cur.endswith(arg):
                            # ★ 去重: 避免重复拼接相同片段
                            _tc['function']['arguments'] = cur + arg
                    # ★ 合并: thought_signature 取第一次非空值
                    if not _tc.get('thought_signature') and hasattr(tc, 'thought_signature') and tc.thought_signature:
                        _tc['thought_signature'] = tc.thought_signature
                    if not _tc['function'].get('thought_signature') and hasattr(tc.function, 'thought_signature') and tc.function.thought_signature:
                        _tc['function']['thought_signature'] = tc.function.thought_signature
                # ★ 发出合并后的快照(前端实时展示)
                _merged = [_tc_by_index[i] for i in _tc_order]
                _emit('tool_call', {'partial': True, 'tools': _merged})
            if hasattr(chunk, 'usage') and chunk.usage:
                try: usage = chunk.usage.model_dump()
                except Exception: pass
            # finish_reason is normally terminal, but several OpenAI-compatible
            # gateways emit it before the final tool-argument delta. Drain a small,
            # bounded tail so web_search calls cannot be cut at ``{"query":`` and
            # the independent usage chunk is not lost. Plain text still exits after
            # one trailing read, preserving the old no-hang behavior.
            if _saw_finish_reason:
                if _tc_by_index and not _tool_calls_are_complete(_tc_by_index):
                    _trailing_drain_chunks += 1
                    if _trailing_drain_chunks >= _MAX_TRAILING_DRAIN:
                        print(f"[_generate_resumable] tool-call tail exceeded {_MAX_TRAILING_DRAIN} chunks", flush=True)
                        break
                    continue
                if usage is None and _trailing_drain_chunks < 1:
                    _trailing_drain_chunks += 1
                    continue
                break

        if _is_cancelled(stream_id):
            # Persist one explicit terminal cancellation. Reconnecting clients must not
            # interpret a finished-but-empty snapshot as an endlessly generating bubble.
            _finalize_cancelled_stream(stream_id)
            with _resumable_lock:
                _resumable.pop(stream_id, None)
            return

        # ★ 最终合并的tool_calls（去重: 按id去重，保留首次出现）
        _seen_tc_order_ids = set()
        _unique_tool_calls = []
        _tc_dup_count = 0
        for i in _tc_order:
            tc = _tc_by_index.get(i)
            if not tc: continue
            tc_id = tc.get('id', '')
            if tc_id and tc_id not in _seen_tc_order_ids:
                _seen_tc_order_ids.add(tc_id)
                _unique_tool_calls.append(tc)
            elif tc_id:
                _tc_dup_count += 1
                print(f"[_generate_resumable] 去重 done tool_call id: {tc_id} (index={i})", flush=True)
        tool_calls = _unique_tool_calls
        if _tc_dup_count > 0:
            print(f"[_generate_resumable] 移除 {_tc_dup_count} 个重复 tool_call", flush=True)
        # 确保每个有id
        for i, tc in enumerate(tool_calls):
            if not tc['id']:
                tc['id'] = f'call_{i}_{int(time.time()*1000)}'
        # ★ 刷新流式状态机缓冲（跨chunk尾部残留 + 未闭合think块）
        if _chunk_carry:
            full += _chunk_carry
        if _in_think and _think_buf:
            reasoning += _think_buf
        # ★ 安全网: 清理 full 中残留的 (think)...(endthink) 标签（流式状态机未覆盖的边界情况）
        import re as _re_tmp
        _inlines = _re_tmp.findall(r'\(think\)([\s\S]*?)\(endthink\)', full)
        if _inlines:
            reasoning += ''.join(_inlines)
            full = _re_tmp.sub(r'', full)
        # 未闭合 (think) 兜底
        _open_m = _re_tmp.search(r'\(think\)([\s\S]*?)$', full)
        if _open_m and len(_open_m.group(1)) < 3000:
            reasoning += _open_m.group(1)
            full = _re_tmp.sub(r'', full)
        if not full.strip() and not reasoning.strip() and not tool_calls:
            # A provider connection may finish without yielding a parseable SSE event.
            # Treat it as a terminal typed error: replaying the already-billable request
            # through the browser HTTP fallback can duplicate both cost and side effects.
            raise AgentRuntimeError(
                ErrorCode.EMPTY_RESPONSE,
                '模型流已结束，但未返回正文、思考或工具调用',
                status=502,
                retryable=False,
            )
        truncated = stop_reason in ('max_tokens', 'length')
        done_data = {'full_text': full.strip(), 'reasoning_text': reasoning.strip(),
                     'tool_calls': tool_calls, 'usage': usage,
                     'stop_reason': stop_reason, 'truncated': truncated}
        _emit('done', done_data)
        _resumable[stream_id]['finished'] = True
        buf.done()
        _complete_task_from_stream(stream_id, 'completed')
        try: q.put(None)
        except Exception: pass

    except Exception as e:
        _stream_error = e if isinstance(e, AgentRuntimeError) else classify_provider_error(e, provider=locals().get('_provider_info').name if locals().get('_provider_info') else '')
        _resumable[stream_id]['error'] = _stream_error.to_dict()
        _emit('error', {'error': _stream_error.to_dict()})
        _resumable[stream_id]['finished'] = True
        buf.done()
        _complete_task_from_stream(stream_id, 'failed')
        try: q.put(None)
        except Exception: pass
    finally:
        _client_to_close = locals().get('client')
        if _client_to_close is not None:
            try:
                _client_to_close.close()
            except Exception:
                pass
        _http_to_close = locals().get('_http_client')
        if _http_to_close is not None and _http_to_close is not getattr(_client_to_close, '_client', None):
            try:
                _http_to_close.close()
            except Exception:
                pass


@app.post("/engine/chat/create")
async def chat_create(request: Request, user_id: str = Query("")):
    """创建可恢复流 — 接收消息，返回 stream_id，后台线程调 OpenAI"""
    try: body = await request.json()
    except Exception: return JSONResponse({"error": "invalid JSON"}, status_code=400)
    if not body.get('api_key'):
        return JSONResponse({"error": "api_key required"}, status_code=400)

    sid = f"stream_{uuid.uuid4().hex[:12]}"
    q = queue.Queue()
    task_id = f"task_{uuid.uuid4().hex[:12]}"
    chat_id = body.get("chat_id", "")
    raw_msg_id = body.get("msg_id", f"msg_{int(time.time()*1000)}")
    msg_id = re.sub(r'[^a-zA-Z0-9_-]', '', str(raw_msg_id or ''))[:128]
    if not msg_id:
        msg_id = f"msg_{int(time.time()*1000)}"
    model = body.get("model", "")
    trace_id = str(body.get('trace_id', '') or request.headers.get('x-oneapichat-trace', ''))[:96]
    _sync_trace('stream_create_received', {
        'trace_id': trace_id, 'chat_id': chat_id, 'msg_id': msg_id,
        'user_id': user_id, 'model': model, 'message_count': len(body.get('messages') or []),
    })
    if not user_id:
        return JSONResponse({"error": {"code": "UNAUTHORIZED", "message": "user_id required"}}, status_code=401)
    runtime_chat_id = chat_id or f"_stream_{msg_id}"

    # ★ DSH-style authoritative single producer / idempotency barrier.
    # A browser retry, a second tab, or a delayed observer must attach to the
    # running server task instead of starting a second model/tool loop.
    try:
        _existing_store = get_chat_store(user_id)
        _existing_task = _existing_store.find_running_task(user_id, chat_id=runtime_chat_id, msg_id=msg_id)
        if not _existing_task and chat_id:
            _existing_task = _existing_store.find_running_task(user_id, chat_id=chat_id)
        if _existing_task:
            _existing_sid = str(_existing_task.get('stream_id') or '')
            _existing_msg = str(_existing_task.get('msg_id') or msg_id)
            _existing_task_id = str(_existing_task.get('task_id') or '')
            _sync_trace('stream_create_attached_existing', {
                'trace_id': trace_id, 'chat_id': chat_id, 'msg_id': _existing_msg,
                'stream_id': _existing_sid, 'task_id': _existing_task_id,
                'user_id': user_id, 'model': _existing_task.get('model') or model,
            })
            if chat_id and _existing_sid:
                _broadcast_to_user(user_id, 'chat:stream_started', {
                    'chat_id': chat_id, 'stream_id': _existing_sid,
                    'msg_id': _existing_msg, 'model': _existing_task.get('model') or model,
                    'source': str(body.get('client_source') or ''),
                    'ts': time.time(), 'trace_id': trace_id, 'attached': True,
                })
            return {
                'stream_id': _existing_sid,
                'task_id': _existing_task_id,
                'msg_id': _existing_msg,
                'attached': True,
            }
    except Exception as _dedupe_e:
        print(f"[chat_create] single-producer lookup warning: {_dedupe_e}")

    provider_info = detect_provider(
        body.get('base_url', '') or body.get('anthropic_url', ''),
        model,
        anthropic_format=bool(body.get('anthropic_format')),
    )
    # ★ 对齐 DSH：流状态先落盘（磁盘权威，内存只是缓存）。
    #   只要 create 通过参数校验，就在磁盘留下该 msg 记录；后续 ensure_session/create_job/
    #   register_task 即使失败，引擎重启/刷新后按 msg 恢复也不会 "stream not found"，而是
    #   返回磁盘快照（或 producer_lost 可恢复标记），实现 DSH 式"刷新不影响"。
    try:
        _persist_buf = _get_stream_buffer(msg_id)
        _persist_buf.set_meta(sid, chat_id, user_id)
        _persist_buf._save()
    except Exception as _persist_e:
        print(f"[chat_create] early persist warning: {_persist_e}")
    runtime_session = agent_runtime.ensure_session(
        user_id, runtime_chat_id, provider=provider_info.name, model=model, origin='chat'
    )
    runtime_session_id = runtime_session['session_id']
    latest_user_message = next((
        item for item in reversed(body.get('messages') or [])
        if isinstance(item, dict) and item.get('role') == 'user'
    ), None)
    agent_runtime.append_event(runtime_session_id, 'request/header', {
        'stream_id': sid, 'task_id': task_id, 'message_id': msg_id,
        'provider': provider_info.name, 'model': model,
    })
    agent_runtime.append_event(runtime_session_id, 'turn/start', {
        'turn_id': task_id, 'stream_id': sid, 'message_id': msg_id,
    })
    agent_runtime.append_event(runtime_session_id, 'step/start', {
        'step_id': task_id, 'kind': 'llm', 'model': model,
    })
    if latest_user_message:
        agent_runtime.append_event(runtime_session_id, 'user/message', {
            'message': redact_secrets(latest_user_message), 'message_id': latest_user_message.get('id', ''),
        })
    runtime_job = agent_runtime.create_job(user_id, 'chat-stream', {
        'stream_id': sid, 'task_id': task_id, 'message_id': msg_id,
        'chat_id': runtime_chat_id, 'provider': provider_info.name, 'model': model,
    }, session_id=runtime_session_id)
    agent_runtime.update_job(runtime_job['job_id'], 'running', user_id=user_id)
    body['runtime_session_id'] = runtime_session_id
    body['runtime_job_id'] = runtime_job['job_id']

    with _resumable_lock:
        # 在启动后台线程前写全关联信息。旧逻辑先 start()，快速响应可能在
        # msg_id/task_id 注册前完成，导致磁盘缓冲写到 stream_id 且任务永不收尾。
        _resumable[sid] = {
            'queue': q,
            'chunks': [],
            'finished': False,
            'created': time.time(),
            'task_id': task_id,
            'user_id': user_id,
            'msg_id': msg_id,
            'chat_id': chat_id,
            'runtime_session_id': runtime_session_id,
            'runtime_job_id': runtime_job['job_id'],
            'runtime_terminal': False,
            'trace_id': trace_id,
            'created_at': time.time(),
            'finished_at': 0.0,
            'provider': provider_info.name,
            'model': model,
            'base_url': body.get('base_url', ''),
            'anthropic_url': body.get('anthropic_url', ''),
            'anthropic_format': bool(body.get('anthropic_format')),
            # 完成态必须能在没有浏览器在线时提交到会话文件；仅驻留当前后台线程内存，
            # active_tasks 仍只保存 safe_request_snapshot，不落敏感凭据。
            'request_data': body,
            'projection_committed': False,
            'projection_result': {},
        }
        # 清理过期
        now = time.time()
        for k in list(_resumable.keys()):
            v = _resumable[k]
            if now - v.get('created', 0) > _RESUMABLE_TTL:
                if not v.get('finished'):
                    v['cancel'] = True
                    try: v['queue'].put_nowait(None)
                    except Exception: pass
                    # 生成线程仍可能从上游醒来并访问该条目，先保留到它自行收尾。
                    continue
                # 聚合快照已落盘，完成流无需常驻内存；统一端点仍可按 msg_id 恢复。
                del _resumable[k]

    stream_buf = _get_stream_buffer(msg_id)
    stream_buf.set_meta(sid, chat_id, user_id)
    stream_buf._save()

    # 先注册恢复任务，再允许生成线程完成并调用 _complete_task_from_stream。
    # 即使调用方漏传 chat_id，也用稳定的内部 stream chat namespace 建立恢复索引，
    # 不能因为缺少可选字段而产生无主任务。
    persist_chat_id = chat_id or runtime_chat_id
    if user_id:
        try:
            store = get_chat_store(user_id)
            claim_res = store.claim_task(task_id, sid, persist_chat_id, msg_id, user_id, model, body)
            if not claim_res.get("claimed"):
                # 并发赢家已存在，附着到现有运行任务上
                _attached_task = claim_res.get("task", {})
                _existing_sid = str(_attached_task.get('stream_id') or sid)
                _existing_task_id = str(_attached_task.get('task_id') or task_id)
                _existing_msg = str(_attached_task.get('msg_id') or msg_id)
                _sync_trace('stream_create_attached_concurrent_claim', {
                    'trace_id': trace_id, 'chat_id': chat_id, 'msg_id': _existing_msg,
                    'stream_id': _existing_sid, 'task_id': _existing_task_id,
                })
                with _resumable_lock:
                    _resumable.pop(sid, None)
                return {
                    'stream_id': _existing_sid,
                    'task_id': _existing_task_id,
                    'msg_id': _existing_msg,
                    'attached': True,
                }
        except Exception as e:
            # Durable ACK is a safety barrier: never acknowledge a stream the recovery
            # index cannot see, otherwise a browser refresh creates an undiscoverable job.
            with _resumable_lock:
                _resumable.pop(sid, None)
            try:
                agent_runtime.update_job(runtime_job['job_id'], 'failed', error={
                    'code': 'PERSISTENCE_UNAVAILABLE', 'message': 'task registration failed'
                }, user_id=user_id)
            except Exception:
                pass
            return JSONResponse({"error": {"code": "PERSISTENCE_UNAVAILABLE", "message": "无法持久化可恢复任务，请稍后重试"}}, status_code=503)

    threading.Thread(target=_generate_resumable, args=(body, sid), daemon=True).start()
    _sync_trace('stream_registered', {
        'trace_id': trace_id, 'chat_id': chat_id, 'stream_id': sid,
        'task_id': task_id, 'msg_id': msg_id, 'user_id': user_id, 'model': model,
    })

    # ★ 多端同步: 广播流开始事件到其他浏览器/设备
    if user_id and chat_id:
        _broadcast_to_user(user_id, 'chat:stream_started', {
            'chat_id': chat_id, 'stream_id': sid, 'msg_id': msg_id,
            'task_id': task_id, 'model': model, 'source': str(body.get('client_source') or ''),
            'ts': time.time(), 'trace_id': trace_id
        })

    return {
        "stream_id": sid, "task_id": task_id, "msg_id": msg_id,
        "runtime_session_id": runtime_session_id, "runtime_job_id": runtime_job['job_id'],
    }


@app.get("/engine/chat/stream")
async def chat_stream_offset(
    request: Request,
    msg_id: str = Query(""),
    since: int = Query(0),
    stream_id: str = Query(""),
    user_id: str = Query(""),
    snapshot: int = Query(0),
):
    """
    统一 SSE 流端点（支持 offset 断点续传）：
    - msg_id + since: 从磁盘补发 since 之后的 chunks，再连接实时流
    - stream_id: 兼容旧 ResumeStream（直接消费 live stream）
    """
    if not msg_id and not stream_id:
        return JSONResponse({"error": "msg_id or stream_id required"}, status_code=400)

    # 同时给出 stream_id + msg_id 时优先使用精确流，并允许在引擎重启后
    # 仅凭 msg_id 从磁盘快照恢复最终显示。
    s = None
    if stream_id:
        with _resumable_lock:
            s = _resumable.get(stream_id)
        if s and s.get('user_id') != user_id:
            return JSONResponse({"error": "stream not found", "finished": True}, status_code=404)
    if not s and msg_id:
        with _resumable_lock:
            for _sid, _entry in _resumable.items():
                if _entry.get('msg_id') == msg_id and _entry.get('user_id') == user_id:
                    s = _entry
                    stream_id = _sid
                    break

    resolved_msg_id = msg_id or (s.get('msg_id', stream_id) if s else stream_id)
    with _stream_buffers_lock:
        known_in_memory = resolved_msg_id in _stream_buffers
    persisted_path = STREAM_DIR / f"{resolved_msg_id}.json"
    if not s and not known_in_memory and not persisted_path.exists():
        return JSONResponse({"error": "stream not found", "finished": True}, status_code=404)
    buf = _get_stream_buffer(resolved_msg_id)
    trusted_internal = bool(getattr(request.state, 'trusted_internal', False))
    if buf.user_id and buf.user_id != user_id:
        return JSONResponse({"error": "stream not found", "finished": True}, status_code=404)
    if not buf.user_id and not s and not trusted_internal:
        return JSONResponse({"error": "stream ownership unavailable", "finished": True}, status_code=404)

    async def gen():
        # snapshot=1：先发聚合首屏，正文/思考/工具状态无需等待新 token。
        # 快照的 offset 同时是后续实时流起点，因此不会重复拼接历史 chunks。
        if snapshot:
            snap = buf.snapshot()
            if s:
                snap['finished'] = bool(snap.get('finished') or s.get('finished'))
            elif not snap.get('finished'):
                # 引擎重启后上游生产者已丢失，但快照仍是可见事实；不要伪造 provider error，
                # 让客户端完成当前气泡并以 recoverable/producer_lost 标记提示下一步。
                snap['finished'] = True
                snap['recovery_state'] = 'producer_lost'
                snap['recoverable'] = True
                snap['stop_reason'] = snap.get('stop_reason') or 'producer_lost'
                snap['error'] = ''
            yield f"event: snapshot\ndata: {json.dumps(snap, ensure_ascii=False)}\n\n"
            replay_offset = int(snap.get('offset', 0) or 0)
        else:
            replay_offset = max(0, since)
            missed = buf.since(replay_offset)
            for c in missed:
                yield c
                replay_offset += 1
                await asyncio.sleep(0.001)

        # 阶段2: 连接实时流（如果还在生成中）。定期 ping，避免模型长时间
        # 思考/等待上游时代理把一个健康的 SSE 连接判定为空闲并断开。
        if s and not s.get('finished'):
            # ★ 多消费者广播：用索引轮询 s['chunks']（避免 q.get() 瓜分 chunk）
            _live_idx = replay_offset
            _last_emit = time.monotonic()
            while not s.get('finished'):
                _cur_len = len(s['chunks'])
                while _live_idx < _cur_len:
                    yield s['chunks'][_live_idx]
                    _live_idx += 1
                    _last_emit = time.monotonic()
                    await asyncio.sleep(0.001)
                if s.get('finished'):
                    break
                if time.monotonic() - _last_emit >= 10:
                    yield f"event: ping\ndata: {json.dumps({'offset': _live_idx, 'ts': time.time()})}\n\n"
                    _last_emit = time.monotonic()
                await asyncio.sleep(0.05)
            # ★ 发送finished=true后可能残留的chunks(含done事件)
            _cur_len = len(s['chunks'])
            while _live_idx < _cur_len:
                yield s['chunks'][_live_idx]
                _live_idx += 1
                await asyncio.sleep(0.001)

    return StreamingResponse(gen(), media_type="text/event-stream",
                            headers={"X-Accel-Buffering": "no", "Cache-Control": "no-cache"})


@app.get("/engine/chat/stream/{stream_id}")
async def chat_stream_get(stream_id: str, user_id: str = Query("")):
    """消费 SSE 流 — 已完成流返回JSON，活跃流回放chunks+实时推送"""
    s = _resumable.get(stream_id)
    if not s or s.get('user_id') != user_id:
        return JSONResponse({"error": "stream not found", "finished": True}, status_code=404)

    # ★ 流已完成:直接返回JSON(避免SSE的TCP分片导致done事件丢失)
    if s.get('finished'):
        cached = list(s['chunks'])
        for c in reversed(cached):
            if 'event: done' in c or '"full_text"' in c:
                try:
                    _parts = c.split('\ndata: ', 1)
                    if len(_parts) > 1:
                        _data = json.loads(_parts[1].strip())
                        return JSONResponse(_data)
                except Exception:
                    pass
                break
        return JSONResponse({"full_text": "", "tool_calls": [], "finished": True})

    async def gen():
        yield f"event: start\ndata: {json.dumps({'stream_id': stream_id})}\n\n"
        idx = 0
        cached = list(s['chunks'])
        for c in cached:
            yield c
            idx += 1
            await asyncio.sleep(0.001)
        if not s.get('finished'):
            # ★ 多消费者广播：用索引轮询 s['chunks'] 而非 q.get()
            while not s.get('finished'):
                current_len = len(s['chunks'])
                while idx < current_len:
                    yield s['chunks'][idx]
                    idx += 1
                    await asyncio.sleep(0.001)
                if s.get('finished'):
                    break
                await asyncio.sleep(0.05)
        # ★ 发送finished=true后可能残留的chunks(done事件可能在最后一次poll和finished之间写入)
        current_len = len(s['chunks'])
        while idx < current_len:
            yield s['chunks'][idx]
            idx += 1
            await asyncio.sleep(0.001)

    return StreamingResponse(gen(), media_type="text/event-stream",
                            headers={"X-Accel-Buffering": "no", "Cache-Control": "no-cache"})


@app.delete("/engine/chat/stream/{stream_id}")
async def chat_stream_delete(stream_id: str, msg_id: str = Query(""), user_id: str = Query("")):
    """取消/清理指定流 — 标记 cancel, 通知后台生成线程停止
    (用户点停止键后调用, 避免"停止了还在等模型思考完"、继续消耗token)"""
    with _resumable_lock:
        entry = _resumable.get(stream_id)
        if entry and entry.get('user_id') != user_id:
            entry = None
        if not entry and msg_id:
            for candidate in _resumable.values():
                if candidate.get('msg_id') == msg_id and candidate.get('user_id') == user_id:
                    entry = candidate
                    break
        if entry:
            entry['cancel'] = True
    if not entry:
        return JSONResponse({"error": "stream not found"}, status_code=404)
    _finalize_cancelled_stream(stream_id)
    return {"cleaned": True}


def _complete_task_from_stream(stream_id: str, status: str):
    """流结束时标记任务状态并广播"""
    try:
        with _resumable_lock:
            entry = _resumable.get(stream_id, {})
        task_id = entry.get('task_id', '')
        uid = entry.get('user_id', '')
        chat_id = entry.get('chat_id', '')
        runtime_session_id = entry.get('runtime_session_id', '')
        runtime_job_id = entry.get('runtime_job_id', '')
        trace_id = entry.get('trace_id', '')
        _sync_trace('stream_complete_begin', {
            'trace_id': trace_id, 'stream_id': stream_id, 'task_id': task_id,
            'chat_id': chat_id, 'user_id': uid, 'status': status,
        })
        if runtime_session_id and not entry.get('runtime_terminal'):
            with _resumable_lock:
                live = _resumable.get(stream_id, {})
                live['runtime_terminal'] = True
                was_cancelled = bool(live.get('cancel'))
            kind = 'cancelled' if was_cancelled else ('interrupted' if status == 'interrupted' else ('completed' if status == 'completed' else 'error'))
            try:
                agent_runtime.flush_stream_delta(runtime_session_id)
                if kind in ('error', 'interrupted'):
                    agent_runtime.append_event(runtime_session_id, 'session/error', {
                        'error': {'code': kind.upper(), 'message': f'stream {kind}', 'retryable': kind == 'interrupted'}
                    })
                agent_runtime.append_event(runtime_session_id, 'step/end', {'reason': {'kind': kind}})
                agent_runtime.append_event(runtime_session_id, 'turn/end', {'reason': {'kind': kind}})
                if runtime_job_id:
                    job_status = 'cancelled' if kind == 'cancelled' else ('interrupted' if kind == 'interrupted' else status)
                    agent_runtime.update_job(runtime_job_id, job_status, error=None if job_status == 'completed' else {
                        'code': kind.upper(), 'message': f'stream {kind}'
                    }, user_id=uid or None)
            except Exception as runtime_error:
                print(f"[_complete_task] runtime finalize error: {runtime_error}")
        projection = {}
        if uid and chat_id:
            try:
                msg_id = str(entry.get('msg_id') or '')
                snapshot = _get_stream_buffer(msg_id).snapshot() if msg_id else {}
                projection = chat_projection_store.commit_assistant(
                    user_id=uid,
                    chat_id=chat_id,
                    msg_id=msg_id,
                    request_data=entry.get('request_data') or {},
                    content=snapshot.get('full_text') or '',
                    reasoning=snapshot.get('reasoning_text') or '',
                    tool_calls=snapshot.get('tool_calls') or [],
                    usage=snapshot.get('usage') or {},
                    error=snapshot.get('error') if status != 'completed' else None,
                    stop_reason=snapshot.get('stop_reason') or '',
                    truncated=bool(snapshot.get('truncated')),
                    model=str(entry.get('model') or ''),
                )
                with _resumable_lock:
                    live = _resumable.get(stream_id, {})
                    live['projection_committed'] = bool(projection.get('ok'))
                    live['projection_result'] = projection
                _sync_trace('stream_projection_committed', {
                    'trace_id': trace_id, 'stream_id': stream_id, 'chat_id': chat_id,
                    'user_id': uid, 'msg_id': msg_id, 'ok': bool(projection.get('ok')),
                    'revision': projection.get('revision', 0), 'msg_count': projection.get('msg_count', 0),
                })
            except Exception as projection_error:
                projection = {'ok': False, 'error': type(projection_error).__name__}
                print(f"[_complete_task] chat projection error={type(projection_error).__name__}")
        if task_id and uid:
            store = get_chat_store(uid)
            # 完成态只有在会话投影提交成功后才退出 active_tasks；若磁盘暂时不可写，保留
            # recoverable 入口，客户端可继续从 StreamBuffer 恢复，绝不出现“回复已完成但找不到”。
            persisted_status = status if projection.get('ok') or not chat_id else 'recoverable'
            store.complete_task(task_id, persisted_status)
            # Commit-before-broadcast: 收到 done 的客户端回源时，权威文件已包含最终回复。
            _broadcast_to_user(uid, 'chat:stream_done', {
                'task_id': task_id, 'stream_id': stream_id, 'msg_id': entry.get('msg_id', ''),
                'chat_id': chat_id, 'status': status, 'persisted': bool(projection.get('ok')),
                'revision': projection.get('revision', 0), 'updated_at': projection.get('updated_at', 0),
                'msg_count': projection.get('msg_count', 0), 'ts': time.time(), 'trace_id': trace_id
            })
            _sync_trace('stream_complete_broadcasted', {
                'trace_id': trace_id, 'stream_id': stream_id, 'task_id': task_id,
                'chat_id': chat_id, 'user_id': uid, 'status': status,
            })
    except Exception as e:
        print(f"[_complete_task] error={type(e).__name__}")


@app.get("/engine/tasks/active")
async def active_tasks(user_id: str = Query("")):
    """获取用户活跃任务（用于跨浏览器刷新恢复）"""
    if not user_id:
        return JSONResponse({"ok": True, "tasks": []})
    store = get_chat_store(user_id)
    with _resumable_lock:
        live_stream_ids = {
            sid for sid, entry in _resumable.items()
            if not entry.get('finished') and not entry.get('cancel')
        }
    reconciled = store.reconcile_active_tasks(user_id, live_stream_ids)
    if reconciled:
        with _resumable_lock:
            for item in reconciled:
                if item.get('status') != 'failed':
                    continue
                entry = _resumable.get(item.get('stream_id', ''))
                if entry and not entry.get('finished'):
                    entry['cancel'] = True
                    try:
                        entry['queue'].put_nowait(None)
                    except Exception:
                        pass
    tasks = store.get_active_tasks(user_id)
    return {"ok": True, "tasks": tasks}


@app.post("/engine/tasks/{task_id}/abandon")
async def abandon_active_task(task_id: str, user_id: str = Query("")):
    """客户端确认无法恢复后终止任务，避免刷新时反复接续。"""
    if not user_id:
        return JSONResponse({"ok": False, "error": "user_id required"}, status_code=400)
    store = get_chat_store(user_id)
    changed = store.abandon_task(task_id, user_id)
    with _resumable_lock:
        for entry in _resumable.values():
            if entry.get('task_id') == task_id and entry.get('user_id') == user_id:
                entry['cancel'] = True
                try:
                    entry['queue'].put_nowait(None)
                except Exception:
                    pass
                break
    return {"ok": True, "changed": changed}


# ═══════════════════════════════════════════════════════════════
# SSE 事件总线 — 用户级实时推送通道（跨浏览器同步）
# ═══════════════════════════════════════════════════════════════

_user_event_queues: dict = {}  # {user_id: [(event_loop, asyncio.Queue), ...]}
_user_event_queues_lock = threading.Lock()
_user_agent_modes: dict = {}  # ★ 多端同步: 存储每个用户的 agent 模式状态
# SSE 事件短期回放缓存：EventSource 重连会携带 Last-Event-ID，避免断线窗口丢失完成/错误事件。
_user_event_history: dict[str, deque] = {}
_user_event_seq: dict[str, int] = {}
_USER_EVENT_HISTORY_LIMIT = 256


def _resolve_sse_replay_cursor(raw_last_id, current_cursor: int) -> tuple[int, bool]:
    """Resolve replay start without treating a fresh page load as cursor zero.

    EventSource sends Last-Event-ID only when reconnecting the same stream. A new
    page has no cursor and must start at the current tail; otherwise every refresh
    replays the whole bounded history and re-fires old UI notifications.
    """
    if raw_last_id is None or str(raw_last_id).strip() == "":
        return max(0, int(current_cursor or 0)), False
    try:
        return max(0, int(raw_last_id)), True
    except (TypeError, ValueError):
        return max(0, int(current_cursor or 0)), False


def _broadcast_to_user(user_id: str, event_type: str, data: dict):
    """向指定用户的所有活跃 SSE 连接推送事件，并保存短期游标回放。"""
    if not user_id:
        return
    with _user_event_queues_lock:
        seq = _user_event_seq.get(user_id, 0) + 1
        _user_event_seq[user_id] = seq
        sse_payload = (
            f"id: {seq}\n"
            f"event: {event_type}\n"
            f"data: {json.dumps(data, ensure_ascii=False)}\n\n"
        )
        history = _user_event_history.setdefault(
            user_id, deque(maxlen=_USER_EVENT_HISTORY_LIMIT)
        )
        history.append((seq, sse_payload))
        queues = list(_user_event_queues.get(user_id, []))
    _sync_trace('sse_broadcast', {
        'trace_id': data.get('trace_id', '') if isinstance(data, dict) else '',
        'event_type': event_type,
        'chat_id': data.get('chat_id', '') if isinstance(data, dict) else '',
        'user_id': user_id,
        'source': data.get('source', '') if isinstance(data, dict) else '',
        'seq': seq,
        'subscriber_count': len(queues),
        'history_count': len(history),
    })
    delivered = 0
    for loop, q in queues:
        try:
            # This helper is also called by model/background threads.  Capturing
            # the loop when the SSE subscriber connects avoids creating an
            # un-awaited coroutine against the caller's (or a missing) loop.
            if not loop.is_closed():
                loop.call_soon_threadsafe(q.put_nowait, sse_payload)
                delivered += 1
        except RuntimeError:
            pass
    _sync_trace('sse_queued', {
        'trace_id': data.get('trace_id', '') if isinstance(data, dict) else '',
        'event_type': event_type,
        'chat_id': data.get('chat_id', '') if isinstance(data, dict) else '',
        'user_id': user_id,
        'seq': seq,
        'delivered_count': delivered,
    })


@app.get("/engine/events")
async def user_events_stream(request: Request, user_id: str = Query(""), agent_mode: str = Query("")):
    """用户级持久 SSE 通道，支持 Last-Event-ID 短期回放。"""
    if not user_id:
        return JSONResponse({"error": "user_id required"}, status_code=400)

    # ★ 多端同步: 存储客户端上报的 agent 模式
    # ★ 关键: 'off' 不能覆盖另一个设备上报的活跃模式(agent/plan/yolo)
    if agent_mode:
        _existing = _user_agent_modes.get(user_id, {})
        _existing_mode = _existing.get('mode', '')
        _existing_ts = _existing.get('ts', 0)
        # 允许更新条件: 1) 新模式非off 2) 无现有模式 3) 现有模式已过期(>30s无心跳视为过期)
        if agent_mode != 'off' or not _existing_mode or _existing_mode == 'off' or (time.time() - _existing_ts > 30):
            _user_agent_modes[user_id] = {'mode': agent_mode, 'ts': time.time()}
    current_mode = _user_agent_modes.get(user_id, {}).get('mode', '')

    # 浏览器自动重连会发送 Last-Event-ID；query 兼容手写客户端/测试客户端。
    # 新页面没有游标，必须从当前队尾开始，不能把历史 256 条事件全部重放。
    raw_last_id = request.headers.get('last-event-id')
    if raw_last_id is None:
        raw_last_id = request.query_params.get('since')

    q = asyncio.Queue()
    queue_entry = (asyncio.get_running_loop(), q)
    connection_id = request.headers.get('x-oneapichat-source', '') or request.query_params.get('source', '')
    with _user_event_queues_lock:
        cursor = _user_event_seq.get(user_id, 0)
        last_event_id, replay_requested = _resolve_sse_replay_cursor(raw_last_id, cursor)
        replay = [payload for seq, payload in _user_event_history.get(user_id, ()) if seq > last_event_id]
        # 在把连接暴露给实时广播前，先在同一锁内排入全部历史事件。这样 seq=N+1
        # 不可能先于回放的 seq<=N 进入队列，客户端状态机始终单调前进。
        for payload in replay:
            q.put_nowait(payload)
        _user_event_queues.setdefault(user_id, []).append(queue_entry)
        subscriber_count = len(_user_event_queues.get(user_id, []))
    _sync_trace('sse_subscribed', {
        'trace_id': request.query_params.get('sync_trace', ''),
        'user_id': user_id,
        'source': connection_id,
        'cursor': cursor,
        'since': last_event_id,
        'replay': replay_requested,
        'replay_count': len(replay),
        'subscriber_count': subscriber_count,
    })
    async def event_gen():
        try:
            connected_data = {
                'user_id': user_id,
                'ts': time.time(),
                'agent_mode': current_mode,
                'cursor': cursor,
                'replay': replay_requested,
                'replay_count': len(replay),
                'since': last_event_id,
            }
            yield f"event: connected\ndata: {json.dumps(connected_data)}\n\n"
            while True:
                try:
                    payload = await asyncio.wait_for(q.get(), timeout=12)
                    yield payload
                except asyncio.TimeoutError:
                    yield f"event: heartbeat\ndata: {json.dumps({'hb': True, 'cursor': _user_event_seq.get(user_id, 0)})}\n\n"
        except asyncio.CancelledError:
            pass
        finally:
            with _user_event_queues_lock:
                queues = _user_event_queues.get(user_id, [])
                if queue_entry in queues:
                    queues.remove(queue_entry)
                subscriber_count = len(queues)
            _sync_trace('sse_unsubscribed', {
                'trace_id': request.query_params.get('sync_trace', ''),
                'user_id': user_id,
                'source': connection_id,
                'subscriber_count': subscriber_count,
            })

    return StreamingResponse(event_gen(), media_type="text/event-stream",
                            headers={"X-Accel-Buffering": "no", "Cache-Control": "no-cache", "Connection": "keep-alive"})


@app.post("/engine/events/broadcast")
async def events_broadcast(request: Request, user_id: str = Query("")):
    """接收前端发来的广播请求，转发给同用户的其他连接"""
    try:
        body = await request.json()
    except Exception:
        body = {}
    event_type = body.get("event_type", "")
    data = body.get("data", {})
    if not event_type or not user_id:
        return JSONResponse({"ok": False, "error": "event_type and user_id required"}, status_code=400)
    _sync_trace('broadcast_received', {
        'trace_id': data.get('trace_id', '') if isinstance(data, dict) else '',
        'event_type': event_type,
        'chat_id': data.get('chat_id', '') if isinstance(data, dict) else '',
        'user_id': user_id,
        'source': data.get('source', '') if isinstance(data, dict) else '',
    })
    # ★ 多端同步: 存储 agent 模式变更到服务端
    if event_type == 'agent:mode_changed' and data.get('mode'):
        _user_agent_modes[user_id] = {'mode': data['mode'], 'ts': data.get('ts', time.time())}
    _broadcast_to_user(user_id, event_type, data)
    return {"ok": True, "trace_id": data.get('trace_id', '') if isinstance(data, dict) else ''}


# ═══════════════════════════════════════════════════════════════
# WebSocket 网关 — 持久连接，后端管理 AI 流，多端同步
# ═══════════════════════════════════════════════════════════════

class WSConnectionManager:
    """管理用户级别的 WebSocket 连接池"""
    def __init__(self):
        self.connections: dict = {}  # {user_id: [WebSocket, ...]}
        self._lock = threading.Lock()

    def add(self, user_id: str, ws: WebSocket):
        with self._lock:
            self.connections.setdefault(user_id, []).append(ws)

    def remove(self, user_id: str, ws: WebSocket):
        with self._lock:
            if user_id in self.connections:
                try:
                    self.connections[user_id].remove(ws)
                except ValueError:
                    pass

    async def broadcast(self, user_id: str, data: dict):
        """向用户的所有连接广播消息"""
        with self._lock:
            conns = list(self.connections.get(user_id, []))
        for ws in conns:
            try:
                await ws.send_json(data)
            except Exception:
                pass

    def user_online(self, user_id: str) -> bool:
        with self._lock:
            return len(self.connections.get(user_id, [])) > 0


ws_mgr = WSConnectionManager()


class ChatStream:
    """一个 AI 对话流 — 后端管理 LLM 调用，广播 token 到所有 WS 连接"""
    def __init__(self, stream_id: str, user_id: str, chat_id: str, msg_id: str, request: dict):
        self.sid = stream_id
        self.user_id = user_id
        self.chat_id = chat_id
        self.msg_id = msg_id
        self.request = request
        self.chunks: list = []  # 已产生的 token chunks
        self.finished = False
        self.error = None
        self._created_at = time.time()

    async def run(self):
        """调用 LLM API，逐 token 广播并缓存"""
        from openai import OpenAI
        try:
            # ★ 代理配置: 请求级优先
            _httpc = None
            _rp = self.request.get('proxy_url', '')
            if self.request.get('proxy_enabled') and _rp:
                import httpx; _httpc = httpx.Client(proxy=_rp)
            elif _requires_auto_proxy(self.request.get('base_url', '')):
                import httpx; _httpc = httpx.Client(proxy=_PROXY_URL)
            client = OpenAI(
                api_key=self.request.get('api_key', ''),
                base_url=self.request.get('base_url', '').strip().rstrip('/') or None,
                http_client=_httpc
            )
            params = {
                'model': self.request.get('model', 'deepseek-chat'),
                'messages': self.request.get('messages', []),
                'stream': True,
                'temperature': self.request.get('temperature', 0.7),
                'max_tokens': self.request.get('max_tokens', 4096),
            }
            if self.request.get('tools'):
                params['tools'] = self.request['tools']
            _eb = self.request.get('extra_body')
            if isinstance(_eb, dict) and _eb and 'thinking' not in _eb:
                params['extra_body'] = _eb
            if self.request.get('reasoning_effort'):
                params['reasoning_effort'] = self.request['reasoning_effort']
            if self.request.get('thinking_level'):
                params['thinking_level'] = self.request['thinking_level']

            for chunk in client.chat.completions.create(**params):
                delta = chunk.choices[0].delta
                c = delta.content or ''
                r = getattr(delta, 'reasoning_content', '') or ''
                tc_data = None
                if delta.tool_calls:
                    for tc in delta.tool_calls:
                        tc_data = {
                            'index': getattr(tc, 'index', 0),
                            'id': tc.id or f'call_{len(self.chunks)}',
                            'function': {'name': tc.function.name or '', 'arguments': tc.function.arguments or ''}
                        }
                if c:
                    self.chunks.append({'type': 'content', 'delta': c})
                    await ws_mgr.broadcast(self.user_id, {'event': 'content', 'data': {'delta': c, 'stream_id': self.sid}})
                if r:
                    self.chunks.append({'type': 'reasoning', 'delta': r})
                    await ws_mgr.broadcast(self.user_id, {'event': 'reasoning', 'data': {'delta': r, 'stream_id': self.sid}})
                if tc_data:
                    self.chunks.append({'type': 'tool_call', 'data': tc_data})
                    await ws_mgr.broadcast(self.user_id, {'event': 'tool_call', 'data': tc_data})
                if hasattr(chunk, 'usage') and chunk.usage:
                    try:
                        u = chunk.usage.model_dump()
                        self.chunks.append({'type': 'usage', 'data': u})
                    except Exception:
                        pass

            self.finished = True
            full_text = ''.join(str(item.get('delta') or '') for item in self.chunks if item.get('type') == 'content')
            reasoning_text = ''.join(str(item.get('delta') or '') for item in self.chunks if item.get('type') == 'reasoning')
            projection = chat_projection_store.commit_assistant(
                user_id=self.user_id, chat_id=self.chat_id, msg_id=self.msg_id,
                request_data=self.request, content=full_text, reasoning=reasoning_text,
                model=str(self.request.get('model') or ''),
            )
            await ws_mgr.broadcast(self.user_id, {
                'event': 'done', 'data': {
                    'stream_id': self.sid, 'finished': True, 'persisted': bool(projection.get('ok')),
                    'revision': projection.get('revision', 0), 'updated_at': projection.get('updated_at', 0),
                }
            })
        except Exception as e:
            self.error = str(e)
            self.finished = True
            await ws_mgr.broadcast(self.user_id, {
                'event': 'error', 'data': {'stream_id': self.sid, 'error': str(e)}
            })

    def get_snapshot(self, since: int = 0) -> list:
        """获取 since 之后的 chunks（用于断线重连）"""
        if since >= len(self.chunks):
            return []
        return self.chunks[since:]


# 活跃流注册表 {stream_id: ChatStream}
_active_streams: dict = {}
_active_streams_lock = threading.Lock()


@app.websocket("/engine/ws/{user_id}")
async def ws_chat(ws: WebSocket, user_id: str):
    """WebSocket 聊天网关 — 持久连接，收发 AI 消息"""
    try:
        user_id = require_scope_owner(ws.scope, user_id)
    except AgentRuntimeError:
        await ws.close(code=1008, reason='owner mismatch')
        return
    await ws.accept()
    ws_mgr.add(user_id, ws)
    try:
        while True:
            msg = await ws.receive_json()
            action = msg.get('action', '')

            if action == 'chat':
                # 创建新的 AI 流
                sid = f"ws_{uuid.uuid4().hex[:12]}"
                chat_id = msg.get('chat_id', '')
                msg_id = msg.get('msg_id', '')
                request = msg.get('request', {})
                stream = ChatStream(sid, user_id, chat_id, msg_id, request)
                with _active_streams_lock:
                    _active_streams[sid] = stream
                # 通知客户端流已创建
                await ws.send_json({'event': 'stream_created', 'data': {'stream_id': sid, 'msg_id': msg_id}})
                # 后台执行 LLM 调用
                asyncio.create_task(stream.run())

            elif action == 'resume':
                # 续接已有流：补发 missed chunks
                sid = msg.get('stream_id', '')
                since = msg.get('since', 0)
                with _active_streams_lock:
                    stream = _active_streams.get(sid)
                if stream and stream.user_id == user_id:
                    missed = stream.get_snapshot(since)
                    for chunk in missed:
                        await ws.send_json({'event': chunk['type'], 'data': chunk.get('data', chunk.get('delta', '')), 'stream_id': sid})
                    if stream.finished:
                        await ws.send_json({'event': 'done', 'data': {'stream_id': sid, 'finished': True}})
                    else:
                        await ws.send_json({'event': 'resumed', 'data': {'stream_id': sid, 'since': len(stream.chunks)}})
                else:
                    await ws.send_json({'event': 'error', 'data': {'error': 'stream not found', 'stream_id': sid}})

            elif action == 'ping':
                await ws.send_json({'event': 'pong', 'data': {}})

    except WebSocketDisconnect:
        pass
    except Exception as e:
        print(f"[WS] connection error={type(e).__name__}")
    finally:
        ws_mgr.remove(user_id, ws)


# 定期清理过期流（30 分钟）
def _cleanup_old_streams():
    while True:
        time.sleep(300)
        now = time.time()
        with _active_streams_lock:
            for sid in list(_active_streams.keys()):
                s = _active_streams[sid]
                if s.finished and hasattr(s, '_created_at'):
                    if now - s._created_at > 1800:
                        del _active_streams[sid]

threading.Thread(target=_cleanup_old_streams, daemon=True).start()


# ★ Agent 端点 + 浏览器工具 → engine/agent_endpoints.py
register_agent_endpoints(app, ENGINE_DIR, tool_registry)
from engine.self_api import register_self_endpoints
register_self_endpoints(app, PROJECT_ROOT)

# ★ 记忆系统 v2 → engine/memory_endpoints.py (Phase B: chromadb + FTS5 + 人格)
register_memory_endpoints(app, ENGINE_DIR)

@app.get("/engine/ppt/generate")
def ppt_generate(user_id: str = Query(""), title: str = Query(""), pages: str = Query("[]"), theme: str = Query("default"), filename: str = Query("")):
    """PPT 生成 HTTP 端点"""
    import json as _ppt_json, shutil as _ppt_shutil, os as _ppt_os
    try:
        _ppt_pages = _ppt_json.loads(pages)
    except Exception:
        _ppt_pages = []
    if not _ppt_pages:
        return JSONResponse({"ok": False, "error": "pages 参数为空或格式错误"}, status_code=400)
    try:
        from ppt_engine.build import build_pptx
        _ppt_safe = (filename or 'output').replace('/', '_').replace('\\', '_')
        _ppt_tmp = f"/tmp/ppt_{_ppt_safe}.pptx"
        _ppt_output = build_pptx(_ppt_tmp, title, _ppt_pages, theme)
        _ppt_web_dir = _ppt_os.path.join(PROJECT_ROOT, 'uploads', 'shared')
        _ppt_os.makedirs(_ppt_web_dir, exist_ok=True)
        _ppt_web_path = _ppt_os.path.join(_ppt_web_dir, f"ppt_{_ppt_safe}.pptx")
        _ppt_shutil.copy2(_ppt_output, _ppt_web_path)
        _ppt_size_kb = _ppt_os.path.getsize(_ppt_web_path) // 1024
        _ppt_web_url = f"https://naujtrats.xyz/oneapichat/uploads/shared/ppt_{_ppt_safe}.pptx"
        # ★ 云盘全面结合: 同步到用户 OneAPIChat/generated
        if user_id:
            _import_to_cloudreve(user_id, _ppt_web_path, "generated")
        return {
            "ok": True,
            "download_url": _ppt_web_url,
            "file_size_kb": _ppt_size_kb,
            "pages": len(_ppt_pages),
            "result": f"✅ PPT已生成\\n📥 下载链接: {_ppt_web_url}\\n📦 大小: {_ppt_size_kb} KB | 页数: {len(_ppt_pages)}\\n⚠️ 请直接复制上面的完整链接给用户，不要截断或省略。"
        }
    except ImportError as _ppt_e:
        return JSONResponse({"ok": False, "error": f"PPT生成缺少依赖: {_ppt_e}"}, status_code=500)
    except Exception as _ppt_e:
        return JSONResponse({"ok": False, "error": f"PPT生成失败: {_ppt_e}"}, status_code=500)


# ── 文档生成工具 (Word/Excel/PDF) ──

# ★ 2026-08-03 云盘全面结合: 生成文件同步到用户 Cloudreve 账号 OneAPIChat/{category}
def _import_to_cloudreve(user_id, file_path, category="generated"):
    """调用 PHP 桥接把生成文件导入用户 Cloudreve 云盘（内部 cr_shared 通道, 同机 loopback）"""
    if not user_id or not file_path or not os.path.isfile(file_path):
        return
    try:
        import urllib.parse as _up
        from engine.runtime_auth import get_internal_bridge_secret
        _bridge_secret = get_internal_bridge_secret(PROJECT_ROOT)
        if not _bridge_secret:
            print("[cloudreve] import skipped: bridge credential unavailable", flush=True)
            return
        _url = ("https://127.0.0.1:443/oneapichat/api/cloudreve_api.php"
                "?action=import_file"
                "&user_id=" + _up.quote(str(user_id))
                + "&category=" + _up.quote(category)
                + "&file_path=" + _up.quote(str(file_path)))
        # This is a local HTTPS virtual host whose certificate name is public-host based.
        # The private bridge credential, loopback-only transport and no-proxy policy form
        # the authentication boundary; never put a reusable credential in the URL.
        _r = _http_session.get(
            _url,
            headers={"Host": "naujtrats.xyz", "X-OneAPIChat-Internal": _bridge_secret},
            timeout=120, verify=False, proxies={"http": None, "https": None},
        )
        if _r.status_code != 200:
            print(f"[cloudreve] import failed status={_r.status_code}", flush=True)
    except Exception as _e:
        print(f"[cloudreve] generated-file import failed error={type(_e).__name__}", flush=True)


def _doc_output(safe_name, ext, user_id=""):
    """Copy generated file to web-accessible dir and return download URL."""
    import shutil, os
    web_dir = os.path.join(PROJECT_ROOT, 'uploads', 'shared')
    os.makedirs(web_dir, exist_ok=True)
    tmp_path = f"/tmp/{safe_name}.{ext}"
    web_path = os.path.join(web_dir, f"{safe_name}.{ext}")
    shutil.copy2(tmp_path, web_path)
    os.chmod(web_path, 0o644)
    size_kb = os.path.getsize(web_path) // 1024
    url = f"https://naujtrats.xyz/oneapichat/uploads/shared/{safe_name}.{ext}"
    # ★ 云盘全面结合: 同步到用户 OneAPIChat/generated
    if user_id:
        _import_to_cloudreve(user_id, web_path, "generated")
    return url, size_kb


@app.get("/engine/docx/generate")
def docx_generate(user_id: str = Query(""), title: str = Query(""), content: str = Query("[]"), filename: str = Query("")):
    """Word 文档生成 — content 为 JSON 数组 [{type:'h1'|'h2'|'p'|'bullet', text:''}]"""
    import json as _j, os as _os
    try:
        sections = _j.loads(content)
    except Exception:
        sections = []
    if not sections and title:
        sections = [{"type": "h1", "text": title}, {"type": "p", "text": ""}]
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        doc = Document()
        doc.styles['Normal'].font.size = Pt(11)
        for sec in sections:
            t = sec.get('text', '') if isinstance(sec, dict) else str(sec)
            typ = sec.get('type', 'p') if isinstance(sec, dict) else 'p'
            if typ in ('h1', 'heading1'): doc.add_heading(t, level=1)
            elif typ in ('h2', 'heading2'): doc.add_heading(t, level=2)
            elif typ in ('h3', 'heading3'): doc.add_heading(t, level=3)
            elif typ == 'bullet':
                for line in t.split('\n') if isinstance(t, str) else [t]:
                    doc.add_paragraph(line, style='List Bullet')
            else:
                para = doc.add_paragraph(t)
                if 'bold' in sec and sec['bold']:
                    for run in para.runs: run.bold = True
        safe = (filename or title or 'document').replace('/', '_').replace('\\', '_')[:60]
        tmp = f"/tmp/{safe}.docx"
        doc.save(tmp)
        url, size = _doc_output(safe, 'docx', user_id)
        return {"ok": True, "download_url": url, "file_size_kb": size, "result": f"✅ Word文档已生成\\n📥 {url}\\n📦 {size} KB"}
    except ImportError as e:
        return JSONResponse({"ok": False, "error": f"Word生成缺少依赖: {e}"}, status_code=500)
    except Exception as e:
        return JSONResponse({"ok": False, "error": f"Word生成失败: {e}"}, status_code=500)


@app.get("/engine/xlsx/generate")
def xlsx_generate(user_id: str = Query(""), title: str = Query("Sheet1"), rows: str = Query("[]"), headers: str = Query("[]"), filename: str = Query("")):
    """Excel 表格生成 — rows 为 JSON 二维数组, headers 为 JSON 字符串数组"""
    import json as _j, os as _os
    try:
        _rows = _j.loads(rows)
        _headers = _j.loads(headers)
    except Exception:
        _rows, _headers = [], []
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook()
        ws = wb.active
        ws.title = title or "Sheet1"
        if _headers:
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", size=11)
            for ci, h in enumerate(_headers, 1):
                cell = ws.cell(row=1, column=ci, value=h)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center')
            start_row = 2
        else:
            start_row = 1
        for ri, row in enumerate(_rows, start_row):
            for ci, val in enumerate(row if isinstance(row, (list, tuple)) else [row], 1):
                ws.cell(row=ri, column=ci, value=val)
        # Auto-width
        for col in ws.columns:
            max_len = 0
            for cell in col:
                try: max_len = max(max_len, len(str(cell.value or '')))
                except: pass
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)
        safe = (filename or title or 'spreadsheet').replace('/', '_').replace('\\', '_')[:60]
        tmp = f"/tmp/{safe}.xlsx"
        wb.save(tmp)
        url, size = _doc_output(safe, 'xlsx', user_id)
        return {"ok": True, "download_url": url, "file_size_kb": size, "result": f"✅ Excel表格已生成\\n📥 {url}\\n📦 {size} KB | 行数: {len(_rows)}"}
    except ImportError as e:
        return JSONResponse({"ok": False, "error": f"Excel生成缺少依赖: {e}"}, status_code=500)
    except Exception as e:
        return JSONResponse({"ok": False, "error": f"Excel生成失败: {e}"}, status_code=500)


@app.get("/engine/pdf/generate")
def pdf_generate(user_id: str = Query(""), title: str = Query(""), content: str = Query("[]"), filename: str = Query("")):
    """PDF 文档生成 — content 为 JSON 数组 [{type:'h1'|'p'|'bullet', text:''}]"""
    import json as _j, os as _os
    try:
        sections = _j.loads(content)
    except Exception:
        sections = []
    if not sections and title:
        sections = [{"type": "h1", "text": title}, {"type": "p", "text": ""}]
    try:
        from fpdf import FPDF
        pdf = FPDF()
        pdf.add_page()
        # Register Chinese font if available
        _cn_font_ok = False
        for _fp in ['/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc', '/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf',
                     '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc', '/usr/share/fonts/truetype/arphic/uming.ttc']:
            if _os.path.exists(_fp):
                pdf.add_font('CN', '', _fp, uni=True)   # 常规
                pdf.add_font('CN', 'B', _fp, uni=True)  # 粗体（同文件，fpdf2 用 style='B' 区分）
                _cn_font_ok = True
                break
        def _set_font(style='', size=10):
            if _cn_font_ok:
                pdf.set_font('CN', style, size)
            else:
                pdf.set_font('Helvetica', style, size)
        for sec in sections:
            t = sec.get('text', '') if isinstance(sec, dict) else str(sec)
            typ = sec.get('type', 'p') if isinstance(sec, dict) else 'p'
            if typ in ('h1', 'heading1'):
                _set_font('B', 18)
                pdf.cell(0, 12, t, ln=True); pdf.ln(4)
            elif typ in ('h2', 'heading2'):
                _set_font('B', 14)
                pdf.cell(0, 10, t, ln=True); pdf.ln(3)
            elif typ == 'bullet':
                _set_font('', 10)
                for line in t.split('\n') if isinstance(t, str) else [t]:
                    pdf.cell(8, 7, '•', ln=0)
                    pdf.multi_cell(0, 7, line)
            else:
                _set_font('', 10)
                pdf.multi_cell(0, 7, t); pdf.ln(2)
        safe = (filename or title or 'document').replace('/', '_').replace('\\', '_')[:60]
        tmp = f"/tmp/{safe}.pdf"
        pdf.output(tmp)
        url, size = _doc_output(safe, 'pdf', user_id)
        return {"ok": True, "download_url": url, "file_size_kb": size, "result": f"✅ PDF文档已生成\\n📥 {url}\\n📦 {size} KB"}
    except ImportError as e:
        return JSONResponse({"ok": False, "error": f"PDF生成缺少依赖: {e}"}, status_code=500)
    except Exception as e:
        return JSONResponse({"ok": False, "error": f"PDF生成失败: {e}"}, status_code=500)

@app.post("/engine/video_edit")
async def video_edit_endpoint(request: Request):
    """视频剪辑 HTTP 端点"""
    try:
        body = await request.json()
        action = body.get("action", "")
        params = body.get("params", {})
        input_path = body.get("input_path", "")
        output_path = body.get("output_path", "/tmp/video_output.mp4")
        if not input_path:
            return JSONResponse({"error": "未提供 input_path"}, status_code=400)
        # ★ 自动转换相对路径为绝对路径(支持上传文件的 URL 格式)
        if not os.path.exists(input_path) and input_path.startswith("/"):
            # 处理 /oneapichat/uploads/... 格式
            if input_path.startswith("/oneapichat/"):
                input_path = PROJECT_ROOT + "/" + input_path.replace("/oneapichat/", "", 1)
            elif input_path.startswith("/uploads/"):
                input_path = PROJECT_ROOT + input_path
            elif input_path.startswith("http"):
                return JSONResponse({"error": "不支持远程URL,请先用 server_exec + curl 下载到服务器"}, status_code=400)
            else:
                input_path = PROJECT_ROOT + input_path
        if not os.path.exists(input_path) and action not in ("tts", "voice"):
            return JSONResponse({"error": f"文件不存在: {input_path}"}, status_code=404)
        if action == "info":
            cmd = ["ffprobe", "-v", "quiet", "-print_format", "json", "-show_format", "-show_streams", input_path]
            r = subprocess.run(cmd, capture_output=True, text=True, timeout=30)
            return {"result": r.stdout}
        elif action == "trim":
            start = params.get("start", 0)
            end = params.get("end", None)
            from moviepy import VideoFileClip
            clip = VideoFileClip(input_path)
            if end: clip = clip.subclipped(start, end)
            else: clip = clip.subclipped(start)
            clip.write_videofile(output_path, codec="libx264", audio_codec="aac")
            clip.close()
            return {"result": f"裁剪完成: {output_path}"}
        elif action == "speed":
            factor = float(params.get("factor", 1.0))
            from moviepy import VideoFileClip, vfx
            clip = VideoFileClip(input_path)
            clip = clip.with_effects([vfx.MultiplySpeed(factor)])
            clip.write_videofile(output_path, codec="libx264", audio_codec="aac")
            clip.close()
            return {"result": f"调速完成 (x{factor}): {output_path}"}
        elif action == "resize":
            width = params.get("width", 0); height = params.get("height", 0)
            from moviepy import VideoFileClip
            clip = VideoFileClip(input_path)
            if width and height: clip = clip.resized((width, height))
            elif width: clip = clip.resized(width=width)
            elif height: clip = clip.resized(height=height)
            clip.write_videofile(output_path, codec="libx264", audio_codec="aac")
            clip.close()
            return {"result": f"缩放完成: {output_path}"}
        elif action == "audio":
            from moviepy import VideoFileClip
            clip = VideoFileClip(input_path)
            audio_output = output_path + ".mp3" if not output_path.endswith(".mp3") else output_path
            if clip.audio: clip.audio.write_audiofile(audio_output)
            clip.close()
            return {"result": f"音频提取完成: {audio_output}"}
        elif action == "concat":
            files = params.get("files", [])
            if not files: return JSONResponse({"error": "concat 需要 files 数组"}, status_code=400)
            from moviepy import VideoFileClip, concatenate_videoclips
            clips = [VideoFileClip(f) for f in files]
            final = concatenate_videoclips(clips, method="compose")
            final.write_videofile(output_path, codec="libx264", audio_codec="aac")
            for c in clips: c.close()
            final.close()
            return {"result": f"拼接完成: {output_path}"}
        elif action == "overlay":
            overlay_path = params.get("overlay_path", "")
            if not overlay_path or not os.path.exists(overlay_path):
                return JSONResponse({"error": "overlay_path 无效"}, status_code=400)
            x, y = params.get("x", 10), params.get("y", 10)
            scale = params.get("scale", 0.3)
            from moviepy import VideoFileClip, CompositeVideoClip
            clip = VideoFileClip(input_path)
            ov = VideoFileClip(overlay_path).resized(scale).with_position((x, y))
            final = CompositeVideoClip([clip, ov])
            final.write_videofile(output_path, codec="libx264", audio_codec="aac")
            clip.close(); ov.close(); final.close()
            return {"result": f"画中画完成: {output_path}"}
        elif action == "text":
            result = _apply_subtitle(input_path, output_path, params)
            return {"result": result}
        elif action == "rotate":
            angle = float(params.get("angle", 90))
            from moviepy import VideoFileClip, vfx
            clip = VideoFileClip(input_path)
            clip = clip.with_effects([vfx.Rotate(angle)])
            clip.write_videofile(output_path, codec="libx264", audio_codec="aac")
            clip.close()
            return {"result": f"旋转完成 ({angle}°): {output_path}"}
        elif action in ("filter", "video_filter"):
            return {"result": _apply_ffmpeg_filter(input_path, output_path, params)}
        elif action in ("transition", "video_transition"):
            return {"result": _apply_ffmpeg_transition(input_path, output_path, params)}
        elif action == "tts":
            return {"result": _apply_tts(params)}
        elif action == "voice":
            audio_path2 = params.get("audio_path", "")
            if not audio_path2:
                tts_result2 = _apply_tts(params)
                if "失败" in tts_result2 or "异常" in tts_result2:
                    return {"error": tts_result2}
                audio_path2 = tts_result2.split(": ")[1].split(" ")[0] if ": " in tts_result2 else "/tmp/tts_output.mp3"
            return {"result": _apply_voice_to_video(input_path, audio_path2, output_path, params)}
        elif action == "compose":
            return {"result": _apply_compose(input_path, output_path, params)}
        elif action == "crop":
            return {"result": _apply_crop(input_path, output_path, params)}
        elif action == "reverse":
            return {"result": _apply_reverse(input_path, output_path, params)}
        elif action == "mute":
            return {"result": _apply_mute(input_path, output_path, params)}
        elif action == "bgm":
            return {"result": _apply_bgm(input_path, output_path, params)}
        elif action == "enhance":
            return {"result": _apply_enhance(input_path, output_path, params)}
        elif action == "gif":
            return {"result": _apply_gif(input_path, output_path, params)}
        elif action == "silent_cut":
            return {"result": _apply_silent_cut(input_path, output_path, params)}
        elif action == "stt":
            return {"result": _apply_stt(input_path, output_path, params)}
        elif action == "stt_to_timeline":
            return {"result": _apply_stt_to_timeline(input_path, output_path, params)}
        elif action == "style":
            return {"result": _apply_subtitle_style(input_path, output_path, params)}
        elif action == "frames":
            # 提取关键帧并返回 base64 数组。支持显式 timestamps，确保可精确读取视频结尾。
            count = max(1, min(120, int(params.get("count", 3))))
            duration = max(0.0, float(params.get("duration", 0) or 0))
            scale = max(160, min(1920, int(params.get("scale", 640))))
            requested = params.get("timestamps")
            timestamps = []
            if isinstance(requested, list):
                for value in requested[:120]:
                    try:
                        ts = max(0.0, float(value))
                        if duration > 0:
                            ts = min(ts, max(0.0, duration - 0.04))
                        if not timestamps or abs(ts - timestamps[-1]) >= 0.03:
                            timestamps.append(round(ts, 3))
                    except (TypeError, ValueError):
                        continue
            if not timestamps:
                effective_duration = duration or 10.0
                if count == 1:
                    timestamps = [max(0.0, effective_duration - 0.08)]
                else:
                    end_ts = max(0.0, effective_duration - 0.08)
                    timestamps = [round(end_ts * i / (count - 1), 3) for i in range(count)]

            import base64 as b64
            frames = []
            actual_timestamps = []
            errors = []
            for ts in timestamps:
                cmd = ["ffmpeg", "-v", "error", "-ss", str(ts), "-i", input_path,
                       "-frames:v", "1", "-vf", f"scale={scale}:-2", "-f", "image2pipe",
                       "-q:v", "3", "-vcodec", "mjpeg", "-"]
                r = subprocess.run(cmd, capture_output=True, timeout=45)
                if r.returncode != 0 or len(r.stdout) < 100:
                    errors.append({"timestamp": ts, "error": r.stderr.decode(errors="replace")[:160]})
                    continue
                frames.append("data:image/jpeg;base64," + b64.b64encode(r.stdout).decode())
                actual_timestamps.append(ts)
            if not frames:
                detail = errors[0]["error"] if errors else "未生成画面帧"
                return JSONResponse({"error": f"截图失败: {detail}"}, status_code=500)
            return {"result": json.dumps({"frames": frames, "timestamps": actual_timestamps,
                                          "count": len(frames), "errors": errors})}
        else:
            return JSONResponse({"error": f"未知操作: {action}"}, status_code=400)
    except ImportError as e:
        return JSONResponse({"error": f"缺少依赖: {str(e)}"}, status_code=503)
    except Exception as e:
        return JSONResponse({"error": f"视频剪辑失败: {str(e)}"}, status_code=500)

# ═══════════════════════════════════════════════════════════════
# RAG (Retrieval Augmented Generation) API — 知识库检索
# ═══════════════════════════════════════════════════════════════

@app.get("/engine/rag/collections")
async def rag_collections(user_id: str = Query("")):
    return rag_list_collections(user_id)

@app.post("/engine/rag/collections")
async def rag_create_col(request: Request, user_id: str = Query("")):
    body = await request.json()
    name = _safe_rag_collection(body.get("name", ""))
    return rag_create_collection(name, user_id)

@app.delete("/engine/rag/collections")
async def rag_delete_col(request: Request, user_id: str = Query("")):
    body = await request.json()
    name = _safe_rag_collection(body.get("name", ""))
    return rag_delete_collection(name, user_id)

@app.get("/engine/rag/knowledge")
async def rag_knowledge(collection: str = Query("default"), user_id: str = Query("")):
    return rag_list_documents(_safe_rag_collection(collection), user_id)

@app.post("/engine/rag/upload")
async def rag_upload(request: Request, user_id: str = Query("")):
    try:
        import base64 as _b64
        body = await request.json()
        collection = _safe_rag_collection(body.get("collection", "default"))
        filename = str(body.get("filename", "upload.txt"))[:255]
        content = body.get("content", "")
        # 支持 base64 编码的二进制文件（PDF/DOCX/XLSX等）
        if not content and body.get("content_base64"):
            try:
                content = _b64.b64decode(body["content_base64"]).decode("utf-8", errors="replace")
            except Exception:
                content = ""
        chunk_size = max(128, min(int(body.get("chunk_size", 512)), 4096))
        chunk_overlap = max(0, min(int(body.get("chunk_overlap", 50)), chunk_size // 2))
        api_key = body.get("api_key", "")
        base_url = body.get("base_url", "")
        embed_model = str(body.get("embed_model", "") or _rag_user_settings(user_id, collection).get("embed_model", ""))

        if not content:
            return {"error": "文档内容不能为空"}

        return rag_upload_document(collection, filename, content, user_id,
                                   chunk_size, chunk_overlap, api_key, base_url, embed_model)
    except Exception as e:
        return {"error": str(e)}

@app.get("/engine/rag/search")
async def rag_search_endpoint(q: str = Query(""), collection: str = Query("default"),
                               top_k: int = Query(5), user_id: str = Query(""),
                               api_key: str = Query(""), base_url: str = Query(""),
                               embed_model: str = Query("")):
    if not q:
        return {"results": [], "error": "查询不能为空"}
    collection = _safe_rag_collection(collection)
    top_k = max(1, min(int(top_k), 50))
    embed_model = embed_model or _rag_user_settings(user_id, collection).get("embed_model", "")
    return rag_search(q, collection, top_k, user_id, api_key, base_url, embed_model)

@app.post("/engine/rag/search")
async def rag_search_post(request: Request, user_id: str = Query("")):
    body = await request.json()
    q = body.get("q", body.get("query", ""))
    collection = _safe_rag_collection(body.get("collection", "default"))
    top_k = max(1, min(int(body.get("top_k", 5)), 50))
    api_key = body.get("api_key", "")
    base_url = body.get("base_url", "")
    embed_model = body.get("embed_model", "") or _rag_user_settings(user_id, collection).get("embed_model", "")
    if not q:
        return {"results": [], "error": "查询不能为空"}
    return rag_search(q, collection, top_k, user_id, api_key, base_url, embed_model)

@app.delete("/engine/rag/knowledge")
async def rag_delete_doc(request: Request, user_id: str = Query("")):
    doc_id = request.query_params.get("doc_id", "")
    collection = _safe_rag_collection(request.query_params.get("collection", "default"))
    # 也支持 JSON body
    if not doc_id:
        try:
            body = await request.json()
            doc_id = body.get("doc_id", "")
            collection = _safe_rag_collection(body.get("collection", collection))
        except Exception:
            pass
    if not doc_id:
        return {"error": "doc_id 不能为空"}
    return rag_delete_document(doc_id, collection, user_id)

def _safe_rag_collection(value: str) -> str:
    text = str(value or "default").strip()
    if not re.fullmatch(r"[\w\- \u4e00-\u9fff]{1,80}", text, re.UNICODE):
        raise HTTPException(400, "invalid collection")
    return text


def _rag_user_settings(user_id: str, collection: str) -> dict:
    if not user_id:
        raise AgentRuntimeError(ErrorCode.UNAUTHORIZED, "authenticated user required", status=401)
    data = get_ns("rag_config", user_id).get()
    collections = data.get("collections", {}) if isinstance(data, dict) else {}
    selected = collections.get(collection, {}) if isinstance(collections, dict) else {}
    return selected if isinstance(selected, dict) else {}


@app.get("/engine/rag/embed_config")
async def rag_embed_config(request: Request):
    """Return the authenticated user's collection-scoped embedding settings."""
    user_id = request.query_params.get("user_id", "")
    collection = _safe_rag_collection(request.query_params.get("collection", "default"))
    config_path = Path(PROJECT_ROOT) / "config" / ".mmx_config.json"
    global_cfg = _load_json(config_path) if config_path.exists() else {}
    selected = _rag_user_settings(user_id, collection)
    configured = bool(global_cfg.get("api_key") or global_cfg.get("mmx_api_key"))
    return {
        "embed_model": selected.get("embed_model") or global_cfg.get("embed_model", "text-embedding-3-small"),
        "embed_api_base": "",
        "embed_api_key": "***" if configured else "",
        "embed_api_configured": configured,
        "mode": selected.get("mode", "hybrid"),
        "chunk_size": 512,
        "chunk_overlap": 50,
    }


@app.post("/engine/rag/embed_config")
async def rag_embed_config_post(request: Request):
    """Persist one user's settings and optionally re-embed only their documents."""
    user_id = request.query_params.get("user_id", "")
    collection = _safe_rag_collection(request.query_params.get("collection", "default"))
    embed_model = str(request.query_params.get("embed_model", "")).strip()[:120]
    mode = str(request.query_params.get("mode", "hybrid")).strip().lower()
    if mode not in {"hybrid", "semantic", "keyword"}:
        raise HTTPException(400, "invalid RAG mode")

    store = get_ns("rag_config", user_id)
    settings = store.get()
    if not isinstance(settings, dict):
        settings = {}
    collections = settings.setdefault("collections", {})
    collections[collection] = {"embed_model": embed_model, "mode": mode}
    store.set(settings)

    embedded = 0
    if embed_model:
        config_path = Path(PROJECT_ROOT) / "config" / ".mmx_config.json"
        global_cfg = _load_json(config_path) if config_path.exists() else {}
        docs_file = _docs_path(user_id, collection)
        data = _load_json(docs_file)
        api_key = global_cfg.get("api_key", "") or global_cfg.get("mmx_api_key", "")
        api_base = global_cfg.get("api_base", "") or global_cfg.get("base_url", "")
        for doc in data.get("documents", []):
            for chunk in doc.get("chunks", []):
                try:
                    embedding = _get_embedding(chunk.get("text", ""), api_key, api_base, embed_model)
                    if embedding:
                        chunk["embedding"] = embedding
                        embedded += 1
                except Exception:
                    pass
        if embedded:
            _save_json(docs_file, data)

    return {"success": True, "embedded": embedded, "embed_model": embed_model, "mode": mode}

@app.get("/engine/rag/list_models")
async def rag_list_models():
    return {"models": ["text-embedding-3-small", "text-embedding-3-large", "text-embedding-ada-002"]}


# ═══════════════════════════════════════════════════════════════
# Platform Extract API — 平台特定网页内容提取
# ═══════════════════════════════════════════════════════════════

@app.get("/engine/platform_extract")
async def platform_extract(url: str = Query(""), user_id: str = Query("")):
    """从支持的平台(B站等)提取结构化内容"""
    if not url:
        return {"ok": False, "error": "缺少 url 参数"}
    from engine.browser import is_safe_browser_url
    if not await is_safe_browser_url(url):
        return {"ok": False, "error": "blocked private or invalid URL"}
    try:
        from engine.web_extract import web_extractor as _wex
        bm = None
        try:
            from engine.browser import ensure_browser_connected
            bm = await ensure_browser_connected(user_id)
        except Exception:
            pass
        result = await _wex.extract(url, _http_session, bm)
        return {
            "ok": True,
            "result": result.content,
            "title": result.title,
            "platform": result.platform,
            "structured": result.structured,
        }
    except Exception as e:
        return {"ok": False, "error": str(e)}


# ═══════════════════════════════════════════════════════════════
# Skills API — 可复用提示词模板系统
# ═══════════════════════════════════════════════════════════════

@app.get("/engine/skills/list")
async def skills_list(user_id: str = Query("")):
    """列出用户的所有技能"""
    from engine.skills import list_skills
    return {"skills": list_skills(user_id)}


@app.post("/engine/skills/create")
async def skills_create(request: Request, user_id: str = Query("")):
    """创建新技能"""
    try:
        body = await request.json()
        from engine.skills import create_skill
        return create_skill(user_id, body)
    except Exception as e:
        return {"ok": False, "error": str(e)}


@app.post("/engine/skills/update")
async def skills_update(request: Request, user_id: str = Query("")):
    """更新技能"""
    try:
        body = await request.json()
        name = body.get("name", "")
        from engine.skills import update_skill
        return update_skill(user_id, name, {k: v for k, v in body.items() if k != "name"})
    except Exception as e:
        return {"ok": False, "error": str(e)}


@app.get("/engine/skills/delete")
async def skills_delete(name: str = Query(""), user_id: str = Query("")):
    """删除技能"""
    if not name:
        return {"ok": False, "error": "缺少 name 参数"}
    from engine.skills import delete_skill
    return delete_skill(user_id, name)


@app.post("/engine/skills/run")
async def skills_run(request: Request, user_id: str = Query("")):
    """执行技能: 注入参数 → 返回渲染后的提示词供 Agent 使用"""
    try:
        body = await request.json()
        skill_name = body.get("skill_name", "")
        params = body.get("params", {})
        if not skill_name:
            return {"ok": False, "error": "缺少 skill_name 参数"}

        from engine.skills import get_skill, render_prompt, extract_params

        skill = get_skill(user_id, skill_name)
        if not skill:
            return {"ok": False, "error": f"技能 '{skill_name}' 不存在"}

        prompt = render_prompt(skill.get("prompt_template", ""), params)
        if not prompt:
            return {"ok": False, "error": "技能提示词模板为空"}

        # 返回渲染后的提示词和工具列表, 由调用方(Agent系统的 _execute_tool)创建子代理执行
        return {
            "ok": True,
            "prompt": prompt,
            "tools": skill.get("tools", []),
            "model_tier": skill.get("model_tier", "smart"),
            "max_rounds": int(skill.get("max_rounds", 10)),
        }
    except Exception as e:
        return {"ok": False, "error": str(e)}


# ═══════════════════════════════════════════════════════════════
# ═══════════════════════════════════════════════════════════════
# SRC (StarRailCopilot) REST API — 完整集成端点
# 前端 src-manager.js → /engine/src/* → 控制 StarRailCopilot
# ═══════════════════════════════════════════════════════════════

SRC_DIR = "/home/naujtrats/StarRailCopilot"
SRC_INSTALLED = os.path.isdir(SRC_DIR)

def _safe_src_name(value: str, *, field: str = 'config_name') -> str:
    text = str(value or '').strip()
    if not re.fullmatch(r'[A-Za-z0-9_-]{1,80}', text):
        raise HTTPException(400, f'invalid {field}')
    return text

def _require_src_owner(user_id: str) -> None:
    resource_owners.require('src', str(user_id or ''))

if SRC_INSTALLED and SRC_DIR not in sys.path:
    sys.path.insert(0, SRC_DIR)

# Patch SRC config path resolution (SRC resolves ./config/ relative to cwd, not SRC_DIR)
def _patch_src_path():
    try:
        from module.config import utils as _src_cfg_utils
        _orig_filepath = _src_cfg_utils.filepath_config
        def _patched_filepath(filename, mod_name='alas'):
            filename = _safe_src_name(filename)
            mod_name = _safe_src_name(mod_name, field='module_name')
            if mod_name == 'alas':
                return os.path.join(SRC_DIR, 'config', f'{filename}.json')
            return os.path.join(SRC_DIR, 'config', f'{filename}.{mod_name}.json')
        _src_cfg_utils.filepath_config = _patched_filepath
    except Exception:
        pass

_patch_src_path()

# Thread-safe ProcessManager cache
_src_lock = threading.Lock()
_src_pm_cache = {}

# Task descriptions
SRC_TASK_DESCRIPTIONS = {
    'Alas': '完整调度器，按优先级依次执行所有已启用的任务',
    'Restart': '重启游戏客户端',
    'Dungeon': '刷副本：拟造花萼、侵蚀隧洞、凝滞虚影等，消耗开拓力',
    'Ornament': '刷内圈遗器：差分宇宙·千面xx，消耗开拓力或沉浸器',
    'DailyQuest': '完成每日实训任务，领取活跃度奖励',
    'BattlePass': '领取无名勋礼（大月卡）奖励',
    'Assignment': '收派委托（派遣角色获取材料）',
    'DataUpdate': '更新游戏内数据：信用点、星琼、燃料等资源统计',
    'Freebies': '领取免费奖励：邮件、兑换码、助战奖励',
    'Weekly': '刷历战余响（周本），消耗开拓力',
    'Rogue': '刷模拟宇宙（差分宇宙），可设置祝福/奇物/事件策略',
    'Daemon': '后台托管模式：自动启停模拟器+游戏，循环清体力',
    'PlannerScan': '角色养成规划扫描：读取角色/光锥材料需求',
}

SRC_TASK_GROUPS = [
    {'name': 'Main', 'label': '基础', 'tasks': ['Alas', 'Restart']},
    {'name': 'Daily', 'label': '日常', 'tasks': ['Dungeon', 'Ornament', 'DailyQuest', 'BattlePass', 'Assignment', 'DataUpdate', 'Freebies']},
    {'name': 'Weekly', 'label': '周常', 'tasks': ['Weekly', 'Rogue']},
    {'name': 'Tool', 'label': '工具', 'tasks': ['Daemon', 'PlannerScan']},
]

SRC_DEFAULT_CONFIG = 'src'


def _src_get_config(config_name=None):
    """Lazy-load AzurLaneConfig for a given instance name."""
    if not SRC_INSTALLED:
        return None
    name = config_name or SRC_DEFAULT_CONFIG
    try:
        from module.config.config import AzurLaneConfig
        return AzurLaneConfig(name)
    except Exception:
        return None


def _src_ensure_state():
    """Initialize SRC State multiprocessing.Manager."""
    if not SRC_INSTALLED:
        return
    try:
        from module.webui.setting import State
        if State.manager is None:
            State.init()
    except Exception:
        pass


def _src_have_pm(config_name=None):
    name = config_name or SRC_DEFAULT_CONFIG
    return name in _src_pm_cache


def _src_get_pm(config_name=None):
    """Get or create ProcessManager. Returns None if unavailable."""
    if not SRC_INSTALLED:
        return None
    name = config_name or SRC_DEFAULT_CONFIG
    with _src_lock:
        if name not in _src_pm_cache:
            try:
                _src_ensure_state()
                from module.webui.process_manager import ProcessManager
                _src_pm_cache[name] = ProcessManager.get_manager(name)
            except Exception:
                return None
        return _src_pm_cache.get(name)


def _src_config_dict(config):
    """Extract safe dict from AzurLaneConfig."""
    if config is None:
        return {}
    try:
        return config.data if hasattr(config, 'data') else {}
    except Exception:
        return {}


def _src_task_status(config, task_name):
    """Extract task scheduler status from config data."""
    data = _src_config_dict(config)
    task_data = data.get(task_name, {})
    scheduler = task_data.get('Scheduler', {})
    return {
        'name': task_name,
        'enable': scheduler.get('Enable', False),
        'command': scheduler.get('Command', task_name),
        'next_run': str(scheduler.get('NextRun', '')),
        'description': SRC_TASK_DESCRIPTIONS.get(task_name, ''),
    }


def _src_safe_int(v, default=0):
    try: return int(v)
    except (TypeError, ValueError): return default


def _src_stored_value(config, path_key, default=0):
    """Safe read from config stored data."""
    if config is None:
        return default
    try:
        from module.config.deep import deep_get
        return deep_get(config.data, keys=path_key, default=default)
    except Exception:
        return default


# ── Endpoints ──

@app.get("/engine/src/status")
async def src_status(config_name: str = Query("src"), user_id: str = Query("")):
    _require_src_owner(user_id)
    config_name = _safe_src_name(config_name)
    if not SRC_INSTALLED:
        return {"ok": True, "status": "not_installed", "alive": False, "state": 0,
                "state_label": "not_installed", "config_name": config_name,
                "message": f"StarRailCopilot 未安装。git clone 到 {SRC_DIR}"}
    try:
        if not _src_have_pm(config_name):
            return {"ok": True, "config_name": config_name,
                    "alive": False, "state": 2, "state_label": "stopped"}
        pm = _src_get_pm(config_name)
        if pm is None:
            return {"ok": True, "config_name": config_name,
                    "alive": False, "state": 0, "state_label": "unavailable"}
        alive = pm.alive
        state = pm.state
        labels = {1: 'running', 2: 'stopped', 3: 'error', 4: 'updating'}
        return {"ok": True, "config_name": config_name,
                "alive": alive, "state": state,
                "state_label": labels.get(state, 'unknown')}
    except Exception as e:
        return {"ok": False, "error": str(e), "alive": False, "state": 0, "state_label": "error"}


@app.get("/engine/src/ping")
async def src_ping(user_id: str = Query("")):
    _require_src_owner(user_id)
    return {"ok": True, "installed": SRC_INSTALLED, "dir": SRC_DIR}


@app.get("/engine/src/dashboard")
async def src_dashboard(config_name: str = Query("src"), user_id: str = Query("")):
    _require_src_owner(user_id)
    config_name = _safe_src_name(config_name)
    if not SRC_INSTALLED:
        return {"ok": True, "resources": {}, "message": "SRC 未安装"}
    try:
        config = _src_get_config(config_name)
        resources = {}
        stored_paths = {
            'trailblaze_power': ('Dungeon.DungeonStorage.TrailblazePower',),
            'daily_activity': ('DailyQuest.DailyStorage.DailyActivity',),
            'credit': ('DataUpdate.DataUpdateStorage.Credit',),
            'stellar_jade': ('DataUpdate.DataUpdateStorage.StallerJade',),
            'fuel': ('Dungeon.DungeonStorage.Fuel',),
            'reserved_power': ('Dungeon.DungeonStorage.Reserved',),
            'immersifier': ('Dungeon.DungeonStorage.Immersifier',),
            'echo_of_war': ('Weekly.WeeklyStorage.EchoOfWar',),
            'simulated_universe': ('Rogue.RogueStorage.SimulatedUniverse',),
            'battle_pass_level': ('BattlePass.BattlePassStorage.BattlePassLevel',),
        }
        for key, path in stored_paths.items():
            val = _src_stored_value(config, list(path) + ['value'], None)
            total = _src_stored_value(config, list(path) + ['total'], None)
            t = _src_stored_value(config, list(path) + ['time'], '')
            resources[key] = {
                'value': _src_safe_int(val) if val is not None else 0,
                'total': _src_safe_int(total) if total is not None else None,
                'time': str(t) if t else '',
            }
        return {"ok": True, "resources": resources, "updated_at": str(datetime.now())}
    except Exception as e:
        return {"ok": False, "error": str(e), "resources": {}}


@app.get("/engine/src/tasks")
async def src_tasks(config_name: str = Query("src"), user_id: str = Query("")):
    _require_src_owner(user_id)
    config_name = _safe_src_name(config_name)
    if not SRC_INSTALLED:
        return {"ok": True, "groups": [], "tasks": [], "message": "SRC 未安装"}
    try:
        config = _src_get_config(config_name)
        groups_out = []
        all_tasks = []
        for group in SRC_TASK_GROUPS:
            tasks = [_src_task_status(config, t) for t in group['tasks']]
            groups_out.append({'name': group['name'], 'label': group['label'], 'tasks': tasks})
            all_tasks.extend(tasks)
        return {"ok": True, "groups": groups_out, "tasks": all_tasks}
    except Exception as e:
        return {"ok": False, "error": str(e), "groups": [], "tasks": []}


@app.post("/engine/src/run")
async def src_run(request: Request, user_id: str = Query("")):
    _require_src_owner(user_id)
    if not SRC_INSTALLED:
        return {"ok": False, "error": "SRC 未安装，无法启动"}
    try:
        data = await request.json() if request.headers.get("content-type", "").startswith("application/json") else {}
    except Exception:
        data = {}
    config_name = _safe_src_name(data.get('config_name', SRC_DEFAULT_CONFIG))
    task = data.get('task', 'Alas')
    valid_tasks = [t for g in SRC_TASK_GROUPS for t in g['tasks']]
    if task not in valid_tasks:
        return {"ok": False, "error": f"未知任务: {task}", "valid_tasks": valid_tasks}
    try:
        pm = _src_get_pm(config_name)
        if pm is None:
            return {"ok": False, "error": "SRC 后端不可用（ProcessManager 初始化失败）"}
        if pm.alive:
            return {"ok": False, "error": "已有任务在运行，请先停止"}
        import inflection
        func = inflection.underscore(task) if task != 'Alas' else 'alas'
        pm.start(func)
        return {"ok": True, "config_name": config_name, "task": task, "func": func}
    except Exception as e:
        return {"ok": False, "error": str(e)}


@app.post("/engine/src/stop")
async def src_stop(request: Request, user_id: str = Query("")):
    _require_src_owner(user_id)
    if not SRC_INSTALLED:
        return {"ok": False, "error": "SRC 未安装"}
    try:
        data = await request.json() if request.headers.get("content-type", "").startswith("application/json") else {}
    except Exception:
        data = {}
    config_name = _safe_src_name(data.get('config_name', SRC_DEFAULT_CONFIG))
    try:
        if not _src_have_pm(config_name):
            return {"ok": False, "error": "没有运行中的任务"}
        pm = _src_get_pm(config_name)
        if pm is None or not pm.alive:
            return {"ok": False, "error": "没有运行中的任务"}
        pm.stop()
        return {"ok": True, "config_name": config_name, "message": "已停止"}
    except Exception as e:
        return {"ok": False, "error": str(e)}


@app.get("/engine/src/config/{config_name}")
async def src_get_config(config_name: str, user_id: str = Query("")):
    _require_src_owner(user_id)
    config_name = _safe_src_name(config_name)
    if not SRC_INSTALLED:
        return {"ok": True, "data": {}, "config_name": config_name, "message": "SRC 未安装"}
    try:
        config = _src_get_config(config_name)
        if config is None:
            try:
                from module.config.utils import read_file, filepath_config
                raw = read_file(filepath_config(config_name))
                if raw:
                    return {"ok": True, "config_name": config_name, "data": raw}
            except Exception:
                pass
            return {"ok": False, "error": "配置不存在", "config_name": config_name}
        return {"ok": True, "config_name": config_name, "data": _src_config_dict(config)}
    except Exception as e:
        return {"ok": False, "error": str(e), "config_name": config_name}


@app.put("/engine/src/config/{config_name}")
async def src_set_config(config_name: str, request: Request, user_id: str = Query("")):
    _require_src_owner(user_id)
    config_name = _safe_src_name(config_name)
    if not SRC_INSTALLED:
        return {"ok": False, "error": "SRC 未安装"}
    try:
        data = await request.json()
    except Exception:
        return {"ok": False, "error": "Invalid JSON body"}
    updates = data.get('updates', [])
    if not updates and 'path' in data:
        updates = [{'path': data['path'], 'value': data['value']}]
    if not updates:
        return {"ok": False, "error": "Missing path/value or updates"}
    try:
        config = _src_get_config(config_name)
        if config is None:
            return {"ok": False, "error": "配置不存在"}
        from module.config.deep import deep_set
        applied = []
        for upd in updates:
            path = upd['path']
            value = upd['value']
            deep_set(config.data, keys=path, value=value)
            config.modified[path] = value
            applied.append({'path': path, 'value': value})
        config.save()
        return {"ok": True, "applied": applied}
    except Exception as e:
        return {"ok": False, "error": str(e)}


@app.get("/engine/src/logs")
async def src_logs(config_name: str = Query("src"), limit: int = Query(50), user_id: str = Query("")):
    _require_src_owner(user_id)
    config_name = _safe_src_name(config_name)
    limit = max(1, min(int(limit), 500))
    if not SRC_INSTALLED:
        return {"ok": True, "lines": ["[SRC] StarRailCopilot 未安装"], "count": 1}
    try:
        if not _src_have_pm(config_name):
            return {"ok": True, "lines": [], "count": 0, "note": "未启动"}
        pm = _src_get_pm(config_name)
        if pm is None:
            return {"ok": True, "lines": [], "count": 0}
        renderables = getattr(pm, 'renderables', []) or []
        entries = []
        for r in renderables[-limit:]:
            try:
                if hasattr(r, 'markup'): entries.append(str(r.markup))
                elif hasattr(r, 'text'): entries.append(r.text)
                else: entries.append(str(r))
            except Exception:
                entries.append(str(r))
        return {"ok": True, "lines": entries, "count": len(entries)}
    except Exception as e:
        return {"ok": False, "error": str(e), "lines": []}


@app.get("/engine/src/configs")
async def src_configs(user_id: str = Query("")):
    """List available SRC config instances."""
    _require_src_owner(user_id)
    if not SRC_INSTALLED:
        return {"ok": True, "instances": []}
    try:
        from module.config.utils import alas_instance
        return {"ok": True, "instances": alas_instance()}
    except Exception as e:
        return {"ok": False, "error": str(e), "instances": []}


@app.get("/engine/src/args/{task_name}")
async def src_args(task_name: str, user_id: str = Query("")):
    """Get argument schema for a task."""
    _require_src_owner(user_id)
    task_name = _safe_src_name(task_name, field='task_name')
    if not SRC_INSTALLED:
        return {"ok": False, "error": "SRC 未安装"}
    try:
        from module.config.utils import read_file
        args_file = os.path.join(SRC_DIR, 'module', 'config', 'argument', 'args.json')
        all_args = read_file(args_file)
        if isinstance(all_args, str):
            all_args = json.loads(all_args)
        task_args = all_args.get(task_name, {})
        if not task_args:
            return {"ok": False, "error": f'任务 "{task_name}" 不存在'}
        return {"ok": True, "task": task_name, "args": task_args}
    except Exception as e:
        return {"ok": False, "error": str(e), "args": {}}


if __name__ == "__main__":
    host = os.getenv("ENGINE_HOST", "127.0.0.1")
    port = int(os.getenv("ENGINE_PORT", "8766"))
    print(f"[引擎] 启动 http://{host}:{port}")
    print(f"[引擎] Cron task count={len(cron_store.get())}")
    print(f"[引擎] subagent count={len(agent_store.get())}")
    # HTTP request targets can contain chat/user query values. Keep operational logs
    # metadata-only and rely on the durable redacted event store for diagnostics.
    uvicorn.run(app, host=host, port=port, access_log=False)
